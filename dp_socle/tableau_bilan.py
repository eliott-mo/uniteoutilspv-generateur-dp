"""Lecture du tableau bilan Excel du bureau d'études (lot 2bis, étape B).

L'onglet « 2. Caractéristiques du projet » est organisé en colonnes, une par
indice de révision : le libellé en colonne A, puis `IND05` en B, `IND06` en C…

Deux pièges, tous deux capables de produire des valeurs fausses sans lever la
moindre erreur :

- **repérer les paramètres par leur libellé, jamais par leur numéro de ligne.**
  Le fichier est déjà en version 6 et la mise en page bouge d'une version à
  l'autre ;
- **les valeurs sont typées de façon irrégulière** : `0°` avec le symbole
  degré, `13 / 26` pour deux longueurs de table, des dates en `datetime`. Les
  convertisseurs sont tolérants sur la forme et **échouent explicitement** quand
  le motif ne correspond pas, plutôt que de rendre une valeur par défaut.

Mesuré sur `20260825_SCV_Tableau_Bilan_V6.xlsx` le 03/09/2026.
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .erreurs import ErreurTableauBilan
from .import_be import normaliser

#: Noms d'onglet acceptés pour les caractéristiques du projet, dans l'ordre de
#: préférence. Le tableau a été renommé entre ses versions : « 2.
#: Caractéristiques du projet » en V6, « Projet » en V10. La comparaison passe
#: par `normaliser`, donc la casse et les accents ne comptent pas. Un nom
#: inconnu est refusé en listant les onglets présents, jamais deviné.
ONGLETS_CARACTERISTIQUES = ("2. Caractéristiques du projet", "Projet")
ONGLET_CARACTERISTIQUES = ONGLETS_CARACTERISTIQUES[0]
ONGLET_DIMENSIONS = "Dimensions postes et pieux"
ONGLET_STANDARDS = "Standards UNITe"
#: La même table des cotes, sous le nom que lui donne le classeur des gabarits
#: UNITe dont le tableau bilan la recopie. Relevé le 23/09/2026 : même mise en
#: page, mêmes sections, mêmes en-têtes « Dimensions (…) ».
ONGLET_GABARITS = "Dimensions postes"

#: Le classeur des gabarits UNITe, livré avec l'outil.
#:
#: Défini ici, auprès de `lire_gabarits` qui le lit, et non dans le module du
#: plan PDF qui n'en est qu'un usager : les cotes manquantes d'un tableau bilan
#: s'y complètent aussi désormais, et deux copies du chemin dériveraient.
CHEMIN_GABARITS = Path(__file__).resolve().parent / "ressources" / "gabarits_unite.xlsx"

#: Noms sous lesquels l'onglet des cotes se cherche dans un tableau bilan.
#:
#: Certains tableaux recopient la table des gabarits **sous son nom d'origine** :
#: relevé le 27/09/2026 sur celui d'Auzainvilliers (V3), qui porte un onglet
#: « Dimensions postes » aux mêmes sections et aux mêmes lignes que le
#: « Dimensions postes et pieux » de Saint-Cyr, une colonne plus à droite. Ne
#: chercher que le premier nom privait le dossier de toutes ses cotes DP 4, et
#: l'avertissement accusait un onglet absent qui était là.
ONGLETS_DIMENSIONS = (ONGLET_DIMENSIONS, ONGLET_GABARITS)

#: Motif de l'indice de révision, tel qu'il apparaît en en-tête de colonne et
#: dans le nom du fichier DXF (`20260903_SCV_IND06.dxf`).
#:
#: Le suffixe alphabétique est indispensable, et son absence a coûté cher : un
#: motif sans lui écartait `IND10A` et `IND10B` de l'en-tête du tableau de
#: Sarnois **sans rien dire**, et l'outil concluait que le plan était en avance
#: sur le tableau alors que les deux colonnes étaient là. Ces suffixes
#: distinguent deux variantes d'un même indice — à Sarnois, deux réductions du
#: projet passant sous les 3 MWc.
#:
#: La lettre, ou le nombre lorsqu'il n'y en a pas, ne doit être suivie d'aucun
#: autre caractère alphanumérique : sans quoi `IND06final` rendrait l'indice
#: `IND06F`. Un nom ambigu ne rend rien plutôt que de rendre au hasard.
MOTIF_INDICE = re.compile(r"IND\s*0*(\d+)\s*([A-Za-z])?(?![A-Za-z0-9])", re.IGNORECASE)


def _format_indice(trouve: re.Match) -> str:
    """Forme canonique d'un indice : « IND06 », « IND10A »."""
    lettre = (trouve.group(2) or "").upper()
    return f"IND{int(trouve.group(1)):02d}{lettre}"


# ---------------------------------------------------------------------------
# Convertisseurs
# ---------------------------------------------------------------------------


def _texte(valeur, libelle: str) -> str:
    if valeur is None or not str(valeur).strip():
        raise ErreurTableauBilan(f"« {libelle} » est vide dans le tableau bilan.")
    return str(valeur).strip()


def _nombre(valeur, libelle: str) -> float:
    """Nombre décimal, en tolérant la virgule, les espaces et un symbole d'unité."""
    if isinstance(valeur, bool):
        raise ErreurTableauBilan(f"« {libelle} » vaut un booléen, un nombre est attendu.")
    if isinstance(valeur, (int, float)):
        return float(valeur)
    if valeur is None:
        raise ErreurTableauBilan(f"« {libelle} » est vide, un nombre est attendu.")
    texte = str(valeur).strip().replace(" ", "").replace(" ", "")
    texte = texte.replace(",", ".")
    trouve = re.fullmatch(r"[-+]?\d*\.?\d+\s*[^\d]*", texte)
    if not trouve:
        raise ErreurTableauBilan(
            f"« {libelle} » vaut {valeur!r} : impossible d'y lire un nombre."
        )
    return float(re.match(r"[-+]?\d*\.?\d+", texte).group())


def _entier(valeur, libelle: str) -> int:
    nombre = _nombre(valeur, libelle)
    if abs(nombre - round(nombre)) > 1e-9:
        raise ErreurTableauBilan(
            f"« {libelle} » vaut {valeur!r} : un nombre entier est attendu."
        )
    return int(round(nombre))


def _angle(valeur, libelle: str) -> float:
    """Angle en degrés, avec ou sans le symbole °.

    La valeur numérique seule ne suffit pas à interpréter l'azimut du tableau :
    voir `azimut_brut` dans PARAMETRES. Ce convertisseur ne rend que le nombre,
    et la cellule telle qu'écrite est conservée à côté.
    """
    return _nombre(valeur, libelle)


def _date(valeur, libelle: str) -> _dt.date:
    if isinstance(valeur, _dt.datetime):
        return valeur.date()
    if isinstance(valeur, _dt.date):
        return valeur
    texte = _texte(valeur, libelle)
    for motif in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return _dt.datetime.strptime(texte, motif).date()
        except ValueError:
            continue
    raise ErreurTableauBilan(
        f"« {libelle} » vaut {valeur!r} : format de date non reconnu "
        "(attendu jj/mm/aaaa)."
    )


#: Séparateurs des valeurs composites. Le tableau décrit les projets à deux
#: formats de table par une paire, écrite « 13 / 26 » en V6 et « 13 & 26 » en
#: V10. Les deux sont acceptées ; la forme retenue n'est pas devinée ligne à
#: ligne, elles sont simplement toutes reconnues.
SEPARATEURS_COMPOSITES = r"[/;&+]"

#: Marques d'une valeur explicitement non renseignée. Le tableau V10 met « / »
#: dans la case GCR des indices où elle n'a pas été calculée. Ces marques ne
#: sont acceptées que là où le paramètre le prévoit : ailleurs, une valeur
#: illisible reste une erreur.
MARQUEURS_ABSENCE = frozenset({"/", "na", "n/a", "-", "—", "sansobjet"})


def _couple(valeur, libelle: str) -> tuple[int, ...]:
    """Valeur composite « 13 / 26 » ou « 13 & 26 », rendue en tuple d'entiers.

    Le tableau y met les deux longueurs de table du projet. Une valeur simple
    est acceptée et rendue en tuple d'un élément.
    """
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        return (_entier(valeur, libelle),)
    texte = _texte(valeur, libelle)
    morceaux = [m.strip() for m in re.split(SEPARATEURS_COMPOSITES, texte) if m.strip()]
    if not morceaux:
        raise ErreurTableauBilan(f"« {libelle} » vaut {valeur!r} : valeur illisible.")
    return tuple(_entier(m, libelle) for m in morceaux)


def _somme(valeur, libelle: str) -> int:
    """Total d'une valeur composite : « 269 & 27 » vaut 296 tables.

    Un projet à deux formats de table décrit son parc en deux nombres. Les
    contrôles croisés portent sur le total, et la chaîne d'origine est conservée
    à côté — voir `nb_tables_brut` — pour que la décomposition reste lisible.
    """
    return sum(_couple(valeur, libelle))


def _nombre_ou_absent(valeur, libelle: str) -> float | None:
    """Nombre, ou None quand la case porte une marque d'absence explicite.

    Réservé aux paramètres dont aucun contrôle ne dépend. Ailleurs, rendre None
    en silence sur une case illisible masquerait le problème.
    """
    if valeur is not None and normaliser(valeur) in MARQUEURS_ABSENCE:
        return None
    return _nombre(valeur, libelle)


# ---------------------------------------------------------------------------
# Description des paramètres attendus
# ---------------------------------------------------------------------------

#: (clé, libellé attendu en colonne A, convertisseur, obligatoire)
#: Le libellé est comparé après normalisation tolérante : les espaces de fin
#: (« Phase », « Format ») et les accents ne comptent pas.
PARAMETRES = {
    "generalites": (
        ("phase", "Phase", _texte, True),
        ("date", "Date (jj/mm/aaaa)", _date, True),
        ("type_projet", "Type de projet", _texte, True),
        ("surface_cloturee_ha", "Surface clôturée (ha)", _nombre, True),
        ("lineaire_cloture_m", "Linéaire de clôture (m)", _nombre, True),
        ("nb_portails", "Nombre de portails d'accès", _entier, True),
        ("largeur_portails_m", "Largeur des portails (m)", _nombre, True),
    ),
    "structures": (
        ("type_structure", "Type de structure", _texte, True),
        ("format_table", "Format table", _texte, True),
        ("inclinaison_deg", "Inclinaison (°)", _angle, True),
        ("azimut_deg", "Azimut (°)", _angle, True),
        # La même cellule, lue deux fois : en nombre et telle qu'écrite.
        #
        # Ce champ n'a pas de convention stable d'un projet à l'autre. Relevé
        # le 03/09/2026 sur trois dossiers : « 0° » pour un champ plein sud
        # (0 = sud), « 41° SE » avec la direction en toutes lettres, et
        # « -24,5 » sans direction et de signe contraire pour une orientation
        # de la même famille. Un convertisseur numérique seul avalerait le
        # « SE » sans rien dire : la chaîne d'origine est donc conservée, et
        # c'est elle qui est montrée à l'écran de contrôle.
        ("azimut_brut", "Azimut (°)", _texte, True),
        # Deux formats de table sur un même projet donnent « 269 & 27 » :
        # les contrôles portent sur le total, la chaîne reste lisible à côté.
        ("nb_tables", "Nombre de tables", _somme, True),
        ("nb_tables_brut", "Nombre de tables", _texte, True),
        ("inter_table_m", "Inter-table (m)", _nombre, True),
        ("pitch_m", "Pitch (m)", _nombre, True),
        ("point_bas_m", "Point bas (m)", _nombre, True),
        ("point_haut_m", "Point haut (m)", _nombre, True),
        ("nb_modules_rampant", "Nb de mod en rampant", _entier, True),
        ("nb_modules_longueur", "Nb de mod en longueur", _couple, True),
        ("nb_modules_table", "Nb de mod/table", _couple, True),
        ("type_fondation", "Type de fondation", _texte, True),
        ("nb_pieux", "Nombre de pieux", _entier, True),
    ),
    "modules": (
        ("format_module", "Format", _texte, True),
        ("puissance_unitaire_wc", "Puissance unitaire (Wc)", _nombre, True),
        ("surface_unitaire_m2", "Surface unitaire (m²)", _nombre, True),
        ("nb_modules", "Nombre installé", _entier, True),
        # « / » dans les indices où le GCR n'a pas été calculé. Aucun
        # contrôle n'en dépend, l'absence est donc admise et rendue None.
        ("gcr", "GCR (%)", _nombre_ou_absent, True),
        ("surface_modules_m2", "Surface modules (m²)", _nombre, True),
        ("surface_projetee_m2", "Surface module projetée au sol (m²)", _nombre, True),
        ("puissance_mwc", "Puissance projet (MWc)", _nombre, True),
    ),
    "postes": (
        ("nb_ptr", "Nombre de PTR", _entier, True),
        ("surface_ptr_m2", "Surface des PTR (m²)", _nombre, True),
        ("surface_plateforme_ptr_m2", "Surface plateforme PTR (m²)", _nombre, True),
        ("nb_pdl_ptr", "Nombre de PDL /PTR", _entier, True),
        ("surface_pdl_ptr_m2", "Surface des PDL/PTR (m²)", _nombre, True),
        (
            "surface_plateforme_pdl_ptr_m2",
            "Surface plateforme PDL/PTR (m²)",
            _nombre,
            True,
        ),
        ("nb_pdl", "Nombre de PDL", _entier, True),
        ("surface_pdl_m2", "Surface des PDL (m²)", _nombre, True),
        ("surface_plateforme_pdl_m2", "Surface plateforme PDL (m²)", _nombre, True),
        ("nb_bess", "Nombre de conteneurs de batteries", _entier, True),
        ("surface_bess_m2", "Surface des conteneurs (m²)", _nombre, True),
        (
            "surface_plateforme_bess_m2",
            "Surface plateforme batteries (m²)",
            _nombre,
            True,
        ),
        ("nb_citerne_refroidissement", "Citerne de refroidissement (120m3)", _entier, True),
        ("nb_bac_retention", "Bac de rétention (120m3)", _entier, True),
        ("nb_local_stockage", "Nombre de locaux de stockage matériel", _entier, True),
        (
            "surface_local_stockage_m2",
            "Surface locaux de stockage matériel (m²)",
            _nombre,
            True,
        ),
    ),
    "pistes": (
        # Le total, supplément d'aire de grutage compris. Ce libellé-ci n'est
        # présent qu'une fois, contrairement à son équivalent enherbé.
        ("surface_piste_lourde_m2", "Surface piste lourde (m²)", _nombre, True),
        # Sur la mise en page de Saint-Cyr et de Sarnois, « Surface piste
        # enherbée (m²) » apparaît **deux fois** — une fois en cours de section,
        # une fois au total. Le lecteur refuse par principe un libellé ambigu, et
        # le total se recompose donc depuis les deux lignes qui ne le sont pas.
        #
        # Toutes les mises en page ne découpent pas l'enherbé en interne et
        # externe : celle du tableau d'Auzainvilliers (V3, onglet « Projet »)
        # n'en porte qu'une ligne, et sans ambiguïté. Ces deux paramètres ne
        # sont donc plus obligatoires — le refus arrêtait un dossier entier sur
        # une ligne qui, sur ce projet, était vide — et `_piste_enherbee_unique`
        # lit la ligne unique quand elle existe. Voir `surface_piste_legere_m2`,
        # par où passent tous les consommateurs.
        (
            "surface_piste_legere_interne_m2",
            "Surface piste enherbée interne (m²)",
            _nombre,
            False,
        ),
        (
            "surface_piste_legere_externe_m2",
            "Surface piste enherbée externe (m²)",
            _nombre,
            False,
        ),
    ),
}

#: Libellé de la surface enherbée sur les mises en page qui ne la découpent pas.
LIBELLE_PISTE_ENHERBEE = "Surface piste enherbée (m²)"


def surface_piste_legere_m2(pistes: dict) -> float:
    """Total de piste légère déclaré, quelle que soit la mise en page du tableau.

    Un seul endroit sait que ce total s'écrit tantôt en deux lignes — interne et
    externe — tantôt en une seule. Les contrôles croisés passent par ici plutôt
    que d'additionner eux-mêmes deux clés : c'est la répartition qui change d'un
    tableau à l'autre, pas le total, et un contrôle qui referait l'addition
    retomberait à zéro sur les tableaux qui ne la découpent pas.
    """
    detail = (
        pistes.get("surface_piste_legere_interne_m2"),
        pistes.get("surface_piste_legere_externe_m2"),
    )
    if any(valeur is not None for valeur in detail):
        return sum(valeur or 0.0 for valeur in detail)
    return pistes.get("surface_piste_legere_m2") or 0.0


# ---------------------------------------------------------------------------
# Modèle
# ---------------------------------------------------------------------------


@dataclass
class CoteNormalisee:
    """Une ligne de l'onglet « Dimensions postes et pieux »."""

    ouvrage: str
    dimensions: str
    #: Ordre des cotes tel qu'annoncé par l'en-tête de la section. Il change
    #: d'une section à l'autre dans le fichier de référence (« largeur x
    #: longueur x hauteur » pour le PTR, « longueur x largeur x hauteur » pour
    #: le PDL/PTR) : le lot 4 doit lire cet ordre et non le supposer.
    ordre_cotes: str
    surface_m2: float | None
    surface_plateforme_m2: float | None


@dataclass
class TableauBilan:
    """Paramètres techniques du projet, pour un indice de révision donné."""

    indice: str
    indices_disponibles: list[str]
    nom_projet: str
    generalites: dict
    structures: dict
    modules: dict
    postes: dict
    pistes: dict
    cotes: list[CoteNormalisee]
    standards: dict
    source: str
    #: Ouvrages dont la cote ne vient pas du tableau mais du classeur UNITe.
    #:
    #: Vide dans le cas courant. Renseigné, il dit exactement ce que l'outil a
    #: apporté : c'est ce qui distingue un complément d'un repli muet.
    cotes_completees: tuple[str, ...] = ()
    avertissements: list[str] = field(default_factory=list)

    def tous_parametres(self) -> dict:
        return {
            **self.generalites,
            **self.structures,
            **self.modules,
            **self.postes,
            **self.pistes,
        }


# ---------------------------------------------------------------------------
# Sélection de l'indice
# ---------------------------------------------------------------------------


def indice_depuis_nom(nom: str | Path) -> str | None:
    """Indice de révision porté par un nom de fichier DXF, ou None.

    Sert uniquement à *proposer* un appariement : il est toujours confirmé par
    l'utilisateur, un DXF pouvant très bien être exporté sous un autre nom.
    """
    trouve = MOTIF_INDICE.search(Path(str(nom)).name)
    return _format_indice(trouve) if trouve else None


def _classeur(chemin: Path):
    if not chemin.exists():
        raise ErreurTableauBilan(f"Tableau bilan introuvable : {chemin}")
    try:
        return openpyxl.load_workbook(str(chemin), data_only=True)
    except Exception as exc:  # noqa: BLE001 - remonté en erreur nommée
        raise ErreurTableauBilan(
            f"Tableau bilan illisible ({chemin.name}) : {exc}"
        ) from exc


def _onglet(classeur, nom: str, chemin: Path):
    if nom not in classeur.sheetnames:
        attendus = ", ".join(classeur.sheetnames)
        raise ErreurTableauBilan(
            f"Onglet « {nom} » absent de {chemin.name}. Onglets présents : {attendus}."
        )
    return classeur[nom]


def _onglet_caracteristiques(classeur, chemin: Path):
    """Onglet des caractéristiques, sous l'un de ses noms connus."""
    presents = {normaliser(n): n for n in classeur.sheetnames}
    for candidat in ONGLETS_CARACTERISTIQUES:
        reel = presents.get(normaliser(candidat))
        if reel is not None:
            return classeur[reel], reel
    raise ErreurTableauBilan(
        f"Aucun onglet de caractéristiques dans {chemin.name}. Noms acceptés : "
        f"{', '.join(ONGLETS_CARACTERISTIQUES)}. Onglets présents : "
        f"{', '.join(classeur.sheetnames)}."
    )


def indices_disponibles(chemin: str | Path) -> list[str]:
    """Indices de révision présents en en-tête de l'onglet caractéristiques."""
    chemin = Path(chemin)
    feuille, _ = _onglet_caracteristiques(_classeur(chemin), chemin)
    return _indices_de(feuille)[0]


def _indices_de(feuille) -> tuple[list[str], dict[str, int]]:
    """Indices de la ligne 1 et colonne où chacun se trouve."""
    colonnes: dict[str, int] = {}
    for colonne in range(2, feuille.max_column + 1):
        valeur = feuille.cell(1, colonne).value
        if valeur is None or not str(valeur).strip():
            continue
        trouve = MOTIF_INDICE.fullmatch(str(valeur).strip())
        if trouve:
            colonnes[_format_indice(trouve)] = colonne
    return list(colonnes), colonnes


# ---------------------------------------------------------------------------
# Étape B — lecture
# ---------------------------------------------------------------------------


def lire_tableau(chemin: str | Path, indice: str) -> TableauBilan:
    """Lit le tableau bilan pour un indice donné et renvoie ses paramètres."""
    chemin = Path(chemin)
    classeur = _classeur(chemin)
    feuille, nom_onglet = _onglet_caracteristiques(classeur, chemin)

    disponibles, colonnes = _indices_de(feuille)
    if not disponibles:
        raise ErreurTableauBilan(
            f"Aucun indice de révision en ligne 1 de l'onglet "
            f"« {nom_onglet} » de {chemin.name}."
        )
    if indice not in colonnes:
        raise ErreurTableauBilan(
            f"Indice « {indice} » absent du tableau bilan {chemin.name}. "
            f"Indices présents : {', '.join(disponibles)}."
        )
    colonne = colonnes[indice]

    lignes = _index_des_libelles(feuille, chemin.name)
    avertissements: list[str] = []

    valeurs: dict[str, dict] = {}
    for groupe, definitions in PARAMETRES.items():
        valeurs[groupe] = {}
        for cle, libelle, convertir, obligatoire in definitions:
            ligne = lignes.get(normaliser(libelle))
            if ligne is None:
                if obligatoire:
                    voisins = _libelles_voisins(feuille, libelle)
                    raise ErreurTableauBilan(
                        f"Paramètre « {libelle} » introuvable en colonne A de "
                        f"l'onglet « {nom_onglet} » de {chemin.name}."
                        + (
                            " Le tableau porte à la place : "
                            + " ; ".join(f"« {v} »" for v in voisins)
                            + ". C'est une autre mise en page que celle que "
                            "l'outil sait lire — signalez-la plutôt que de "
                            "renommer les lignes, qui en casserait d'autres."
                            if voisins
                            else " La mise en page du tableau a probablement "
                            "changé."
                        )
                    )
                continue
            brut = feuille.cell(ligne, colonne).value
            if brut is None and not obligatoire:
                continue
            valeurs[groupe][cle] = convertir(brut, f"{libelle} ({indice})")

    _piste_enherbee_unique(
        valeurs["pistes"], feuille, lignes, colonne, nom_onglet, avertissements
    )

    nom_projet = str(feuille.cell(1, 1).value or "").strip()
    if not nom_projet:
        avertissements.append(
            f"La cellule A1 de l'onglet « {nom_onglet} » ne porte pas "
            "de nom de projet."
        )

    cotes, avertissements_cotes = _lire_cotes(classeur, chemin)
    avertissements.extend(avertissements_cotes)
    cotes, completees = _completer_depuis_les_gabarits(cotes, avertissements)

    type_projet = valeurs["generalites"].get("type_projet", "")
    standards, avertissements_standards = _lire_standards(classeur, chemin, type_projet)
    avertissements.extend(avertissements_standards)

    return TableauBilan(
        indice=indice,
        indices_disponibles=disponibles,
        nom_projet=nom_projet,
        generalites=valeurs["generalites"],
        structures=valeurs["structures"],
        modules=valeurs["modules"],
        postes=valeurs["postes"],
        pistes=valeurs["pistes"],
        cotes=cotes,
        cotes_completees=completees,
        standards=standards,
        source=chemin.name,
        avertissements=avertissements,
    )


def _piste_enherbee_unique(
    pistes: dict, feuille, lignes: dict, colonne: int, nom_onglet: str,
    avertissements: list,
) -> None:
    """Lit la surface enherbée des tableaux qui ne la découpent pas.

    Relevé le 27/09/2026 sur le tableau d'Auzainvilliers (V3, onglet
    « Projet ») : il ne porte ni « interne » ni « externe », mais une ligne
    unique « Surface piste enherbée (m²) » — et, sur ce projet, vide. Le
    lecteur exigeait les deux lignes découpées et arrêtait le dossier entier.

    L'index des libellés fait le tri tout seul : un libellé qui apparaît deux
    fois en est écarté. Un `get` qui rend une ligne est donc, par construction,
    un libellé sans ambiguïté — celui de Saint-Cyr, présent deux fois, ne
    passera jamais par ici, et le détail découpé continue d'y faire foi.

    La lecture est dite : ce n'est pas la même chose de lire un total découpé en
    deux lignes et de lire une ligne unique, et le contrôle de surface qui en
    découle doit pouvoir se relire.
    """
    if any(
        pistes.get(cle) is not None
        for cle in ("surface_piste_legere_interne_m2", "surface_piste_legere_externe_m2")
    ):
        return

    ligne = lignes.get(normaliser(LIBELLE_PISTE_ENHERBEE))
    if ligne is None:
        avertissements.append(
            f"Aucune surface de piste enherbée lisible dans l'onglet "
            f"« {nom_onglet} » : ni « {LIBELLE_PISTE_ENHERBEE} » sans "
            "ambiguïté, ni son découpage interne/externe. La piste légère "
            "dessinée sera recoupée avec une surface déclarée nulle."
        )
        return

    brut = feuille.cell(ligne, colonne).value
    if brut is None:
        avertissements.append(
            f"« {LIBELLE_PISTE_ENHERBEE} » est vide pour cet indice : le projet "
            "ne déclare pas de piste légère. Toute piste légère dessinée au "
            "plan ressortira donc en écart."
        )
        return

    pistes["surface_piste_legere_m2"] = _nombre(brut, LIBELLE_PISTE_ENHERBEE)
    avertissements.append(
        f"Piste légère lue sur la ligne unique « {LIBELLE_PISTE_ENHERBEE} » "
        f"({pistes['surface_piste_legere_m2']:.0f} m²) : ce tableau ne la "
        "découpe pas en interne et externe."
    )


#: Mots trop communs pour rapprocher deux libellés du tableau.
_MOTS_VIDES = frozenset(
    {"de", "du", "des", "la", "le", "les", "en", "et", "a", "au", "aux", "dont",
     "m", "m2", "ml", "ha", "nombre", "total"}
)


def _mots(libelle: str) -> set:
    """Mots significatifs d'un libellé, accents et ponctuation ôtés.

    `normaliser` ne convient pas ici : elle retire aussi les espaces, et rend
    un seul bloc dont on ne peut plus compter les mots communs.
    """
    import unicodedata

    sans_accent = "".join(
        c
        for c in unicodedata.normalize("NFD", libelle or "")
        if unicodedata.category(c) != "Mn"
    ).casefold()
    bruts = re.split(r"[^a-z0-9²]+", sans_accent)
    return {
        mot.replace("²", "2")
        for mot in bruts
        if mot and mot.replace("²", "2") not in _MOTS_VIDES
    }


def _libelles_voisins(feuille, attendu: str, maximum: int = 6) -> list[str]:
    """Libellés de la colonne A qui parlent de la même chose que `attendu`.

    Un tableau dont la mise en page a changé laissait le chef de projet devant
    « paramètre introuvable » sans savoir quoi regarder. Relevé le 30/09/2026
    sur Auzainvilliers : l'outil cherchait « Surface piste lourde (m²) », et le
    classeur portait cinq lignes voisines — à renforcer, externe, existante,
    d'accès, à créer. Le chef de projet a renommé une ligne pour s'en sortir,
    ce qui en a cassé une autre.

    Deux mots significatifs en commun suffisent à rapprocher : c'est assez pour
    que « Surface piste lourde à renforcer (m²) » sorte, et pas assez pour que
    « Surface clôturée (ha) » sorte avec.
    """
    mots_attendus = _mots(attendu)
    if not mots_attendus:
        return []
    voisins: list[tuple[int, str]] = []
    for ligne in range(1, feuille.max_row + 1):
        brut = feuille.cell(ligne, 1).value
        if not isinstance(brut, str) or not brut.strip():
            continue
        communs = len(_mots(brut) & mots_attendus)
        if communs >= 2:
            voisins.append((-communs, brut.strip()))
    return [libelle for _, libelle in sorted(voisins)[:maximum]]


def _index_des_libelles(feuille, nom_fichier: str) -> dict[str, int]:
    """Libellé normalisé de la colonne A → numéro de ligne.

    Un libellé qui apparaît deux fois est écarté de l'index : le tableau porte
    plusieurs « TOTAL », « Nombre » et « Surface (m²) » dans des sections
    différentes. Aucun paramètre lu ici n'en fait partie, mais laisser passer
    une ambiguïté reviendrait à lire au hasard l'une des deux lignes.
    """
    index: dict[str, int] = {}
    ambigus: set[str] = set()
    for ligne in range(1, feuille.max_row + 1):
        libelle = normaliser(feuille.cell(ligne, 1).value)
        if not libelle:
            continue
        if libelle in index:
            ambigus.add(libelle)
            continue
        index[libelle] = ligne
    for libelle in ambigus:
        index.pop(libelle, None)

    attendus = {
        normaliser(libelle)
        for definitions in PARAMETRES.values()
        for _, libelle, _, _ in definitions
    }
    collisions = sorted(attendus & ambigus)
    if collisions:
        raise ErreurTableauBilan(
            f"{nom_fichier} : les libellés {', '.join(collisions)} apparaissent "
            "plusieurs fois en colonne A. Impossible de savoir quelle ligne lire."
        )
    return index


# ---------------------------------------------------------------------------
# Onglet « Dimensions postes et pieux »
# ---------------------------------------------------------------------------

#: Sections de l'onglet des cotes, repérées par leur titre. Chaque titre est
#: suivi d'une ligne d'en-tête « Dimensions (…) », puis des lignes de données.
SECTIONS_COTES = (
    "Poste de transformation - PTR",
    "Transformateur -  PT",
    "Poste de livraison et de transformation - PDL/PTR",
    "Poste de livraison - PDL",
    "Aire de charge (BESS)",
    "Local de stockage matériel",
    "Citerne incendie",
)

#: Cote de l'aire d'aspiration : elle n'a pas de section à elle, seulement une
#: phrase libre en fin d'onglet (« Compter une aire d'aspiration de 32m²
#: (8 x 4m) »). Elle est extraite par motif, et son absence est signalée.
MOTIF_AIRE_ASPIRATION = re.compile(
    r"aire\s+d.aspiration.*?(\d+(?:[.,]\d+)?)\s*m.*?\(\s*(\d+(?:[.,]\d+)?)\s*[x×]\s*"
    r"(\d+(?:[.,]\d+)?)\s*m\s*\)",
    re.IGNORECASE | re.DOTALL,
)


def _lire_cotes(
    classeur, chemin: Path, onglet=ONGLETS_DIMENSIONS, lire_valeur=None
) -> tuple[list[CoteNormalisee], list[str]]:
    """Table des cotes normalisées, pour la génération paramétrique du lot 4.

    Cette table ne sert pas aux contrôles de ce lot : son absence n'arrête donc
    pas l'import, mais elle est signalée — le lot 4 ne pourra pas dessiner les
    planches DP 4 sans elle.

    `onglet` et `lire_valeur` servent au classeur des gabarits UNITe, qui porte
    la même table sous un autre nom et l'écrit autrement — voir `lire_gabarits`.
    """
    avertissements: list[str] = []
    noms = (onglet,) if isinstance(onglet, str) else tuple(onglet)
    presents = {normaliser(nom): nom for nom in classeur.sheetnames}
    onglet = next(
        (presents[normaliser(nom)] for nom in noms if normaliser(nom) in presents),
        None,
    )
    if onglet is None:
        return [], [
            f"Onglet « {" » ni « ".join(noms)} » absent de {chemin.name} : les "
            "cotes normalisées des postes ne seront pas disponibles pour les "
            "planches DP 4."
        ]
    feuille = classeur[onglet]

    positions = {}
    for ligne in range(1, feuille.max_row + 1):
        for colonne in range(1, min(feuille.max_column, 6) + 1):
            libelle = normaliser(feuille.cell(ligne, colonne).value)
            for titre in SECTIONS_COTES:
                if libelle == normaliser(titre) and titre not in positions:
                    positions[titre] = (ligne, colonne)

    cotes: list[CoteNormalisee] = []
    for titre in SECTIONS_COTES:
        if titre not in positions:
            avertissements.append(
                f"Section « {titre} » introuvable dans l'onglet "
                f"« {onglet} » : cotes normalisées manquantes."
            )
            continue
        ligne, colonne = positions[titre]
        entete = _entete_dimensions(feuille, ligne, colonne)
        if entete is None:
            avertissements.append(
                f"Section « {titre} » : ligne d'en-tête « Dimensions (…) » "
                "introuvable, cotes non lues."
            )
            continue
        ligne_entete, colonne_dim, ordre, colonne_surface, colonne_plateforme = entete
        cotes.extend(
            _lignes_de_cotes(
                feuille,
                titre,
                ligne_entete,
                colonne_dim,
                ordre,
                colonne_surface,
                colonne_plateforme,
                lire_valeur or _valeur_facultative,
            )
        )

    aire = _aire_aspiration(feuille)
    if aire is None:
        avertissements.append(
            f"Cote de l'aire d'aspiration introuvable dans l'onglet "
            f"« {onglet} » : le lot 4 ne pourra pas la dessiner."
        )
    else:
        cotes.append(aire)
    return cotes, avertissements


def _completer_depuis_les_gabarits(
    cotes: list[CoteNormalisee], avertissements: list[str]
) -> tuple[list[CoteNormalisee], tuple[str, ...]]:
    """Ajoute les cotes que le tableau ne porte pas, prises au classeur UNITe.

    L'onglet des dimensions n'est pas une donnée de projet : c'est le catalogue
    des gabarits UNITe, le même dans tous les tableaux, et l'outil en livre une
    copie — celle que le parcours du plan PDF lit déjà faute de tableau bilan.
    Un tableau bâti sur un modèle ancien en porte une version tronquée, et les
    ouvrages des sections absentes ne sont alors dessinés ni sur leur planche
    DP 4 ni sur la coupe DP 3.

    Mesuré le 06/10/2026 sur Auzainvilliers, dont le tableau (V3) s'arrête après
    « Local de stockage matériel » : il manquait les sections « Aire de charge
    (BESS) » et « Citerne incendie » et la ligne de l'aire d'aspiration, soit
    six ouvrages dessinés au plan et absents du dossier.

    **Le tableau garde la main partout où il parle.** Seuls les libellés qu'il
    ne porte pas sont ajoutés : un projet dont le tableau déclare une citerne
    hors standard garde la sienne. Et la variante d'une famille reste choisie
    sur l'emprise mesurée au plan — il n'y a rien à saisir.

    Le complément est annoncé, et la liste de ce qui a été apporté redescend
    dans le contrat : une substitution muette donnerait un dossier d'apparence
    correcte dont personne ne saurait d'où viennent les hauteurs.
    """
    presentes = {cote.ouvrage for cote in cotes}
    try:
        catalogue, _ecartees = lire_gabarits(CHEMIN_GABARITS)
    except (ErreurTableauBilan, OSError) as exc:
        if presentes:
            # Rien à compléter que l'on sache : le tableau a été lu, et un
            # classeur illisible ne doit pas emporter un import qui tient.
            return cotes, ()
        avertissements.append(
            f"Classeur des gabarits UNITe illisible ({CHEMIN_GABARITS.name}) : "
            f"{exc} Les cotes absentes du tableau bilan ne peuvent pas être "
            "complétées, et les ouvrages concernés ne seront pas dessinés."
        )
        return cotes, ()

    ajoutees = [cote for cote in catalogue if cote.ouvrage not in presentes]
    if not ajoutees:
        return cotes, ()
    avertissements.append(
        f"{len(ajoutees)} cote(s) absente(s) de l'onglet des dimensions du "
        f"tableau bilan, complétée(s) depuis le classeur des gabarits UNITe "
        f"({CHEMIN_GABARITS.name}) : "
        + ", ".join(f"« {cote.ouvrage} »" for cote in ajoutees)
        + ". Ce catalogue est le même dans tous les tableaux ; celui-ci est "
        "bâti sur un modèle qui ne le porte pas en entier. Les ouvrages seront "
        "dessinés à ces dimensions — à vérifier si le projet s'écarte du "
        "standard."
    )
    return cotes + ajoutees, tuple(cote.ouvrage for cote in ajoutees)


def lire_gabarits(chemin: str | Path) -> tuple[list[CoteNormalisee], list[str]]:
    """Cotes normalisées des ouvrages UNITe, lues dans le classeur des gabarits.

    Un projet sans tableau bilan — un plan PDF sur un export HelioScope, le
    lot 2ter — n'a pas d'autre source pour les dimensions de ses ouvrages : son
    plan les situe sans être à l'échelle. Le classeur est une ressource de
    l'outil ; une section qui y manque est un défaut de l'outil, pas du projet,
    et lève. Rend aussi les cases de surface écartées, à dire au rapport.
    """
    chemin = Path(chemin)
    ecartees: list[str] = []

    def surface(feuille, ligne: int, colonne: int | None) -> float | None:
        # Le classeur écrit ses surfaces avec leur unité, « 18m² », que
        # `_nombre` sait lire, et porte ailleurs ce qui n'en est pas une :
        # « Volume 120m3 » ou « N/A » dans la colonne des plateformes de l'aire
        # de charge BESS (relevé le 23/09/2026). Ces cases sont écartées en le
        # disant ; aucune ne porte une dimension d'ouvrage.
        try:
            return _valeur_facultative(feuille, ligne, colonne)
        except ErreurTableauBilan:
            ecartees.append(
                f"« {feuille.cell(ligne, colonne).value} » (ligne {ligne}), qui "
                "n'est pas une surface"
            )
            return None

    cotes, avertissements = _lire_cotes(
        _classeur(chemin), chemin, onglet=ONGLET_GABARITS, lire_valeur=surface
    )
    if avertissements:
        raise ErreurTableauBilan(
            f"Classeur des gabarits UNITe incomplet ({chemin.name}) : "
            + " ; ".join(avertissements)
        )
    return cotes, ecartees


def _entete_dimensions(feuille, ligne_titre: int, colonne_titre: int):
    """Ligne d'en-tête « Dimensions (…) » suivant un titre de section."""
    for ligne in range(ligne_titre + 1, min(ligne_titre + 4, feuille.max_row) + 1):
        for colonne in range(1, min(feuille.max_column, 8) + 1):
            texte = str(feuille.cell(ligne, colonne).value or "")
            if not normaliser(texte).startswith("dimensions"):
                continue
            ordre = ""
            entre_parentheses = re.search(r"\((.*?)\)", texte, re.DOTALL)
            if entre_parentheses:
                ordre = " ".join(entre_parentheses.group(1).split())
            colonne_surface = colonne_plateforme = None
            for suivante in range(colonne + 1, min(feuille.max_column, colonne + 4) + 1):
                libelle = normaliser(feuille.cell(ligne, suivante).value)
                if libelle.startswith("surfaceplateforme"):
                    colonne_plateforme = suivante
                elif libelle.startswith("surface"):
                    colonne_surface = suivante
            return ligne, colonne, ordre, colonne_surface, colonne_plateforme
    return None


def _lignes_de_cotes(
    feuille,
    titre: str,
    ligne_entete: int,
    colonne_dim: int,
    ordre: str,
    colonne_surface: int | None,
    colonne_plateforme: int | None,
    lire_valeur=None,
) -> list[CoteNormalisee]:
    """Lignes de données d'une section, jusqu'à la première ligne sans dimension."""
    lire_valeur = lire_valeur or _valeur_facultative
    resultat = []
    for ligne in range(ligne_entete + 1, feuille.max_row + 1):
        dimensions = feuille.cell(ligne, colonne_dim).value
        if dimensions is None or not str(dimensions).strip():
            break
        # Le nom de la variante est à gauche de la colonne des dimensions
        # (« Conteneurs », « P<=5MWc », le volume de la citerne) ; à défaut,
        # la section n'a qu'une ligne et porte son propre nom.
        variante = feuille.cell(ligne, max(colonne_dim - 1, 1)).value
        ouvrage = f"{titre} — {variante}" if variante not in (None, "") else titre
        resultat.append(
            CoteNormalisee(
                ouvrage=str(ouvrage).strip(),
                dimensions=" ".join(str(dimensions).split()),
                ordre_cotes=ordre,
                surface_m2=lire_valeur(feuille, ligne, colonne_surface),
                surface_plateforme_m2=lire_valeur(feuille, ligne, colonne_plateforme),
            )
        )
    return resultat


def _valeur_facultative(feuille, ligne: int, colonne: int | None) -> float | None:
    if colonne is None:
        return None
    valeur = feuille.cell(ligne, colonne).value
    if valeur is None or not str(valeur).strip():
        return None
    return _nombre(valeur, f"cellule ligne {ligne} colonne {colonne}")


def _aire_aspiration(feuille) -> CoteNormalisee | None:
    for ligne in range(1, feuille.max_row + 1):
        for colonne in range(1, min(feuille.max_column, 8) + 1):
            texte = str(feuille.cell(ligne, colonne).value or "")
            trouve = MOTIF_AIRE_ASPIRATION.search(texte)
            if trouve:
                return CoteNormalisee(
                    ouvrage="Aire d'aspiration",
                    dimensions=f"{trouve.group(2)} x {trouve.group(3)} m",
                    ordre_cotes="longueur x largeur",
                    surface_m2=float(trouve.group(1).replace(",", ".")),
                    surface_plateforme_m2=None,
                )
    return None


# ---------------------------------------------------------------------------
# Onglet « Standards UNITe »
# ---------------------------------------------------------------------------

#: Ligne de l'onglet des standards portant le type d'activité, qui est la valeur
#: à apparier avec « Type de projet » de l'onglet caractéristiques.
LIBELLE_TYPE_ACTIVITE = "Type d'activité"


def _lire_standards(classeur, chemin: Path, type_projet: str) -> tuple[dict, list[str]]:
    """Valeurs de référence UNITe pour le type de projet, à titre de vraisemblance.

    Ces valeurs ne bloquent rien : elles sont affichées en regard des valeurs du
    projet à l'écran de validation. Sur le fichier de référence, quatre d'entre
    elles s'écartent du projet (format de table, inter-table, hauteurs) sans que
    ce soit une anomalie — un standard n'est pas une contrainte.
    """
    if ONGLET_STANDARDS not in classeur.sheetnames:
        return {}, [
            f"Onglet « {ONGLET_STANDARDS} » absent de {chemin.name} : pas de "
            "contrôle de vraisemblance sur les paramètres de structure."
        ]
    feuille = classeur[ONGLET_STANDARDS]

    ligne_types = None
    for ligne in range(1, feuille.max_row + 1):
        for colonne in range(1, min(feuille.max_column, 4) + 1):
            if normaliser(feuille.cell(ligne, colonne).value) == normaliser(
                LIBELLE_TYPE_ACTIVITE
            ):
                ligne_types = (ligne, colonne)
                break
        if ligne_types:
            break
    if ligne_types is None:
        return {}, [
            f"Ligne « {LIBELLE_TYPE_ACTIVITE} » introuvable dans l'onglet "
            f"« {ONGLET_STANDARDS} » : pas de contrôle de vraisemblance."
        ]

    ligne, colonne_libelles = ligne_types
    cible = normaliser(type_projet)
    colonne_projet = None
    for colonne in range(colonne_libelles + 1, feuille.max_column + 1):
        if normaliser(feuille.cell(ligne, colonne).value) == cible:
            colonne_projet = colonne
            break
    if colonne_projet is None:
        return {}, [
            f"Type de projet « {type_projet} » absent de l'onglet "
            f"« {ONGLET_STANDARDS} » : pas de contrôle de vraisemblance."
        ]

    standards = {}
    for ligne_valeur in range(ligne + 1, feuille.max_row + 1):
        libelle = feuille.cell(ligne_valeur, colonne_libelles).value
        if libelle is None or not str(libelle).strip():
            continue
        valeur = feuille.cell(ligne_valeur, colonne_projet).value
        if valeur is None or not str(valeur).strip():
            continue
        standards[" ".join(str(libelle).split())] = valeur
    return standards, []
