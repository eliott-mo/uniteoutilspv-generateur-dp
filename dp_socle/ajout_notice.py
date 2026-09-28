"""Verser la notice DP 11 dans un dossier PowerPoint déjà fini (lot 7).

POURQUOI CE CHEMIN EXISTE
-------------------------
Le chef de projet travaille en trois temps, et ce n'est pas un contournement :
il génère sans notice pour valider son plan au plus tôt et transmettre le
GeoPackage au photomontage ; il finit son `.pptx` à la main — photographies
dans les cadres, photomontage reçu, diapos surnuméraires supprimées — pendant
que la notice se rédige ; et la notice n'arrive qu'à la fin.

Régénérer le dossier à ce moment-là **détruirait son travail**. Le brief du
24/09/2026 posait pourtant l'inverse (« régénération complète, pas
sélective »), et il avait raison tant que la sortie était un PDF que l'outil
assemblait. Le lot 8 a changé ce qu'est la sortie ; la décision s'inverse avec
elle.

CE QUI REND LA CHOSE POSSIBLE, ET QUI A ÉTÉ MESURÉ
--------------------------------------------------
- **Le fichier dit sa propre pagination.** Chaque diapo porte une mise en page
  nommée d'après sa pièce (`ooxml.nouvelle_mise_en_page`). Relevé le
  28/09/2026 sur un dossier revenu de PowerPoint : les noms survivent, et les
  diapos supprimées par le chef de projet ont bien disparu de la liste. C'est
  donc le fichier déposé qui fait foi, pas une mémoire du serveur — qui n'en a
  aucune.
- **PowerPoint conserve les propriétés du document.** Même relevé : un fichier
  enregistré par PowerPoint a gardé son `comments`. L'identité du projet y
  voyage donc, en cent caractères.
- **La photo de couverture survit au redessin de la page de garde.** Le fond
  est posé sur la mise en page, la réservation d'image est une forme sur la
  diapo : échanger l'un ne touche pas l'autre.

CE QUE CE MODULE NE FAIT PAS
----------------------------
Il ne redessine aucune planche, et ne change pas la date du dossier. Elle est
dans le cartouche de chaque planche, dessinée dans l'image : la changer sur la
seule page de garde donnerait un dossier dont la couverture contredit ses
propres planches.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import date as _date
from pathlib import Path

from pptx import Presentation

from . import ooxml, rendu_pptx
from .dossier import PAR_CODE, PIECES
from .erreurs import ErreurReprise
from .planches import page_garde
from .projet import Projet

#: Le code de la notice, seule pièce que ce module sache ajouter.
CODE_NOTICE = "DP 11"

#: La page de garde n'a pas de code : sa mise en page porte son titre. Lu dans
#: `dossier.PIECES` plutôt que recopié — les deux ne peuvent pas diverger.
LIBELLE_PAGE_GARDE = next(piece.titre for piece in PIECES if not piece.code)

#: Ce que le `.pptx` porte de son projet, et qui suffit à redessiner la page de
#: garde. Le `projet.json` entier ferait 2 000 octets ; `core_properties`
#: plafonne à 255 caractères par champ, et ces quatre-là en font cent.
CHAMPS_IDENTITE = ("libelle", "commune", "code_postal", "date")

#: Marque en tête de la propriété, pour distinguer notre charge utile du
#: « generated using python-pptx » que la bibliothèque y écrit par défaut.
MARQUE_IDENTITE = "dossier-dp:"


def ecrire_identite(presentation, projet: Projet) -> None:
    """Grave dans le fichier de quoi le compléter plus tard.

    Appelé à la génération. Sans cette marque, un dossier redéposé ne peut pas
    voir sa page de garde redessinée : ni la commune, ni la date ne se relisent
    d'une image.
    """
    charge = json.dumps(
        {
            "libelle": projet.libelle_affiche,
            "commune": projet.commune,
            "code_postal": projet.code_postal,
            "date": projet.date,
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    proprietes = presentation.core_properties
    proprietes.comments = MARQUE_IDENTITE + charge
    proprietes.title = projet.libelle_affiche


def lire_identite(presentation) -> dict | None:
    """L'identité gravée à la génération, ou None si le fichier n'en porte pas."""
    brut = (presentation.core_properties.comments or "").strip()
    if not brut.startswith(MARQUE_IDENTITE):
        return None
    try:
        identite = json.loads(brut[len(MARQUE_IDENTITE):])
    except ValueError:
        return None
    if not all(champ in identite for champ in CHAMPS_IDENTITE):
        return None
    return identite


#: Ce qui suit le code d'une pièce dans le nom d'une mise en page : le nombre de
#: cadres d'une alternative, son échelle, le rang d'une page de notice.
_SUFFIXE_MISE_EN_PAGE = re.compile(r"\s+(—|au)\s+.*$|,\s*page\s+\d+$")


def code_de_la_mise_en_page(nom: str) -> str:
    """Code de la pièce que porte une mise en page.

    « DP 6 — 2 cadres » et « DP 8 au 1/5000 » sont des alternatives d'une même
    pièce ; « DP 11, page 3 » une page parmi d'autres. Toutes rendent le code
    nu, seul repère que le sommaire connaisse.
    """
    nom = (nom or "").strip()
    if nom == LIBELLE_PAGE_GARDE:
        return ""
    return _SUFFIXE_MISE_EN_PAGE.sub("", nom).strip()


@dataclass
class DossierFini:
    """Ce qu'un `.pptx` déposé dit de lui-même."""

    identite: dict
    #: Code de chaque diapo, dans l'ordre du fichier. La page de garde y est
    #: le code vide, comme dans `dossier.codes_produits`.
    codes: list = field(default_factory=list)

    @property
    def nb_diapos(self) -> int:
        return len(self.codes)

    @property
    def porte_la_notice(self) -> bool:
        return CODE_NOTICE in self.codes

    @property
    def pages_de_la_notice(self) -> int:
        return self.codes.count(CODE_NOTICE)

    @property
    def premiere_page_de_la_notice(self) -> int:
        """Numéro de page qu'occuperait la notice ajoutée en fin de dossier."""
        return self.nb_diapos + 1

    def pagination(self) -> list:
        """(code, titre, page) de chaque pièce, dans l'ordre, sans doublon.

        Les alternatives ayant déjà été tranchées par le chef de projet, une
        pièce n'apparaît qu'une fois — sauf la notice, dont les pages se
        suivent et dont seule la première compte au sommaire.
        """
        lignes, vus = [], set()
        for rang, code in enumerate(self.codes, start=1):
            if code in vus:
                continue
            vus.add(code)
            piece = PAR_CODE.get(code)
            lignes.append((code, piece.titre if piece else LIBELLE_PAGE_GARDE, rang))
        return lignes


def lire_le_dossier(source) -> DossierFini:
    """Ouvre un `.pptx` produit par l'outil et en tire ce qu'il faut pour le compléter.

    `source` est un chemin ou des octets. Les refus nomment ce qu'il y a à
    faire : ce sont des situations ordinaires, pas des défauts.
    """
    presentation = _ouvrir(source)
    identite = lire_identite(presentation)
    if identite is None:
        raise ErreurReprise(
            "Ce fichier ne porte pas l'identité de son projet : il n'a pas été "
            "produit par cet outil, ou il l'a été avant que la reprise n'existe "
            "(28/09/2026). Un dossier regénéré depuis la portera, et pourra "
            "recevoir sa notice sans être refait."
        )
    codes = [
        code_de_la_mise_en_page(diapo.slide_layout.name)
        for diapo in presentation.slides
    ]
    inconnus = sorted({c for c in codes if c and c not in PAR_CODE})
    if inconnus:
        raise ErreurReprise(
            f"Ce dossier porte des planches que l'outil ne reconnaît pas "
            f"({', '.join(inconnus)}) : ses mises en page ont été renommées, et "
            "la pagination ne peut plus s'en déduire."
        )
    if not codes:
        raise ErreurReprise("Ce dossier ne porte aucune diapo.")
    return DossierFini(identite=identite, codes=codes)


def ajouter_la_notice(source, chemin_notice, travail: Path) -> bytes:
    """Rend le `.pptx` complété de sa notice, page de garde remise à jour.

    Trois gestes, et rien d'autre : les pages de la notice à la fin, le
    sommaire de la page de garde redessiné pour qu'il porte enfin son numéro,
    le reste intact. `travail` reçoit le PDF de la page de garde recomposée.
    """
    dossier = lire_le_dossier(source)
    if dossier.porte_la_notice:
        raise ErreurReprise(
            f"Ce dossier porte déjà une notice de {dossier.pages_de_la_notice} "
            "page(s). Remplacer une notice est un autre geste que l'ajouter : "
            "regénérez le dossier avec la bonne notice."
        )

    pages = rendu_pptx.rendre_pages_pdf(Path(chemin_notice), CODE_NOTICE)
    if not pages:
        raise ErreurReprise(
            "La notice déposée ne porte aucune page lisible : rien à ajouter."
        )

    presentation = _ouvrir(source)
    _refaire_la_page_de_garde(presentation, dossier, travail)
    for rang, image in enumerate(pages, start=1):
        libelle = (
            CODE_NOTICE if len(pages) == 1 else f"{CODE_NOTICE}, page {rang}"
        )
        mise_en_page = ooxml.nouvelle_mise_en_page(presentation, libelle)
        ooxml.poser_fond(mise_en_page, image)
        presentation.slides.add_slide(mise_en_page)

    tampon = io.BytesIO()
    presentation.save(tampon)
    return tampon.getvalue()


def _refaire_la_page_de_garde(presentation, dossier: DossierFini, travail: Path) -> None:
    """Redessine la couverture pour que son sommaire porte la notice.

    Seule la page de garde est refaite, et seulement son **fond** : la
    réservation d'image, où le chef de projet a pu déposer sa perspective, est
    une forme de la diapo et n'est pas touchée.
    """
    projet = _projet_de_l_identite(dossier.identite)
    pages = {code: (page, page) for code, _titre, page in dossier.pagination()}
    pages[CODE_NOTICE] = (
        dossier.premiere_page_de_la_notice,
        dossier.premiere_page_de_la_notice,
    )
    garde = page_garde.generer(projet, travail, pages=pages)
    fond = rendu_pptx.rendre_fond(garde.planche, "Page de garde")

    diapo = presentation.slides[0]
    if code_de_la_mise_en_page(diapo.slide_layout.name) != "":
        raise ErreurReprise(
            "La première diapo de ce dossier n'est pas sa page de garde "
            f"(« {diapo.slide_layout.name} ») : le sommaire ne peut pas être "
            "remis à jour sans risquer d'écraser une planche."
        )
    ooxml.remplacer_fond(diapo.slide_layout, fond.image, fond.svg)


def _projet_de_l_identite(identite: dict) -> Projet:
    """Le minimum que `page_garde.generer` demande, reconstruit de l'identité.

    Ni emprise ni notice : la page de garde ne les lit pas. Son image de
    perspective non plus — c'est une réservation dans cette voie, posée sur la
    diapo et non dessinée dans la planche.
    """
    _date.fromisoformat(identite["date"])  # lève si la marque a été trafiquée
    return Projet(
        nom=identite["libelle"],
        commune=identite["commune"],
        code_postal=identite["code_postal"],
        date=identite["date"],
        emprise=None,
        libelle=identite["libelle"],
    )


def _ouvrir(source):
    """Une présentation, depuis un chemin ou des octets."""
    if isinstance(source, (bytes, bytearray)):
        return Presentation(io.BytesIO(source))
    if hasattr(source, "read"):
        return Presentation(source)
    return Presentation(str(source))
