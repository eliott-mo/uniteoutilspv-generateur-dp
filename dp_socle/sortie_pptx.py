"""La sortie PowerPoint que le chef de projet finit lui-même (lot 8).

Le dossier sort aussi en un `.pptx` dont les planches sont déjà dessinées et dont
les pièces photographiques restent à garnir : cadres vides, plan de repérage déjà
imprimé, repères de vue en formes libres à poser à la main. Le chef de projet
peut ainsi télécharger le dossier **sans avoir traité ses photographies**, qui
sont ce qui bloque un dossier : il faut avoir choisi les prises de vue, les avoir
géolocalisées, avoir posé leurs cônes de visée avant que l'outil produise quoi
que ce soit.

CE QUI CESSE D'ÊTRE GARANTI, ET QUI DOIT SE LIRE
-------------------------------------------------
C'est moins robuste que la voie PDF, et c'est assumé. Ce que l'outil garantit
dans le PDF — un cône à la bonne position, au bon azimut, vérifié contre l'EXIF —
devient la responsabilité du chef de projet : « la localisation des cônes de vue
ne sera pas parfaite mais largement suffisante » (25/09/2026). Deux contrôles de
la voie PDF ne s'appliquent plus, parce qu'ils portent sur un fichier que le chef
de projet édite après nous :

- `assemblage._verifier_numerotation`, qui vérifie que le cartouche annonce la
  page où la planche tombe. Les diapos alternatives partagent un numéro, et
  supprimer les surnuméraires est justement ce qui rend la pagination juste ;
- `assemblage.verifier_format`, qui vérifie que chaque page mesure 420 x 297 mm.

Ils sont remplacés par ce que le rapport **dit**, et par le montage : le fond est
posé sur la mise en page, donc hors d'atteinte, et les cadres sont verrouillés en
position. Voir `AVERTISSEMENTS_DE_PRINCIPE`.

CE QUI RESTE DE LA DÉCISION D1, ET CE QUI N'EN EST PAS FAIT
------------------------------------------------------------
D1 voulait éditables les légendes, la table des matières de la page de garde et
les intitulés des cadres photo. Seuls **les intitulés des cadres** le sont ici :
le lot 8 dessine lui-même les cadres photo, il peut donc les dessiner sans leur
titre et le reconstruire en zone de texte. Les légendes des planches
cartographiques et le sommaire de la page de garde sont dans l'image, donc pas
éditables — les rendre éditables demanderait de redessiner ces planches sans eux,
c'est-à-dire de toucher aux modules des lots 1 à 4. Le rapport le dit à chaque
génération plutôt que de le laisser découvrir.
"""

from __future__ import annotations

import tempfile
import warnings
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Mm, Pt

from . import ooxml, rendu_pptx
from .contrat import charger_contrat
from .dossier import codes_produits, piece
from .erreurs import ErreurContrat, ErreurDP, ErreurSortiePPTX
from .geometrie import charger_emprise
from .ign import DPI_DEFAUT
from .planche import HAUTEUR_MM, LARGEUR_MM, Planche
from .planches import (
    dp11_notice,
    dp1_1_situation,
    dp1_2_aerienne,
    dp1_3_cadastre,
    dp2_plan_masse,
    dp3_coupes,
    dp4_ouvrages,
    page_garde,
)
from .planches.commun import nouvelle_planche
from .planches.photographies import (
    composer_le_panneau,
    disposition,
    interieur_du_panneau,
)
from .planches.primitives import union_valide
from .planches.reperage_vues import (
    OUVERTURE_CONE_DEG,
    RAYON_CONE_MM,
    RAYON_REPERE_MM,
    cadrages_sans_vues,
    repere_de_vue,
)
from .polices import POLICE_PRINCIPALE, avertir_si_indisponible
from .projet import Projet

#: Titre de la page de garde au rapport, comme dans le sommaire du dossier.
TITRE_PAGE_GARDE = "Page de garde"

#: Motif du nom du fichier. Il se nomme pour ce qu'il est : un dossier à finir,
#: pas un dossier déposable (décision D9).
MOTIF_SORTIE = "{nom}_DP_a_finaliser.pptx"

#: Les pièces photographiques, ce qu'elles portent, et combien de cadrages au
#: choix elles reçoivent.
#:
#: DP 6 reçoit deux alternatives qui ne diffèrent que par le nombre de cadres —
#: le troisième volet, « projet avec mesures paysagères », n'existe que si le
#: projet porte des mesures paysagères, et le chef de projet garde la diapo qui
#: correspond. DP 8 reçoit trois cadrages successifs, parce que son point de vue
#: peut être loin et qu'on ne sait pas encore où (décision D2).
PIECES_PHOTO = (
    ("DP 6", (2, 3), 1, 1),
    ("DP 7", (2,), 2, 1),
    ("DP 8", (2,), 2, 3),
)

PHOTO_PAR_CODE = {ligne[0]: ligne[1:] for ligne in PIECES_PHOTO}

#: Ce que le chef de projet lit sur le bandeau d'une diapo alternative.
MENTION_ALTERNATIVE = (
    "{nombre} versions au choix — n'en gardez qu'une, supprimez les autres et "
    "ce bandeau avant d'exporter en PDF."
)

#: Rouge du bandeau. Franc, et absent de la palette des ouvrages : un bandeau ne
#: doit pas pouvoir se lire comme une catégorie du plan.
ROUGE_BANDEAU = "C00000"

#: Teinte des visées, reprise de `reperage_vues._VISEE` — le même symbole que le
#: chef de projet a sous les yeux dans `photos-geoloc`, et que porte la voie PDF.
MAGENTA_VISEE = "D81B8C"

#: Opacité du cône, comme au PDF (`Style.opacite_remplissage` y vaut 0,55).
OPACITE_CONE = 0.55

#: Ce qu'un degré vaut dans les poignées d'angle de `python-pptx`.
#:
#: Les poignées d'un `prstGeom` s'écrivent en 60 000e de degré dans le fichier, et
#: `python-pptx` expose la valeur écrite divisée par 100 000. Un degré vaut donc
#: 0,6 dans son API. Mesuré le 26/09/2026 dans le fichier produit : sans ce
#: facteur, un secteur de 50° sortait à `val="6944"`, soit 0,12° — une aiguille
#: invisible, et un cône de visée qui ne montre plus rien.
DEGRE_EN_POIGNEE = 60000 / 100000

#: Corps des textes posés sur la diapo, en points. 7,5 pt est le corps du titre
#: d'un sous-cadre dans `primitives.sous_cadre` : l'intitulé éditable doit
#: ressembler à celui que la voie PDF imprime.
CORPS_TITRE_CADRE_PT = 7.5
CORPS_ETIQUETTE_PT = 6.5
CORPS_BANDEAU_PT = 12.0

#: Premier `idx` des réservations d'image. Les `idx` 10 à 12 sont libérés par
#: `ooxml.nouvelle_mise_en_page`, mais partir au-dessus évite d'avoir à le savoir.
PREMIER_IDX_CADRE = 20

#: Ce que le rapport dit à chaque génération. Aucune de ces phrases n'est
#: conditionnelle : ce sont les garanties que cette sortie ne donne pas, et les
#: taire serait le repli silencieux le plus coûteux du lot.
AVERTISSEMENTS_DE_PRINCIPE = (
    "Cette sortie est un dossier à finir, pas un dossier déposable. Les deux "
    "contrôles de la voie PDF ne s'y appliquent pas : ni la vérification que "
    "chaque cartouche annonce la page où il tombe, ni celle du format 420 x "
    "297 mm de chaque page. C'est vous qui les tenez à partir d'ici.",
    "Les repères et cônes de vue sont posés au centre de chaque plan de "
    "repérage, à déplacer et à orienter à la main. Leur position n'est pas "
    "vérifiée contre l'EXIF des photographies, contrairement au dossier PDF.",
    "Laissez cochée l'option « Ne pas compresser les images dans un fichier » "
    "(Fichier > Options > Options avancées) : décochée, la résolution par "
    "défaut de 96 ppp détruirait les plans à l'enregistrement.",
    "Les légendes des planches et le sommaire de la page de garde sont dans "
    "l'image et ne sont pas éditables ; seuls les intitulés des cadres photo le "
    "sont. Une correction de légende se fait en régénérant le dossier.",
)


@dataclass
class Diapo:
    """Une diapo du fichier produit, et ce qu'il faut en dire au rapport."""

    code: str
    #: Numéro annoncé au cartouche. Les alternatives d'une même pièce le
    #: partagent (décision D3).
    numero: int
    titre: str
    libelle: str
    voie: str
    #: Rang de l'alternative dans son groupe, et taille du groupe. `(1, 1)` pour
    #: une diapo unique.
    alternative: tuple = (1, 1)
    cadres: int = 0
    echelle: int | None = None
    octets: int = 0


@dataclass
class RapportPPTX:
    """Compte rendu d'une génération PowerPoint, destiné à l'interface."""

    dossier: Path
    fichier: Path
    diapos: list = field(default_factory=list)
    sommaire: list = field(default_factory=list)
    taille_mo: float = 0.0
    etat_polices: object | None = None
    avertissements: list = field(default_factory=list)
    origine_contrat: str | None = None
    notice: dict | None = None


def nom_sortie(projet: Projet) -> str:
    return MOTIF_SORTIE.format(nom=projet.nom)


def generer_pptx(
    projet: Projet,
    dossier_sortie: str | Path = "sortie",
    dpi: int = DPI_DEFAUT,
    fond_ign: bool = True,
) -> RapportPPTX:
    """Produit le `.pptx` à finaliser, et le rapport qui dit ce qu'il contient.

    `fond_ign` à faux retire le fond cartographique des plans de repérage, comme
    au lot 6 : c'est la prise qui permet de mesurer le montage sans interroger la
    Géoplateforme.

    L'ordre suit celui du dossier PDF, parce que c'est l'ordre du sommaire et
    celui que les services instructeurs lisent. Les pièces photographiques y
    figurent **toujours**, cadres vides : c'est la raison d'être de cette sortie,
    et un dossier dont les photographies sont prêtes se télécharge en PDF.
    """
    projet.valider()
    etat = avertir_si_indisponible()
    if projet.chemin_notice is not None:
        dp11_notice.examiner(projet.chemin_notice)

    dossier = Path(dossier_sortie) / projet.nom
    dossier.mkdir(parents=True, exist_ok=True)
    emprise = charger_emprise(projet.chemin_emprise)

    avertissements = list(AVERTISSEMENTS_DE_PRINCIPE)
    if not etat.disponible:
        avertissements.append(etat.message)
    contrat, message = _contrat_eventuel(dossier, projet)
    if message:
        avertissements.append(message)

    # Les PDF intermédiaires vont dans un dossier de travail, pas dans celui du
    # projet. Les modules de pièce écrivent toujours un PDF, et ceux d'ici portent
    # la pagination du PPTX — DP 11 y tombe trois pages plus loin que dans le
    # dossier PDF, puisque les pièces photographiques y figurent. Les laisser dans
    # `sortie/{projet}/` remplacerait les PDF de la voie PDF par des planches aux
    # cartouches d'une autre pagination : deux dossiers vrais, mélangés, donc un
    # dossier faux. Seul le `.pptx` reste.
    with (
        tempfile.TemporaryDirectory(prefix="pptx-") as travail,
        warnings.catch_warnings(record=True) as captees,
    ):
        warnings.simplefilter("always", RuntimeWarning)
        groupes = _composer_les_planches(
            projet, contrat, emprise, Path(travail), avertissements, dpi, fond_ign
        )
        notice = _notice_eventuelle(projet, groupes, Path(travail), avertissements)
        fichier = _ecrire(projet, groupes, notice, dossier, Path(travail))
    avertissements.extend(
        str(c.message) for c in captees if issubclass(c.category, RuntimeWarning)
    )

    taille_mo = fichier.stat().st_size / (1024 * 1024)
    return RapportPPTX(
        dossier=dossier,
        fichier=fichier,
        diapos=[diapo for groupe in groupes for diapo in groupe.diapos],
        sommaire=_sommaire(groupes),
        taille_mo=taille_mo,
        etat_polices=etat,
        avertissements=list(dict.fromkeys(avertissements)),
        origine_contrat=contrat.origine if contrat is not None else None,
        notice=notice.details.get("notice") if notice is not None else None,
    )


# ---------------------------------------------------------------------------
# Les planches, et la pagination qu'elles portent
# ---------------------------------------------------------------------------


@dataclass
class Groupe:
    """Une pièce du dossier, et la ou les diapos qu'elle occupe.

    Une pièce tient sur **une page** du dossier, même quand elle sort en
    plusieurs diapos : les alternatives partagent le numéro de leur groupe, et le
    sommaire n'en compte qu'une (décision D3).
    """

    code: str
    numero: int
    titre: str
    planches: list
    diapos: list = field(default_factory=list)
    #: Nombre de pages que la pièce occupe au dossier. Vaut 1 partout sauf pour
    #: la notice, seule pièce qui s'étend.
    pages: int = 1


def _composer_les_planches(projet, contrat, emprise, dossier, avertissements, dpi,
                           fond_ign=True):
    """Toutes les planches du dossier, dans l'ordre, avec leur numéro de page.

    Le numéro de cartouche se compte sur les pièces que **cette** sortie porte.
    Il diffère donc de celui du PDF d'un même projet dont les photographies ne
    sont pas prêtes : le PDF y saute DP 6, DP 7 et DP 8, et sa notice tombe trois
    pages plus tôt. C'est voulu — chacune des deux sorties est paginée pour
    elle-même.
    """
    codes = ["DP 1-1", "DP 1-2", "DP 1-3"]
    emprise_cloturee = None
    if contrat is not None:
        codes.append("DP 2")
        if contrat.profil and (contrat.profil.get("points") or []):
            codes.append("DP 3")
        else:
            avertissements.append(
                "Aucun profil de terrain au contrat : DP 3 n'est pas produite. "
                "Tracez la ligne de coupe A-A' et relevez le profil à l'import."
            )
        avertissements.extend(dp4_ouvrages.ouvrages_ecartes(contrat))
        codes.extend(dp4_ouvrages.planches_necessaires(contrat))
        emprise_cloturee = union_valide(contrat.geometries("cloture"))
        if emprise_cloturee is None or emprise_cloturee.is_empty:
            avertissements.append(
                "Aucune clôture au contrat : les plans de repérage des pièces "
                "photographiques n'ont pas d'emprise à montrer, et DP 6, DP 7 et "
                "DP 8 ne sont pas produites."
            )
        else:
            codes.extend(code for code, _, _, _ in PIECES_PHOTO)
    if projet.chemin_notice is not None:
        codes.append("DP 11")

    produits = codes_produits(codes)
    groupes = []
    for code in codes:
        numero = produits.index(code) + 1
        if code == "DP 11":
            # La notice se compose à part : son PDF est fusionné, pas dessiné.
            groupes.append(Groupe(code, numero, piece(code).titre, planches=[]))
            continue
        try:
            groupes.append(
                _groupe_de_la_piece(
                    projet, code, numero, contrat, emprise, emprise_cloturee,
                    dossier, avertissements, dpi, fond_ign,
                )
            )
        except ErreurDP as exc:
            avertissements.append(f"{code} n'est pas produite : {exc}")
    return groupes


def _groupe_de_la_piece(projet, code, numero, contrat, emprise, emprise_cloturee,
                        dossier, avertissements, dpi, fond_ign=True):
    """Les planches d'une pièce : une seule, ou plusieurs alternatives."""
    if code in PHOTO_PAR_CODE:
        emplacements, vues, alternatives = PHOTO_PAR_CODE[code]
        planches, messages = _planches_photo(
            projet, code, numero, emplacements, vues, alternatives,
            emprise_cloturee, contrat, fond_ign,
        )
        avertissements.extend(messages)
        return Groupe(code, numero, piece(code).titre, planches=planches)

    sortie = _planche_cartographique(
        projet, code, numero, contrat, emprise, dossier, dpi
    )
    avertissements.extend(sortie.details.get("avertissements", []) if sortie.details else [])
    return Groupe(code, numero, piece(code).titre, planches=[PlanchePPTX(sortie.planche)])


def _planche_cartographique(projet, code, numero, contrat, emprise, dossier, dpi):
    """Une planche que les lots 1 à 4 dessinent déjà, produite telle quelle.

    Le module de la pièce écrit toujours un PDF ; celui-ci part dans le dossier de
    travail, et sera effacé. Ce qui nous intéresse est la planche qu'il a
    composée — `Sortie.planche` —, parce que c'est d'elle que la diapo tire son
    image : c'est ce qui rend les deux sorties identiques par construction, et non
    seulement ressemblantes.

    DP 1-1, DP 1-2 et DP 1-3 n'ont pas de `numero` à recevoir : elles ouvrent
    toujours le dossier, aux pages 2, 3 et 4, et leur module le sait.
    """
    if code == "DP 1-1":
        return dp1_1_situation.generer(projet, emprise, dossier, dpi=dpi)
    if code == "DP 1-2":
        return dp1_2_aerienne.generer(projet, emprise, dossier, dpi=dpi)
    if code == "DP 1-3":
        return dp1_3_cadastre.generer(projet, emprise, dossier)
    if code == "DP 2":
        return dp2_plan_masse.generer(
            projet, contrat, emprise, dossier, numero=str(numero)
        )
    if code == "DP 3":
        return dp3_coupes.generer(projet, contrat, dossier, numero=str(numero))
    if code.startswith("DP 4"):
        return dp4_ouvrages.generer(
            projet, contrat, dossier, code, numero=str(numero)
        )
    raise ErreurSortiePPTX(
        f"Aucune façon de composer {code} pour la sortie PowerPoint."
    )


@dataclass
class PlanchePPTX:
    """Une planche prête à devenir une diapo, et ce qu'elle porte de particulier."""

    planche: Planche
    #: Cadres photo à réserver, `(x, y, largeur, hauteur)` en millimètres.
    cadres: list = field(default_factory=list)
    #: Intitulés éditables, un par cadre.
    intitules: list = field(default_factory=list)
    #: Repères de vue à poser en formes déplaçables.
    reperes: list = field(default_factory=list)
    #: Centre du panneau de repérage, où les repères sont posés.
    centre_panneau: tuple | None = None
    #: Cadre du panneau de repérage, `(x, y, largeur, hauteur)`. Le bandeau des
    #: diapos alternatives s'y pose.
    panneau: tuple | None = None
    echelle: int | None = None
    #: Libellé de la diapo, pour le rapport et le nom de la mise en page.
    libelle: str = ""


def _planches_photo(projet, code, numero, emplacements, vues, alternatives,
                    emprise_cloturee, contrat, fond_ign=True):
    """Les diapos d'une pièce photographique : cadres vides et repères à poser.

    Le cadrage ne peut pas se calculer sur les points de vue — il n'y en a pas
    encore. Voir `reperage_vues.cadrages_sans_vues` : emprise plus 150 m, à la
    plus grande échelle qui la contient, déclinée aux crans suivants pour DP 8.
    """
    reperes = [repere_de_vue(code, rang) for rang in range(1, vues + 1)]
    modele = nouvelle_planche(projet, code, numero=str(numero))
    cadrages, messages = cadrages_sans_vues(
        emprise_cloturee,
        interieur_du_panneau(disposition(modele, emplacements[0]).panneau),
        f"{code} — plan de repérage",
        alternatives=alternatives,
    )

    planches = []
    for cadrage in cadrages:
        for nombre_cadres in emplacements:
            planche = nouvelle_planche(projet, code, numero=str(numero))
            plan = disposition(planche, nombre_cadres)
            interieur, messages_panneau = composer_le_panneau(
                planche, plan.panneau, cadrage, emprise_cloturee, contrat,
                fond_ign=fond_ign,
            )
            messages.extend(messages_panneau)
            for pose in plan.emplacements:
                # Le cadre sans son titre : l'intitulé devient une zone de texte
                # éditable, seul reste vivant de la décision D1.
                sous_cadre_nu(planche, pose.cadre)
            planches.append(
                PlanchePPTX(
                    planche=planche,
                    cadres=[pose.image for pose in plan.emplacements],
                    intitules=_intitules(code, reperes, nombre_cadres),
                    reperes=reperes,
                    centre_panneau=(
                        interieur[0] + interieur[2] / 2.0,
                        interieur[1] + interieur[3] / 2.0,
                    ),
                    panneau=plan.panneau,
                    echelle=cadrage.denominateur,
                    libelle=_libelle_alternative(
                        code, cadrage.denominateur, nombre_cadres,
                        len(cadrages) * len(emplacements),
                    ),
                )
            )
    return planches, messages


def sous_cadre_nu(planche: Planche, cadre: tuple) -> None:
    """Le filet d'un sous-cadre, sans son titre.

    `primitives.sous_cadre` appelé sans titre laisse la place du titre vide, ce
    qui est exactement ce qu'il faut : la zone de texte éditable de la diapo s'y
    posera, et la géométrie du cadre ne change pas d'un millimètre.
    """
    from .planches.primitives import sous_cadre

    sous_cadre(planche, *cadre, titre=None)


def _intitules(code: str, reperes: list, nombre_cadres: int) -> list:
    """Les intitulés des cadres, dans la langue du dossier de référence."""
    from .planches.dp6_insertions import INTITULES as VOLETS

    if code == "DP 6":
        return [f"{reperes[0]} : {volet}" for volet in VOLETS[:nombre_cadres]]
    intitule = (
        "Photographie de l'environnement proche" if code == "DP 7"
        else "Photographie du paysage lointain"
    )
    return [f"{repere} : {intitule}" for repere in reperes[:nombre_cadres]]


def _libelle_alternative(code, denominateur, nombre_cadres, total) -> str:
    """Le nom de la diapo, qui doit distinguer les alternatives entre elles."""
    if total == 1:
        return code
    if code == "DP 6":
        return f"{code} — {nombre_cadres} cadres"
    return f"{code} au 1/{denominateur}"


def _contrat_eventuel(dossier: Path, projet: Projet):
    """Même règle que la voie PDF : un contrat absent arrête le dossier au socle."""
    try:
        return charger_contrat(dossier, voirie=projet.voirie), None
    except ErreurContrat as exc:
        if "Aucun contrat" in str(exc):
            return None, (
                "Aucun plan importé pour ce projet : la sortie s'arrête au plan "
                "de cadastre. Importez le plan du bureau d'études ou l'export "
                "HelioScope pour produire DP 2, DP 3, DP 4 et les pièces "
                "photographiques."
            )
        raise


def _notice_eventuelle(projet, groupes, dossier, avertissements):
    """La pièce DP 11, si une notice a été déposée.

    Son absence n'arrête pas la génération : c'est la règle du 14/09/2026, et
    elle vaut d'autant plus ici, où la sortie sert précisément à travailler sur
    un dossier incomplet.
    """
    if projet.chemin_notice is None:
        avertissements.append(
            "Aucune notice DP 11 : la sortie est produite sans elle. Déposez le "
            "PDF de la notice pour qu'elle soit habillée du cadre et du "
            "cartouche et paginée avec les autres pièces."
        )
        return None
    groupe = next((g for g in groupes if g.code == "DP 11"), None)
    if groupe is None:  # pragma: no cover - garde-fou
        raise ErreurSortiePPTX(
            "Une notice est déposée mais DP 11 n'a pas été paginée."
        )
    sortie = dp11_notice.generer(
        projet, projet.chemin_notice, dossier, premier_numero=groupe.numero
    )
    avertissements.extend(sortie.details.get("avertissements", []))
    groupe.pages = len(sortie.numeros) or 1
    return sortie


def _sommaire(groupes) -> list:
    """Le sommaire du dossier : une ligne par pièce, et une seule par groupe.

    Les alternatives n'y comptent pas : supprimer les surnuméraires laisse la
    pagination juste, et c'est tout l'objet de la décision D3.
    """
    return [
        {
            "numero": groupe.code or "—",
            "titre": groupe.titre,
            "page": groupe.numero,
        }
        for groupe in groupes
    ]


def _pages_du_sommaire(groupes) -> dict:
    """Ce que la page de garde attend : un intervalle de pages par code."""
    pages = {"": 1}
    for groupe in groupes:
        pages[groupe.code] = (groupe.numero, groupe.numero + groupe.pages - 1)
    return pages


# ---------------------------------------------------------------------------
# L'écriture du fichier
# ---------------------------------------------------------------------------


def _ecrire(projet, groupes, notice, dossier, travail) -> Path:
    """Monte le fichier, diapo par diapo, dans l'ordre du dossier.

    `dossier` reçoit le `.pptx`, `travail` les PDF intermédiaires — dont celui de
    la page de garde, composée ici parce que son sommaire a besoin de la
    pagination complète.
    """
    presentation = Presentation()
    presentation.slide_width = Mm(LARGEUR_MM)
    presentation.slide_height = Mm(HAUTEUR_MM)

    # La page de garde ferme la composition et ouvre le fichier : son sommaire
    # porte les numéros de page réels, qui ne sont connus qu'une fois les pièces
    # paginées. C'est l'ordre de la voie PDF, et pour la même raison.
    garde = page_garde.generer(projet, travail, pages=_pages_du_sommaire(groupes))
    diapo_garde = Diapo(
        code="", numero=1, titre=TITRE_PAGE_GARDE, libelle=TITRE_PAGE_GARDE, voie="",
    )
    # La perspective de la page de garde est une réservation d'image, comme les
    # cadres photo : le chef de projet l'insère dans le fichier au lieu de la
    # déposer avant de générer. Sans photographie insérée, la planche garde son
    # cadre tireté et sa mention « Photomontage à insérer », que l'image posée
    # recouvre — il n'y a donc rien à dessiner autrement.
    _poser_diapo(
        presentation,
        PlanchePPTX(
            planche=garde.planche,
            cadres=[page_garde.zone_image()],
            libelle=TITRE_PAGE_GARDE,
        ),
        diapo_garde,
    )
    groupes.insert(
        0,
        Groupe("", 1, TITRE_PAGE_GARDE, planches=[], diapos=[diapo_garde]),
    )

    for groupe in groupes:
        if groupe.code == "DP 11":
            groupe.diapos = _poser_la_notice(presentation, notice, groupe)
            continue
        total = len(groupe.planches)
        for rang, planches in enumerate(groupe.planches, start=1):
            diapo = Diapo(
                code=groupe.code,
                numero=groupe.numero,
                titre=groupe.titre,
                libelle=planches.libelle or groupe.code,
                voie="",
                alternative=(rang, total),
                echelle=planches.echelle,
            )
            _poser_diapo(presentation, planches, diapo)
            groupe.diapos.append(diapo)

    fichier = dossier / nom_sortie(projet)
    presentation.save(str(fichier))
    return fichier


def _poser_diapo(presentation, planches: PlanchePPTX, diapo: Diapo) -> None:
    """Une diapo : son fond sur la mise en page, ses cadres et ses repères dessus."""
    fond = rendu_pptx.rendre_fond(planches.planche, diapo.libelle)
    diapo.voie = fond.voie
    diapo.octets = fond.octets
    # Compté ici plutôt qu'annoncé par l'appelant : le rapport disait « 0 cadre »
    # pour la page de garde, qui en porte un depuis que sa perspective est une
    # réservation. Un rapport qui décrit autre chose que le fichier ne sert à rien.
    diapo.cadres = len(planches.cadres)

    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, diapo.libelle)
    ooxml.poser_fond(mise_en_page, fond.image, fond.svg)
    for rang, cadre in enumerate(planches.cadres):
        ooxml.ajouter_reservation_image(
            mise_en_page, PREMIER_IDX_CADRE + rang,
            f"Cadre photo {rang + 1}", *cadre,
        )

    diapositive = presentation.slides.add_slide(mise_en_page)
    for cadre in diapositive.placeholders:
        # Le verrou se pose sur la forme de la diapo : posé sur la mise en page,
        # il ne suit pas le clonage (mesuré le 26/09/2026).
        ooxml.verrouiller(cadre)
    for cadre, intitule in zip(planches.cadres, planches.intitules):
        _poser_intitule(diapositive, cadre, intitule)
    if planches.centre_panneau is not None:
        _poser_les_reperes(diapositive, planches)
    if diapo.alternative[1] > 1:
        _poser_le_bandeau(diapositive, planches, diapo.alternative[1])


def _poser_la_notice(presentation, notice, groupe) -> list:
    """Une diapo par page de la notice, rastérisée depuis le PDF produit."""
    if notice is None:
        return []
    pages = rendu_pptx.rendre_pages_pdf(notice.chemin, "DP 11")
    diapos = []
    for rang, image in enumerate(pages, start=1):
        libelle = "DP 11" if len(pages) == 1 else f"DP 11, page {rang}"
        mise_en_page = ooxml.nouvelle_mise_en_page(presentation, libelle)
        ooxml.poser_fond(mise_en_page, image)
        presentation.slides.add_slide(mise_en_page)
        diapos.append(
            Diapo(
                code=groupe.code,
                numero=groupe.numero + rang - 1,
                titre=groupe.titre,
                libelle=libelle,
                voie=rendu_pptx.MATRICIELLE,
                octets=rendu_pptx.poids_dans_le_paquet(image),
            )
        )
    return diapos


def _poser_intitule(diapositive, cadre: tuple, intitule: str) -> None:
    """L'intitulé d'un cadre photo, en zone de texte éditable.

    Il est posé dans la bande que `sous_cadre` laisse au titre, au même corps que
    celui qu'imprime la voie PDF. Verrouillé en position : laissé libre, il se
    déplaçait au moindre clic (mesuré le 26/09/2026), et un intitulé glissé sur
    la photographie voisine désigne la mauvaise vue.
    """
    from .planches.primitives import HAUTEUR_TITRE_CADRE_MM

    # La bande de titre est celle que `sous_cadre` laisse vide : elle s'arrête où
    # l'image commence, et fait `HAUTEUR_TITRE_CADRE_MM` de haut. Le cadre est
    # posé à la même abscisse que l'image, qui est déjà en retrait de la marge du
    # sous-cadre — retirer cette marge une seconde fois posait l'intitulé sur le
    # filet du cadre, 3,5 mm trop haut (mesuré sur le fichier produit).
    x_mm, y_mm, largeur_mm, _ = cadre
    boite = diapositive.shapes.add_textbox(
        Mm(x_mm), Mm(y_mm - HAUTEUR_TITRE_CADRE_MM),
        Mm(largeur_mm), Mm(HAUTEUR_TITRE_CADRE_MM),
    )
    cadre_texte = boite.text_frame
    cadre_texte.word_wrap = False
    cadre_texte.margin_left = cadre_texte.margin_right = 0
    cadre_texte.margin_top = cadre_texte.margin_bottom = 0
    cadre_texte.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraphe = cadre_texte.paragraphs[0]
    paragraphe.text = intitule
    _habiller(paragraphe, CORPS_TITRE_CADRE_PT, gras=True)
    ooxml.verrouiller(boite)


def _poser_les_reperes(diapositive, planches: PlanchePPTX) -> None:
    """Un groupe « repère + cône + étiquette » par vue attendue.

    Posés au **centre du panneau de repérage**, pas en marge : un cône oublié doit
    se voir (décision D5). Ils sont légèrement décalés les uns des autres pour
    qu'on puisse les attraper séparément — superposés, le second serait
    inatteignable.
    """
    centre_x, centre_y = planches.centre_panneau
    ecart = 2 * RAYON_CONE_MM
    depart = centre_x - ecart * (len(planches.reperes) - 1) / 2.0
    for rang, repere in enumerate(planches.reperes):
        _poser_un_repere(diapositive, depart + rang * ecart, centre_y, repere)


def _poser_un_repere(diapositive, x_mm: float, y_mm: float, repere: str) -> None:
    """Le cône, le disque et l'étiquette, groupés et tournant autour du point.

    Les tailles sont celles de la voie PDF, en **millimètres papier** : 2,4 mm de
    rayon pour le disque, un secteur de 50° et 8 mm pour le cône. Elles ne
    dépendent donc pas de l'échelle du repérage, et le groupe n'a jamais à être
    redimensionné — le redimensionner serait mentir sur ce qu'il mesure.

    Le groupe tourne autour du centre de son cadre. C'est le cône qui fixe ce
    cadre : son `prstGeom` « pie » occupe la boîte du cercle entier, donc un carré
    de 2 x `RAYON_CONE_MM` déjà centré sur le point. Le disque et l'étiquette
    tiennent dedans, et le pivot tombe sur le point de vue.
    """
    groupe = diapositive.shapes.add_group_shape()

    cone = groupe.shapes.add_shape(
        MSO_SHAPE.PIE,
        Mm(x_mm - RAYON_CONE_MM), Mm(y_mm - RAYON_CONE_MM),
        Mm(2 * RAYON_CONE_MM), Mm(2 * RAYON_CONE_MM),
    )
    _regler_le_secteur(cone, OUVERTURE_CONE_DEG)
    cone.fill.solid()
    cone.fill.fore_color.rgb = RGBColor.from_string(MAGENTA_VISEE)
    ooxml.regler_alpha(cone, OPACITE_CONE)
    cone.line.color.rgb = RGBColor.from_string(MAGENTA_VISEE)
    cone.line.width = Pt(0.85)
    cone.shadow.inherit = False

    disque = groupe.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Mm(x_mm - RAYON_REPERE_MM), Mm(y_mm - RAYON_REPERE_MM),
        Mm(2 * RAYON_REPERE_MM), Mm(2 * RAYON_REPERE_MM),
    )
    disque.fill.solid()
    disque.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    disque.line.color.rgb = RGBColor.from_string("1A1A1A")
    disque.line.width = Pt(1.1)
    disque.shadow.inherit = False

    # L'étiquette **sous** le disque, là où le cône ne part pas : le secteur vise
    # le nord au départ, et une étiquette posée au nord comme sur la planche PDF
    # s'y serait superposée d'emblée. Elle tient dans la boîte du cône, pour que le
    # pivot du groupe ne bouge pas. `upright` et `wrap="none"` la gardent
    # horizontale même quand le groupe tourne de côté : le premier seul la replie à
    # une lettre par ligne (mesuré le 25/09/2026).
    etiquette = groupe.shapes.add_textbox(
        Mm(x_mm - RAYON_CONE_MM), Mm(y_mm + RAYON_REPERE_MM + 0.4),
        Mm(2 * RAYON_CONE_MM), Mm(4.4),
    )
    cadre_texte = etiquette.text_frame
    cadre_texte.word_wrap = False
    cadre_texte.margin_left = cadre_texte.margin_right = 0
    cadre_texte.margin_top = cadre_texte.margin_bottom = 0
    paragraphe = cadre_texte.paragraphs[0]
    paragraphe.text = repere
    paragraphe.alignment = PP_ALIGN.CENTER
    _habiller(paragraphe, CORPS_ETIQUETTE_PT, gras=True)
    ooxml.garder_horizontal(etiquette)

    groupe.name = f"Repère {repere}"


def _regler_le_secteur(forme, ouverture_deg: float) -> None:
    """Ouvre le secteur `pie` de `ouverture_deg`, pointé vers le nord.

    Les deux poignées d'un `pie` sont des angles comptés depuis trois heures dans
    le sens des aiguilles : 270° vise le nord de la planche, qui est aussi celui
    de la flèche du cartouche. C'est l'orientation de départ, et le chef de projet
    fait tourner le groupe vers ce que la photographie regarde.

    Les deux poignées restent positives et dans l'ordre croissant — un secteur
    écrit de -25° à +25° enjamberait l'origine, et rien ne garantit que PowerPoint
    le lise comme nous.
    """
    demi = ouverture_deg / 2.0
    forme.adjustments[0] = (270.0 - demi) * DEGRE_EN_POIGNEE
    forme.adjustments[1] = (270.0 + demi) * DEGRE_EN_POIGNEE


def _poser_le_bandeau(diapositive, planches: PlanchePPTX, nombre: int) -> None:
    """Le bandeau rouge d'une diapo alternative, posé sur la diapo donc supprimable.

    Il couvre la bande de titre du panneau de repérage, et rien d'autre : ni un
    cadre photo, ni son intitulé, ni le cartouche. C'est la seule bande de la
    planche dont l'information se retrouve ailleurs — l'échelle du repérage est
    aussi au cartouche —, et un bandeau qui cacherait un cadre empêcherait de voir
    à quoi ce cadre sert.
    """
    from .planches.primitives import HAUTEUR_TITRE_CADRE_MM, MARGE_SOUS_CADRE_MM

    x_mm, y_mm, largeur_mm, _ = planches.panneau
    boite = diapositive.shapes.add_textbox(
        Mm(x_mm), Mm(y_mm),
        Mm(largeur_mm), Mm(HAUTEUR_TITRE_CADRE_MM + 2 * MARGE_SOUS_CADRE_MM),
    )
    boite.fill.solid()
    boite.fill.fore_color.rgb = RGBColor.from_string(ROUGE_BANDEAU)
    cadre_texte = boite.text_frame
    cadre_texte.word_wrap = True
    cadre_texte.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraphe = cadre_texte.paragraphs[0]
    paragraphe.text = MENTION_ALTERNATIVE.format(nombre=nombre)
    paragraphe.alignment = PP_ALIGN.CENTER
    _habiller(paragraphe, CORPS_BANDEAU_PT, gras=True, couleur="FFFFFF")
    boite.name = "Bandeau à supprimer"


def _habiller(paragraphe, corps_pt: float, gras: bool = False,
              couleur: str = "1A1A1A") -> None:
    """La police du dossier sur un paragraphe de la diapo.

    Aptos, et non `polices.famille_active()` — la nuance compte. La planche est
    rastérisée ici, par cairo, qui peut n'avoir qu'un repli sous la main ; le
    texte éditable, lui, sera composé là-bas, par le PowerPoint du chef de projet,
    où Aptos est la police par défaut de Microsoft 365 depuis 2024 (vérifié à
    l'écran le 25/09/2026). Y écrire le repli de notre poste ferait composer du
    Carlito sur une machine qui a l'Aptos du dossier. Si les deux diffèrent,
    `RapportPPTX.etat_polices` le dit, comme pour la voie PDF.
    """
    for passage in paragraphe.runs or [paragraphe.add_run()]:
        passage.font.name = POLICE_PRINCIPALE
        passage.font.size = Pt(corps_pt)
        passage.font.bold = gras
        passage.font.color.rgb = RGBColor.from_string(couleur)
