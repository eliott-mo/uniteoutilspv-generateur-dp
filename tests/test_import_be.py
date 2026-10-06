"""Tests de l'import du plan du bureau d'études (lot 2bis, étapes A et B).

Les valeurs attendues ont été relevées à la main sur les deux fichiers de
référence du projet Saint-Cyr-en-Val avant toute écriture de code.
"""

from __future__ import annotations

from pathlib import Path

import ezdxf
import pytest

from dp_socle.erreurs import ErreurGeoreferencement, ErreurImportBE, ErreurTableauBilan
from dp_socle.import_be import azimut_tables, categorie_proposee, lire_plan_be, normaliser
from dp_socle.tableau_bilan import indice_depuis_nom, indices_disponibles, lire_tableau

REFERENCE = Path("exemples/saint-cyr-DXF")
DXF = REFERENCE / "20260903_SCV_IND06.dxf"
TABLEAU = REFERENCE / "20260825_SCV_Tableau_Bilan_V6.xlsx"


@pytest.fixture(scope="module")
def plan():
    return lire_plan_be(DXF)


@pytest.fixture(scope="module")
def tableau():
    return lire_tableau(TABLEAU, "IND06")


# ---------------------------------------------------------------------------
# Étape A — DXF
# ---------------------------------------------------------------------------


def test_unite_deduite_des_coordonnees_et_non_de_l_entete(plan):
    """$INSUNITS annonce des millimètres, les coordonnées sont en mètres.

    Un outil qui croirait l'en-tête diviserait tout le plan par mille.
    """
    document = ezdxf.readfile(str(DXF))
    assert document.header["$INSUNITS"] == 4  # 4 = millimètres
    assert plan.unite == "mètre"
    assert plan.facteur_unite == 1.0


def test_nombre_de_tables(plan):
    assert plan.nb_tables == 96


def test_surface_cloturee(plan):
    assert plan.surface_cloturee_m2 / 10_000 == pytest.approx(4.522, abs=0.002)


def test_lineaire_de_cloture(plan):
    assert plan.lineaire_cloture_m == pytest.approx(871.9, abs=0.5)


def test_nombre_de_portails(plan):
    """Trois portails, dessinés chacun en cinq entités jointives."""
    assert len(plan.geometries("portail")) == 15
    assert plan.nb_portails == 3


def test_azimut_des_tables(plan):
    """Rangées est-ouest sur le fichier de référence."""
    assert plan.azimut_tables_deg == pytest.approx(0.0, abs=1.0)


def test_les_tables_portent_une_altitude_reelle(plan):
    tables = plan.par_categorie("tables_pv")
    assert all(entite.z_reel for entite in tables)
    assert 90.0 < min(e.z_min for e in tables) < 110.0
    # La clôture est une polyligne 2D : son Z ne porte aucune altitude et le
    # lot 4 ne doit pas le confondre avec le terrain.
    assert not plan.par_categorie("cloture")[0].z_reel


def test_les_hatch_ne_doublent_pas_les_polylignes(plan):
    """Le DXF porte 13 HATCH qui redessinent des polylignes déjà importées."""
    document = ezdxf.readfile(str(DXF))
    assert sum(1 for e in document.modelspace() if e.dxftype() == "HATCH") == 13
    assert len(plan.par_categorie("pdl_ptr")) == 3  # 3 LWPOLYLINE, 3 HATCH ignorés


def test_les_bulges_des_polylignes_sont_discretises(plan):
    """Les plateformes portent des arcs en bulge, invisibles dans les sommets.

    Sans discrétisation, la plateforme du PDL sortirait à 133,8 m² au lieu de
    131,4 : c'est le genre d'écart qui ne se voit pas sur la planche.
    """
    plateformes = sorted(
        (e.geometrie.area for e in plan.par_categorie("plateforme")), reverse=True
    )
    assert plateformes[0] == pytest.approx(131.4, abs=0.5)


def test_les_calques_vides_sont_ignores_sans_message(plan):
    assert "PVcase PV Modules (full frames)" in plan.calques_vides
    assert "PVcase PV Modules (detailed)" in plan.calques_vides
    assert plan.avertissements == []


def test_les_calques_d_un_autre_bureau_d_etudes_sont_reconnus():
    """Tous les plans ne viennent pas du BE interne.

    Celui de Joux-la-Ville (24/09/2026) est dessiné par ORKANE : trois de ses
    cinq calques restaient non appariés, dont la clôture. Sans emprise
    clôturée, le placement de la coupe s'arrêtait sur un `AttributeError`, et
    le chef de projet recevait une trace Python au lieu d'un dossier.

    La tolérance de `normaliser` vaut ici comme ailleurs : accent, casse et
    séparateurs ne doivent pas décider de ce qui entre au dossier.
    """
    assert categorie_proposee("ORKANE_Cloture") == "cloture"
    assert categorie_proposee("ORKANE_Clôture") == "cloture"
    assert categorie_proposee("orkane cloture") == "cloture"
    assert categorie_proposee("ORKANE_Local Technique") == "local_technique"
    # « PVcase Road » allait à rien le 28/09, au motif que lourde ou légère ne se
    # devine pas. La décision était juste et la conséquence fausse : un calque
    # non apparié n'est pas importé du tout, et deux plans sont sortis sans
    # leur voirie — Joux-la-Ville et Bédarieux. La catégorie `voirie` existe
    # précisément pour ça : elle fait entrer la géométrie au dossier **et**
    # pose la question au chef de projet, au lieu de perdre l'élément en
    # silence.
    assert categorie_proposee("PVcase Road") == "voirie"


def test_correspondance_tolerante_aux_accents_et_separateurs():
    """L'apostrophe de « aire d'aspiration » est un tiret dans le fichier BE."""
    assert normaliser("UNI_Clôture") == normaliser("UNI_Cloture")
    assert normaliser("UNI_SDIS_Aire_d-aspiration") == normaliser(
        "UNI SDIS Aire d'aspiration"
    )
    assert categorie_proposee("UNI-CLOTURE") == "cloture"
    assert categorie_proposee("uni_sdis_aire d'aspiration") == "aire_aspiration"


CALQUE_TABLES = "PVcase PV Modules (optimised)"


def _ajouter_table(espace, origine) -> None:
    """Une table est-ouest : sans elle, l'azimut n'est pas calculable."""
    x, y = origine
    espace.add_lwpolyline(
        [(x, y), (x + 15, y), (x + 15, y + 4.6), (x, y + 4.6)],
        close=True,
        dxfattribs={"layer": CALQUE_TABLES},
    )


def _ajouter_table_en_segments(espace, origine, altitude=100.0) -> None:
    """La même table, mais éclatée en quatre segments cotés en Z.

    C'est ce que porte le plan d'Auzainvilliers du 23/09/2026 : 316 `LINE` là
    où Saint-Cyr porte 96 `POLYLINE`. Les altitudes montent d'un coin à l'autre
    comme sur un terrain en pente, pour que `z_reel` ait quelque chose à lire.
    """
    x, y = origine
    coins = [
        (x, y, altitude),
        (x + 15, y, altitude + 0.3),
        (x + 15, y + 4.6, altitude + 0.5),
        (x, y + 4.6, altitude + 0.2),
    ]
    for depart, arrivee in zip(coins, coins[1:] + coins[:1]):
        espace.add_line(depart, arrivee, dxfattribs={"layer": CALQUE_TABLES})


def _dxf_de_tables(chemin: Path, en_segments: bool, refermees=True) -> Path:
    """Un plan réduit à sa clôture et à deux rangées, dessinées d'une façon ou de l'autre."""
    document = ezdxf.new(setup=True)
    document.layers.add(CALQUE_TABLES)
    espace = document.modelspace()
    x, y = 622_900.0, 6_750_700.0
    espace.add_lwpolyline(
        [(x - 5, y - 5), (x + 25, y - 5), (x + 25, y + 35), (x - 5, y + 35)],
        close=True,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    for rang in range(2):
        origine = (x, y + rang * 12)
        if not en_segments:
            _ajouter_table(espace, origine)
        elif refermees:
            _ajouter_table_en_segments(espace, origine)
        else:
            # Trois côtés sur quatre : rien ne se referme, rien n'est une table.
            xa, ya = origine
            coins = [(xa, ya, 100.0), (xa + 15, ya, 100.0), (xa + 15, ya + 4.6, 100.0)]
            for depart, arrivee in zip(coins, coins[1:]):
                espace.add_line(depart, arrivee, dxfattribs={"layer": CALQUE_TABLES})
    document.saveas(str(chemin))
    return chemin


def test_les_tables_livrees_en_segments_sont_refermees(tmp_path):
    """Un calque de tables éclaté en segments donne quand même des rangées.

    Relevé sur le plan d'Auzainvilliers le 26/09/2026 : les 316 `LINE` du
    calque formaient 79 rectangles complets, mais une ligne n'a pas de surface.
    L'import ne trouvait aucune table, ne pouvait pas orienter la coupe A-A' et
    refusait le plan entier, alors que la géométrie y était tout entière.
    """
    plan = lire_plan_be(_dxf_de_tables(tmp_path / "segments.dxf", en_segments=True))

    assert len(plan.tables) == 2
    for table in plan.tables:
        assert table.area == pytest.approx(15.0 * 4.6, rel=1e-6)
    assert plan.azimut_tables_deg == pytest.approx(0.0, abs=1e-9)

    # Le Z des segments est rendu aux sommets : sans lui, la coupe DP 3
    # reposerait les tables sur des hauteurs de catalogue.
    entites = [e for e in plan.entites if e.categorie == "tables_pv"]
    assert len(entites) == 2
    assert all(e.z_reel for e in entites)
    assert min(e.z_min for e in entites) == pytest.approx(100.0)
    assert max(e.z_max for e in entites) == pytest.approx(100.5)

    # Et la reconstruction se dit, avec ce qu'il faut demander au BE.
    recompositions = [m for m in plan.avertissements if "recomposée" in m]
    assert len(recompositions) == 1
    assert "2 rangée(s)" in recompositions[0]
    # La formule que l'interface reconnaît pour rassembler les demandes en un
    # bloc recopiable, comme pour les contrôles croisés plus bas.
    assert "à demander au bureau d'études" in recompositions[0]


def test_un_contour_de_table_dessine_fait_foi(tmp_path):
    """Une polyligne fermée ne passe pas par la recomposition, et rien ne le dit.

    Même garde-fou que pour les aires de grutage : la reconstruction ne vaut
    qu'à défaut de polygone. Un plan qui passait doit continuer de passer par
    le même chemin qu'avant.
    """
    plan = lire_plan_be(_dxf_de_tables(tmp_path / "polylignes.dxf", en_segments=False))

    assert len(plan.tables) == 2
    assert not [m for m in plan.avertissements if "recomposée" in m]


def test_des_segments_qui_ne_referment_rien_ne_deviennent_pas_des_tables(tmp_path):
    """Trois côtés sur quatre : le plan est refusé, pas complété d'office.

    `polygonize` ne rend une surface que de ce qui se referme. Inventer le
    quatrième côté donnerait une rangée que le BE n'a pas dessinée, et une
    surface de modules fausse dans les contrôles croisés.
    """
    chemin = _dxf_de_tables(tmp_path / "ouverts.dxf", en_segments=True, refermees=False)
    with pytest.raises(ErreurImportBE, match="Aucune table de modules"):
        lire_plan_be(chemin)


def _dxf_minimal(chemin: Path, calque: str, origine=(622_900.0, 6_750_700.0)) -> Path:
    """DXF d'un rectangle de 10 x 4 m sur `calque`, plus une table, pour les cas d'erreur."""
    document = ezdxf.new(setup=True)
    document.layers.add(calque)
    espace = document.modelspace()
    x, y = origine
    espace.add_lwpolyline(
        [(x, y), (x + 10, y), (x + 10, y + 4), (x, y + 4)],
        close=True,
        dxfattribs={"layer": calque},
    )
    _ajouter_table(espace, (x, y + 20))
    document.saveas(str(chemin))
    return chemin


def test_dxf_hors_des_bornes_lambert_93_refuse(tmp_path):
    chemin = _dxf_minimal(
        tmp_path / "hors_bornes.dxf", "UNI_Clôture", origine=(1_200.0, 3_400.0)
    )
    with pytest.raises(ErreurGeoreferencement) as erreur:
        lire_plan_be(chemin)
    assert "Lambert 93" in str(erreur.value)


def test_calque_cloture_avec_variante_d_accentuation_reconnu(tmp_path):
    chemin = _dxf_minimal(tmp_path / "variante.dxf", "uni-cloture")
    plan_variante = lire_plan_be(chemin)
    assert plan_variante.correspondance["uni-cloture"] == "cloture"
    assert plan_variante.surface_cloturee_m2 == pytest.approx(40.0)


def test_calque_uniquement_en_hatch_signale(tmp_path):
    """Un élément qui n'existe qu'en remplissage serait perdu en silence."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    espace.add_lwpolyline(
        [(622_900, 6_750_700), (622_910, 6_750_700), (622_910, 6_750_704)],
        close=True,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    remplissage = espace.add_hatch(dxfattribs={"layer": "UNI_VRD_Plateforme"})
    remplissage.paths.add_polyline_path(
        [(622_920, 6_750_700), (622_930, 6_750_700), (622_930, 6_750_710)], is_closed=True
    )
    _ajouter_table(espace, (622_900, 6_750_720))
    chemin = tmp_path / "hatch_seul.dxf"
    document.saveas(str(chemin))

    plan_hatch = lire_plan_be(chemin)
    messages = " ".join(plan_hatch.avertissements)
    assert "UNI_VRD_Plateforme" in messages and "HATCH" in messages
    assert plan_hatch.par_categorie("plateforme") == []


def test_sommets_dupliques_en_plan_dedoublonnes_dans_l_ordre(tmp_path):
    """Un `set()` détruirait la géométrie en perdant l'ordre du tracé."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    sommets = [
        (622_900, 6_750_700, 10.0),
        (622_910, 6_750_700, 11.0),
        (622_910, 6_750_700, 12.0),  # doublon en plan, extrusion 3D
        (622_910, 6_750_704, 12.0),
        (622_900, 6_750_704, 11.0),
    ]
    espace.add_polyline3d(
        sommets, close=True, dxfattribs={"layer": "PVcase PV Modules (optimised)"}
    )
    chemin = tmp_path / "doublons.dxf"
    document.saveas(str(chemin))

    plan_doublons = lire_plan_be(chemin)
    table = plan_doublons.tables[0]
    assert len(table.exterior.coords) == 5  # 4 sommets + fermeture
    assert table.area == pytest.approx(40.0)
    assert any("dupliqué" in m for m in plan_doublons.avertissements)


def test_azimut_robuste_au_repliement_a_90_degres():
    """Des rangées nord-sud ne doivent pas donner une médiane à 0°.

    Les angles voisins de ±90° tombent aux deux bouts de l'intervalle : une
    médiane naïve rendrait la direction perpendiculaire à la vraie.
    """
    from shapely.geometry import box
    from shapely.affinity import rotate

    tables = [
        rotate(box(0, 0, 15, 4.6), 90 + ecart, origin="centroid")
        for ecart in (-0.3, -0.1, 0.0, 0.1, 0.4)
    ]
    assert abs(azimut_tables(tables)) == pytest.approx(90.0, abs=0.5)


def test_calque_non_apparie_signale(tmp_path):
    chemin = _dxf_minimal(tmp_path / "inconnu.dxf", "UNI_Chose_Inconnue")
    plan_inconnu = lire_plan_be(chemin)
    assert plan_inconnu.calques_ignores == ["UNI_Chose_Inconnue"]
    assert any("non apparié" in m for m in plan_inconnu.avertissements)


def test_plan_sans_table_refuse(tmp_path):
    """Sans table, ni l'azimut ni la ligne de coupe n'ont de sens."""
    document = ezdxf.new(setup=True)
    document.modelspace().add_lwpolyline(
        [(622_900, 6_750_700), (622_910, 6_750_700), (622_910, 6_750_704)],
        close=True,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    chemin = tmp_path / "sans_table.dxf"
    document.saveas(str(chemin))
    with pytest.raises(ErreurImportBE) as erreur:
        lire_plan_be(chemin)
    assert "Aucune table" in str(erreur.value)


# ---------------------------------------------------------------------------
# Étape B — tableau bilan
# ---------------------------------------------------------------------------


def test_indice_propose_depuis_le_nom_du_dxf():
    assert indice_depuis_nom("20260903_SCV_IND06.dxf") == "IND06"
    assert indice_depuis_nom("plan_sans_indice.dxf") is None


def test_indices_disponibles():
    assert indices_disponibles(TABLEAU) == ["IND05", "IND06"]


def test_indice_absent_refuse():
    with pytest.raises(ErreurTableauBilan) as erreur:
        lire_tableau(TABLEAU, "IND09")
    assert "IND05, IND06" in str(erreur.value)


def test_parametres_de_l_indice_06(tableau):
    assert tableau.generalites["phase"] == "APS"
    assert tableau.generalites["date"].isoformat() == "2026-08-25"
    assert tableau.generalites["surface_cloturee_ha"] == pytest.approx(4.52)
    assert tableau.generalites["lineaire_cloture_m"] == pytest.approx(875.0)
    assert tableau.generalites["nb_portails"] == 3
    assert tableau.structures["nb_tables"] == 96
    assert tableau.structures["inclinaison_deg"] == pytest.approx(15.0)
    assert tableau.structures["point_bas_m"] == pytest.approx(2.5)
    assert tableau.structures["point_haut_m"] == pytest.approx(4.0)
    assert tableau.modules["nb_modules"] == 4628
    assert tableau.modules["puissance_mwc"] == pytest.approx(2.93878)


def test_les_colonnes_d_indice_ne_sont_pas_confondues():
    """IND05 et IND06 diffèrent : lire la mauvaise colonne serait invisible."""
    ind05 = lire_tableau(TABLEAU, "IND05")
    assert ind05.structures["nb_tables"] == 126
    assert ind05.modules["puissance_mwc"] == pytest.approx(3.99399)


def test_azimut_avec_symbole_degre(tableau):
    """La cellule vaut la chaîne « 0° », pas le nombre 0."""
    assert tableau.structures["azimut_deg"] == pytest.approx(0.0)


def test_valeur_composite_rendue_en_couple(tableau):
    assert tableau.structures["nb_modules_longueur"] == (13, 26)
    assert tableau.structures["nb_modules_table"] == (26, 52)


def test_cotes_normalisees_avec_leur_ordre_de_lecture(tableau):
    """L'ordre des cotes change d'une section à l'autre et doit être conservé."""
    par_ouvrage = {c.ouvrage: c for c in tableau.cotes}
    ptr = par_ouvrage["Poste de transformation - PTR"]
    pdl_ptr = par_ouvrage["Poste de livraison et de transformation - PDL/PTR"]
    assert ptr.dimensions == "10 x 3 x 3m"
    assert ptr.ordre_cotes == "largeur x longueur x hauteur"
    assert pdl_ptr.dimensions == "12 x 3 x 3m"
    assert pdl_ptr.ordre_cotes == "longueur x largeur x hauteur"
    assert par_ouvrage["Aire d'aspiration"].surface_m2 == pytest.approx(32.0)
    assert par_ouvrage["Aire de charge (BESS) — Conteneurs"].dimensions == "6 x 3 x 3m"


def test_standards_du_type_de_projet(tableau):
    assert tableau.standards["Format modules"] == "G12R"
    assert tableau.standards["Inclinaison"] == 15


#: De quoi remplir chaque paramètre selon son convertisseur.
_VALEURS_D_ESSAI = {
    "_texte": "APS",
    "_entier": 1,
    "_nombre": 1.0,
    "_date": "01/01/2026",
}


def _tableau_fabrique(chemin, decoupees: bool, onglet_cotes: str | None = None):
    """Un tableau bilan complet, écrit depuis la liste des paramètres attendus.

    Fabriqué et non recopié : le jeu réel pèse 7 Mo, et un fixture qui figerait
    sa mise en page cesserait de suivre `PARAMETRES` le jour où un paramètre s'y
    ajoute. `decoupees` choisit la mise en page de la piste enherbée — deux
    lignes interne et externe, comme Saint-Cyr, ou une seule, comme le tableau
    d'Auzainvilliers.
    """
    import openpyxl

    from dp_socle.tableau_bilan import PARAMETRES, normaliser

    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Projet"
    feuille["A1"] = "Projet d'essai"
    feuille["B1"] = "IND03"

    rang, vus = 2, set()
    for definitions in PARAMETRES.values():
        for _cle, libelle, convertir, obligatoire in definitions:
            if normaliser(libelle) in vus or (not obligatoire and not decoupees):
                continue
            vus.add(normaliser(libelle))
            feuille.cell(rang, 1, libelle)
            feuille.cell(rang, 2, _VALEURS_D_ESSAI.get(convertir.__name__, 1.0))
            rang += 1
    if not decoupees:
        feuille.cell(rang, 1, "Surface piste enherbée (m²)")
        feuille.cell(rang, 2, 420.0)

    if onglet_cotes is not None:
        cotes = classeur.create_sheet(onglet_cotes)
        cotes.cell(3, 3, "Poste de transformation - PTR")
        cotes.cell(4, 3, "Dimensions  (largeur x longueur x hauteur)")
        cotes.cell(4, 4, "Surface (m²)")
        cotes.cell(4, 5, "Surface plateforme (m²)")
        cotes.cell(5, 3, "10 x 3 x 3m")
        cotes.cell(5, 4, 30)
        cotes.cell(5, 5, 115)

    classeur.save(chemin)
    return chemin


def test_un_tableau_qui_ne_decoupe_pas_la_piste_enherbee_se_lit(tmp_path):
    """Une seule ligne « Surface piste enherbée (m²) » suffit.

    Le tableau d'Auzainvilliers (V3, onglet « Projet ») ne porte ni interne ni
    externe. Le lecteur exigeait les deux et arrêtait le dossier entier sur une
    ligne qui, sur ce projet, était vide — alors que tout le reste du tableau
    était lisible (relevé le 27/09/2026).
    """
    from dp_socle.tableau_bilan import surface_piste_legere_m2

    chemin = _tableau_fabrique(tmp_path / "unique.xlsx", decoupees=False)
    tableau = lire_tableau(chemin, "IND03")

    assert surface_piste_legere_m2(tableau.pistes) == pytest.approx(420.0)
    # La lecture se dit : lire un total sur une ligne unique n'est pas la même
    # chose que l'additionner depuis deux lignes.
    assert any("ligne unique" in m for m in tableau.avertissements)


def test_le_decoupage_interne_externe_fait_foi_quand_il_existe(tmp_path):
    """Sur la mise en page de Saint-Cyr, rien ne change : le détail est lu."""
    from dp_socle.tableau_bilan import surface_piste_legere_m2

    chemin = _tableau_fabrique(tmp_path / "decoupe.xlsx", decoupees=True)
    tableau = lire_tableau(chemin, "IND03")

    assert tableau.pistes["surface_piste_legere_interne_m2"] == pytest.approx(1.0)
    assert surface_piste_legere_m2(tableau.pistes) == pytest.approx(2.0)
    assert not [m for m in tableau.avertissements if "ligne unique" in m]


@pytest.mark.parametrize("onglet", ["Dimensions postes et pieux", "Dimensions postes"])
def test_les_cotes_se_trouvent_sous_les_deux_noms_d_onglet(tmp_path, onglet):
    """Certains tableaux recopient la table des gabarits sous son nom d'origine.

    Celui d'Auzainvilliers porte « Dimensions postes », aux mêmes sections et
    aux mêmes lignes que le « Dimensions postes et pieux » de Saint-Cyr. Ne
    chercher que le second privait le dossier de toutes ses cotes DP 4, et
    l'avertissement accusait un onglet absent qui était là.
    """
    chemin = _tableau_fabrique(
        tmp_path / f"cotes_{onglet[-6:]}.xlsx", decoupees=True, onglet_cotes=onglet
    )
    tableau = lire_tableau(chemin, "IND03")

    # Les cotes **lues** sont celles du tableau ; le reste du catalogue est
    # complété depuis le classeur des gabarits depuis le 06/10/2026, et se
    # distingue par `cotes_completees`.
    lues = [c for c in tableau.cotes if c.ouvrage not in tableau.cotes_completees]
    assert [c.ouvrage for c in lues] == ["Poste de transformation - PTR"]
    assert lues[0].dimensions == "10 x 3 x 3m"
    assert lues[0].surface_m2 == pytest.approx(30.0)


def test_paramètre_absent_leve_plutot_que_de_rendre_une_valeur_par_defaut(tmp_path):
    """Un tableau dont la mise en page a changé doit échouer, pas deviner."""
    import openpyxl

    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "2. Caractéristiques du projet"
    feuille["A1"] = "Projet incomplet"
    feuille["B1"] = "IND06"
    feuille["A2"] = "Phase "
    feuille["B2"] = "APS"
    chemin = tmp_path / "incomplet.xlsx"
    classeur.save(chemin)

    with pytest.raises(ErreurTableauBilan) as erreur:
        lire_tableau(chemin, "IND06")
    assert "introuvable" in str(erreur.value)


# ---------------------------------------------------------------------------
# Références de bloc — révélées par les plans de Sarnois
# ---------------------------------------------------------------------------


class _NamespaceSansCalque:
    """Espace de noms DXF dont `layer` lève, comme celui d'un GEOMAPIMAGE."""

    @property
    def layer(self):
        raise ezdxf.DXFAttributeError(
            'Invalid DXF attribute "layer" for entity GEOMAPIMAGE'
        )


class _EntiteSansCalque:
    """Entité DXF sans attribut de calque.

    Une doublure plutôt qu'un vrai GEOMAPIMAGE : ezdxf n'offre pas de fabrique
    pour cette entité, et le plan de Sarnois qui en porte un pèse 47 Mo, trop
    lourd pour l'historique du dépôt.
    """

    def __init__(self):
        self.dxf = _NamespaceSansCalque()

    def dxftype(self):
        return "GEOMAPIMAGE"


def test_calque_absent_ne_leve_pas_d_exception_nue():
    """Toutes les entités DXF ne portent pas d'attribut de calque.

    Le plan de Sarnois contient un `GEOMAPIMAGE`, l'image géoréférencée du fond
    de plan, qui n'en a pas. Y accéder levait une `DXFAttributeError` nue,
    remontée jusqu'à l'écran en trace brute.
    """
    from dp_socle.import_be import _calque_de

    assert _calque_de(_EntiteSansCalque()) is None


def test_entites_sans_calque_ecartees_et_comptees(tmp_path):
    from dp_socle.import_be import _developper

    anomalies: dict = {}
    entites = [_EntiteSansCalque(), _EntiteSansCalque()]
    assert list(_developper(entites, None, 0, anomalies)) == []
    assert anomalies["sans_calque"] == {"GEOMAPIMAGE": 2}


def _dxf_avec_bloc(chemin: Path, calque_contenu: str, rotation: float = 0.0) -> Path:
    """DXF dont la table est une référence de bloc, comme chez le BE.

    Sur Sarnois, les 85 tables, les 4 portails et la citerne sont des `INSERT` :
    dessinés à plat, ils seraient invisibles à l'import.
    """
    document = ezdxf.new(setup=True)
    bloc = document.blocks.new("TABLE_PV")
    bloc.add_lwpolyline(
        [(0, 0), (15, 0), (15, 4.6), (0, 4.6)],
        close=True,
        dxfattribs={"layer": calque_contenu},
    )
    espace = document.modelspace()
    espace.add_blockref(
        "TABLE_PV",
        (622_900, 6_750_700),
        dxfattribs={"layer": CALQUE_TABLES, "rotation": rotation},
    )
    espace.add_lwpolyline(
        [(622_880, 6_750_680), (622_960, 6_750_680), (622_960, 6_750_760)],
        close=True,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    document.saveas(str(chemin))
    return chemin


def test_table_en_reference_de_bloc_importee(tmp_path):
    chemin = _dxf_avec_bloc(tmp_path / "bloc.dxf", "PVcase PV Modules (optimised)")
    plan_bloc = lire_plan_be(chemin)
    assert plan_bloc.nb_tables == 1
    assert plan_bloc.tables[0].area == pytest.approx(69.0, abs=0.1)


def test_calque_zero_dans_un_bloc_herite_de_l_insertion(tmp_path):
    """Convention AutoCAD qu'ezdxf 1.4.4 n'applique pas de lui-même.

    Vérifié le 03/09/2026 : `virtual_entities()` rend une sous-entité déclarée
    sur « 0 » toujours sur « 0 ». Sans cette règle, tout le contenu des blocs
    dessinés sur le calque 0 tomberait dans un fourre-tout et serait perdu.
    """
    chemin = _dxf_avec_bloc(tmp_path / "herite.dxf", "0")
    plan_herite = lire_plan_be(chemin)
    assert plan_herite.nb_tables == 1
    assert plan_herite.par_categorie("tables_pv")[0].calque == CALQUE_TABLES


def test_rotation_de_la_reference_de_bloc_appliquee(tmp_path):
    """`virtual_entities()` applique la transformation de l'insertion.

    Les 85 tables de Sarnois sont posées à 245,5° : sans cette résolution,
    elles ressortiraient toutes parallèles à l'axe du dessin.
    """
    from dp_socle.import_be import azimut_tables

    droit = lire_plan_be(_dxf_avec_bloc(tmp_path / "droit.dxf", "0", rotation=0.0))
    tourne = lire_plan_be(_dxf_avec_bloc(tmp_path / "tourne.dxf", "0", rotation=24.5))
    assert azimut_tables(droit.tables) == pytest.approx(0.0, abs=0.01)
    assert azimut_tables(tourne.tables) == pytest.approx(24.5, abs=0.01)


def test_blocs_imbriques_developpes(tmp_path):
    document = ezdxf.new(setup=True)
    interieur = document.blocks.new("CONTOUR")
    interieur.add_lwpolyline(
        [(0, 0), (15, 0), (15, 4.6), (0, 4.6)], close=True, dxfattribs={"layer": "0"}
    )
    exterieur = document.blocks.new("TABLE")
    exterieur.add_blockref("CONTOUR", (0, 0), dxfattribs={"layer": "0"})
    espace = document.modelspace()
    espace.add_blockref("TABLE", (622_900, 6_750_700), dxfattribs={"layer": CALQUE_TABLES})
    espace.add_lwpolyline(
        [(622_880, 6_750_680), (622_960, 6_750_680), (622_960, 6_750_760)],
        close=True,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    chemin = tmp_path / "imbrique.dxf"
    document.saveas(str(chemin))

    plan_imbrique = lire_plan_be(chemin)
    assert plan_imbrique.nb_tables == 1
    assert plan_imbrique.tables[0].area == pytest.approx(69.0, abs=0.1)


def test_cloture_ouverte_signalee(tmp_path):
    """Refermer un contour ouvert invente un segment que le BE n'a pas dessiné.

    Le drapeau `closed` du DXF ne suffit pas à trancher : sur Sarnois la
    clôture l'a à faux alors que son tracé revient sur son point de départ. Ce
    qui compte est la distance entre les deux extrémités, pas le drapeau.
    """
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    espace.add_lwpolyline(
        [
            (622_900, 6_750_700),
            (622_980, 6_750_700),
            (622_980, 6_750_780),
            (622_900, 6_750_760),
        ],
        close=False,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    _ajouter_table(espace, (622_910, 6_750_710))
    chemin = tmp_path / "cloture_ouverte.dxf"
    document.saveas(str(chemin))

    plan_ouvert = lire_plan_be(chemin)
    messages = " ".join(plan_ouvert.avertissements)
    assert "contour de clôture" in messages and "ouvert" in messages
    assert "60.0 m" in messages  # écart entre les deux extrémités


def test_cloture_fermee_sans_drapeau_ne_declenche_rien(tmp_path):
    """Le cas réel de Sarnois : drapeau à faux, tracé refermé à la main."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    espace.add_lwpolyline(
        [
            (622_900, 6_750_700),
            (622_980, 6_750_700),
            (622_980, 6_750_780),
            (622_900, 6_750_700),
        ],
        close=False,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    _ajouter_table(espace, (622_910, 6_750_710))
    chemin = tmp_path / "cloture_refermee.dxf"
    document.saveas(str(chemin))

    plan_ferme = lire_plan_be(chemin)
    assert not any("ouvert" in m for m in plan_ferme.avertissements)
    assert plan_ferme.surface_cloturee_m2 == pytest.approx(3200.0, abs=1.0)


# ---------------------------------------------------------------------------
# Tableau bilan V10 — révélé par Sarnois
# ---------------------------------------------------------------------------

TABLEAU_V10 = Path("exemples/sarnois-1-DXF/20260827_Sarnois_Tableau_Bilan_V10.xlsx")


def test_onglet_renomme_entre_les_versions_du_tableau():
    """« 2. Caractéristiques du projet » en V6, « Projet » en V10."""
    import openpyxl

    from dp_socle.tableau_bilan import ONGLETS_CARACTERISTIQUES

    v10 = openpyxl.load_workbook(TABLEAU_V10, read_only=True).sheetnames
    assert "2. Caractéristiques du projet" not in v10
    assert "Projet" in v10
    assert set(ONGLETS_CARACTERISTIQUES) >= {"2. Caractéristiques du projet", "Projet"}

    assert indices_disponibles(TABLEAU_V10) == [
        "IND06",
        "IND07",
        "IND08",
        "IND09",
        "IND10A",
        "IND10B",
    ]
    assert indices_disponibles(TABLEAU) == ["IND05", "IND06"]


def test_onglet_inconnu_refuse_en_listant_les_onglets(tmp_path):
    import openpyxl

    classeur = openpyxl.Workbook()
    classeur.active.title = "Autre chose"
    chemin = tmp_path / "inconnu.xlsx"
    classeur.save(chemin)

    with pytest.raises(ErreurTableauBilan) as erreur:
        indices_disponibles(chemin)
    assert "Autre chose" in str(erreur.value)


def test_valeur_composite_au_perluete_sommee():
    """Un projet à deux formats de table décrit son parc en deux nombres.

    « 269 & 27 » vaut 296 tables. La chaîne d'origine reste lisible à côté :
    la décomposition ne doit pas disparaître dans la somme.
    """
    tableau_v10 = lire_tableau(TABLEAU_V10, "IND09")
    assert tableau_v10.structures["nb_tables"] == 296
    assert tableau_v10.structures["nb_tables_brut"] == "269 & 27"
    assert tableau_v10.structures["nb_modules_longueur"] == (13, 26)


def test_valeur_explicitement_non_renseignee_admise_sur_le_gcr():
    """Le tableau met « / » là où le GCR n'a pas été calculé.

    Admis parce qu'aucun contrôle n'en dépend. Ailleurs, une case illisible
    reste une erreur — rendre None en silence masquerait le problème.
    """
    from dp_socle.tableau_bilan import _nombre, _nombre_ou_absent

    assert lire_tableau(TABLEAU_V10, "IND09").modules["gcr"] is None
    assert lire_tableau(TABLEAU_V10, "IND06").modules["gcr"] == pytest.approx(0.4214, abs=1e-4)
    assert _nombre_ou_absent("/", "GCR (%)") is None
    with pytest.raises(ErreurTableauBilan):
        _nombre("/", "GCR (%)")


def test_azimut_negatif_du_tableau_v10():
    """Le projet SE dont le tableau écrit « -24.5 » sans direction."""
    tableau_v10 = lire_tableau(TABLEAU_V10, "IND09")
    assert tableau_v10.structures["azimut_deg"] == pytest.approx(-24.5)
    assert tableau_v10.structures["azimut_brut"] == "-24.5"


def test_indice_a_suffixe_alphabetique_reconnu():
    """`IND10A` et `IND10B` étaient écartés de l'en-tête **sans rien dire**.

    Le motif n'acceptait qu'un nombre. Les deux colonnes de Sarnois, deux
    réductions du projet sous les 3 MWc, disparaissaient de la liste des
    indices, et l'outil laissait croire que le plan était en avance sur le
    tableau alors que les colonnes étaient là.
    """
    assert indice_depuis_nom("2026_08_025-IMP-DEV-Fixe-IND10a.dxf") == "IND10A"
    assert indice_depuis_nom("2026_08_27-IMP-DEV-Fixe-IND10b.dxf") == "IND10B"
    assert indice_depuis_nom("20260903_SCV_IND06.dxf") == "IND06"


def test_indice_ambigu_ne_rend_rien_plutot_qu_au_hasard():
    """Sans garde, « IND06final » rendrait l'indice « IND06F »."""
    assert indice_depuis_nom("truc_IND06final.dxf") is None
    assert indice_depuis_nom("plan_sans_indice.dxf") is None


def test_variantes_d_un_meme_indice_lues_separement():
    """Les deux réductions de Sarnois diffèrent, et les confondre serait invisible."""
    a = lire_tableau(TABLEAU_V10, "IND10A")
    b = lire_tableau(TABLEAU_V10, "IND10B")

    assert a.structures["nb_tables"] == 85 and a.structures["nb_tables_brut"] == "84 & 1"
    assert b.structures["nb_tables"] == 90 and b.structures["nb_tables_brut"] == "79 & 11"
    assert a.generalites["surface_cloturee_ha"] == pytest.approx(4.87)
    assert b.generalites["surface_cloturee_ha"] == pytest.approx(3.48)
    assert a.generalites["lineaire_cloture_m"] == pytest.approx(949.0)
    assert b.generalites["lineaire_cloture_m"] == pytest.approx(738.0)
    # Deux variantes calibrées pour passer sous le seuil de 3 MWc.
    for variante in (a, b):
        assert variante.modules["puissance_mwc"] == pytest.approx(2.98792)
        assert variante.generalites["nb_portails"] == 1
        assert variante.structures["azimut_deg"] == pytest.approx(-24.5)


# ---------------------------------------------------------------------------
# Charte des calques — élargie sur les plans de Sarnois
# ---------------------------------------------------------------------------


def test_ecran_de_correspondance_voit_ce_que_l_import_apparie(tmp_path):
    """Les deux lectures doivent voir les mêmes calques.

    `calques_du_dxf` ne développait pas les blocs : les tables et les portails
    de Sarnois n'apparaissaient donc pas à l'écran où l'on confirme
    l'appariement, alors que l'import les trouvait.
    """
    from dp_socle.import_be import calques_du_dxf

    chemin = _dxf_avec_bloc(tmp_path / "ecran.dxf", "PVcase PV Modules (optimised)")
    vus = {c.nom for c in calques_du_dxf(chemin)}
    apparies = set(lire_plan_be(chemin).correspondance)
    assert apparies <= vus
    assert "PVcase PV Modules (optimised)" in vus


def test_calques_ecartes_regroupes_en_une_ligne(tmp_path):
    """45 calques non appariés faisaient 71 avertissements, donc aucun lu."""
    from dp_socle.import_be import motif_ecart

    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_700))
    for calque in ("CAD_1PARCELLE", "CAD_3BATIDUR", "UNI_Legende", "UNI_Echelle"):
        espace.add_lwpolyline(
            [(622_800, 6_750_600), (622_810, 6_750_600), (622_810, 6_750_610)],
            close=True,
            dxfattribs={"layer": calque},
        )
    chemin = tmp_path / "ecartes.dxf"
    document.saveas(str(chemin))

    plan_ecarte = lire_plan_be(chemin)
    groupes = [m for m in plan_ecarte.avertissements if "calque(s) écarté(s)" in m]
    assert len(groupes) == 2  # fond cadastral, mobilier de dessin
    assert not [m for m in plan_ecarte.avertissements if "non apparié" in m]
    assert motif_ecart("CAD_1PARCELLE") == "fond cadastral du BE, remplacé par le WFS IGN"
    assert motif_ecart("UNI_Clôture") is None


def test_calque_vraiment_inconnu_garde_son_avertissement(tmp_path):
    """Écarter en groupe ne doit pas noyer ce qui demande une décision."""
    chemin = _dxf_minimal(tmp_path / "inconnu.dxf", "UNI_Chose_Inconnue")
    plan_inconnu = lire_plan_be(chemin)
    assert any(
        "UNI_Chose_Inconnue" in m and "non apparié" in m
        for m in plan_inconnu.avertissements
    )


def test_annotations_regroupees(tmp_path):
    """Textes et cotations produisaient un avertissement par calque."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_700))
    for indice, calque in enumerate(("UNI_Clôture", "UNI_PDL", "UNI_VRD_Plateforme")):
        espace.add_text(
            f"étiquette {indice}", dxfattribs={"layer": calque}
        ).set_placement((622_800 + indice, 6_750_600))
    chemin = tmp_path / "annotations.dxf"
    document.saveas(str(chemin))

    plan_annote = lire_plan_be(chemin)
    groupes = [m for m in plan_annote.avertissements if "annotation(s) écartée(s)" in m]
    assert len(groupes) == 1
    assert "3 TEXT" in groupes[0]


def test_portail_d_exploitation_distinct_du_portail_d_acces():
    """Le tableau ne compte que les portails d'accès.

    Les confondre donnait 4 portails au plan contre 1 au tableau sur Sarnois,
    et faisait échouer un contrôle qui avait raison de se plaindre.
    """
    from dp_socle.import_be import CATEGORIES, categorie_proposee

    assert categorie_proposee("UNI_portail") == "portail"
    assert categorie_proposee("UNI_Portail exploitant") == "portail_exploitant"
    assert "portail_exploitant" in CATEGORIES


def test_le_calque_des_chemins_existants_est_reconnu_sans_gonfler_la_voie_lourde():
    """Le béton déjà en place entre au dossier sans entrer dans les comptes.

    Saint-Cyr, 29/09/2026 : le chef de projet demande que les chemins bétonnés
    existants figurent au plan. Le tableau bilan ne les compte pas — 3 362 m²
    de voie lourde au plan pour 3 365 déclarés, soit l'anneau seul. Les
    additionner à la voie lourde ferait échouer un contrôle qui a raison.
    """
    from dp_socle.import_be import (
        CATEGORIES,
        CATEGORIES_VOIE_LOURDE,
        categorie_proposee,
    )

    assert categorie_proposee("UNI_VRD_Chemin_existant") == "chemin_existant"
    assert "chemin_existant" in CATEGORIES
    assert "chemin_existant" not in CATEGORIES_VOIE_LOURDE


def test_noms_de_calque_de_sarnois_apparies():
    """La charte du BE n'est pas figée d'un projet à l'autre."""
    from dp_socle.import_be import categorie_proposee

    attendus = {
        "UNI_Cloture": "cloture",
        "UNI_Haies": "haie",
        "UNI_Haies existantes": "haie_existante",
        "UNI_Local_Stockage": "local_technique",
        "UNI_VRD_Voirie": "voirie",
        "UNI_PDT": "ptr",
        "UNI_BESS_Batterie": "bess",
        "UNI_BESS_Rétention": "bac_retention",
    }
    for calque, categorie in attendus.items():
        assert categorie_proposee(calque) == categorie, calque


def test_saint_cyr_reste_sans_avertissement(plan):
    """La charte élargie ne doit rien changer au jeu de référence."""
    assert plan.avertissements == []
    assert len(plan.correspondance) == 9


def test_le_beton_en_place_de_saint_cyr_ne_gonfle_pas_la_voie_lourde():
    """« ENV_Surface bétonnée » va au chemin existant, pas à la piste lourde.

    Relevé le 01/10/2026 : le bureau d'études a ajouté le calque à la demande
    du chef de projet, qui l'appariait à la piste lourde existante — le choix
    naturel au vu du libellé. Mais c'est du béton déjà en place, que le tableau
    bilan ne compte pas : le contrôle de voie lourde passait à **135,7 %**
    d'écart pour 5 % admis. En chemin existant il tombe à 0,1 %, et les
    7 930 m² paraissent quand même au plan, sous leur propre entrée de légende.
    """
    from dp_socle.import_be import CATEGORIES_VOIE_LOURDE, categorie_proposee

    assert categorie_proposee("ENV_Surface bétonnée") == "chemin_existant"
    assert categorie_proposee("UNI_VRD_Piste_enherbée") == "piste_legere"
    assert "chemin_existant" not in CATEGORIES_VOIE_LOURDE


def test_un_remplissage_sans_contour_porte_quand_meme_sa_surface():
    """Six surfaces bétonnées sur sept disparaissaient sans un mot.

    Les HATCH sont écartés partout, et pour une raison mesurée : sur le fichier
    de référence, chacun doublait une polyligne du même calque, à l'aire
    identique au centième de mètre carré. Relevé le 30/09/2026 sur Saint-Cyr :
    le calque « ENV_Surface bétonnée » porte **sept** remplissages — 2 796,
    3 261, 394, 415, 412, 460 et 192 m² — et **une seule** polyligne, qui
    double le premier. 5 134 m² de béton étaient perdus, et le chef de projet
    ne voyait pas ses pistes.

    Ce qui est repris, c'est le contour que le remplissage porte lui-même :
    lire, et non inventer. Trois garde-fous le bornent, mesurés ici.
    """
    from shapely.geometry import LineString, Polygon

    from dp_socle.import_be import EntiteBE, _contours_des_remplissages

    class _Chemin:
        def __init__(self, sommets):
            self.vertices = sommets

    class _Remplissage:
        def __init__(self, sommets):
            self.paths = [_Chemin(sommets)]

    def carre(x, y, cote):
        return [(x, y), (x + cote, y), (x + cote, y + cote), (x, y + cote)]

    def entite(geometrie):
        return EntiteBE(
            categorie="piste_lourde_existante",
            calque="ENV_Surface bétonnée",
            geometrie=geometrie,
            z_reel=False,
        )

    messages = []
    # Un remplissage à l'écart de tout tracé : il entre.
    repris = _contours_des_remplissages(
        [_Remplissage(carre(0.0, 0.0, 10.0)), _Remplissage(carre(100.0, 0.0, 20.0))],
        [entite(Polygon(carre(0.0, 0.0, 10.0)))],
        1.0,
        "ENV_Surface bétonnée",
        messages,
    )
    assert [round(p.area) for p in repris] == [400], "le doublon reste dehors"
    assert any("1 remplissage(s) repris" in m for m in messages)

    # Le doublon se mesure sur le plus petit des deux : une polyligne
    # discrétisée est plus petite que le remplissage aux bords droits.
    messages = []
    assert not _contours_des_remplissages(
        [_Remplissage(carre(0.0, 0.0, 10.0))],
        [entite(Polygon(carre(0.5, 0.5, 9.0)))],
        1.0,
        "ENV_Surface bétonnée",
        messages,
    )
    assert not messages

    # Un calque qui n'a que des traits est un calque de symbole : le
    # remplissage n'y est que la teinte du dessin.
    messages = []
    assert not _contours_des_remplissages(
        [_Remplissage(carre(0.0, 0.0, 10.0))],
        [entite(LineString([(0.0, 0.0), (10.0, 10.0)]))],
        1.0,
        "UNI_portail",
        messages,
    )
    assert not messages


def test_variantes_de_nommage_du_premier_lot_de_dossiers_reels():
    """Quatre calques non appariés, quatre éléments absents du plan de masse.

    Relevé le 29/09/2026 sur Auzainvilliers et Bédarieux : le chef de projet
    remontait « pas de PDL », « pas de zone de remisage », « piste lourde
    inexistante », « haie à créer inexistante » comme autant de défauts de
    l'outil. C'étaient quatre noms de calque que la correspondance ne
    connaissait pas, et l'import les avait tous nommés.

    `PVcase Road` va à `voirie` et non à une piste : son nom ne dit pas si
    elle est lourde ou légère, et `voirie` est la catégorie qui pose la
    question plutôt que d'y répondre.
    """
    from dp_socle.import_be import categorie_proposee

    attendus = {
        "UNI_PDL-PTR": "pdl_ptr",
        "UNI_Zone de remise": "zone_remise",
        "UNI_VRD_Piste_lourde_créée": "piste_lourde_a_creer",
        "UNI_Haie créée": "haie",
        "PVcase Road": "voirie",
    }
    for calque, categorie in attendus.items():
        assert categorie_proposee(calque) == categorie, calque

    # Et les noms voisins déjà connus n'ont pas bougé : les deux chartes
    # coexistent, chaque projet employant la sienne.
    assert categorie_proposee("UNI_PDL") == "pdl_ptr"
    assert categorie_proposee("UNI_VRD_Piste_lourde_à_créer") == "piste_lourde_a_creer"
    assert categorie_proposee("UNI_Haie") == "haie"


def test_categories_de_chantier_agri_et_bess():
    """Vocabulaire élargi aux objets réellement présents chez le BE."""
    from dp_socle.import_be import CATEGORIES, categorie_proposee

    attendus = {
        "UNI_VRD_Base_vie": "base_vie",
        "UNI_VRD_Stockage_Logistique": "stockage_chantier",
        "UNI_Limite paddock": "limite_paddock",
        "UNI_Bac d'équarissage": "bac_equarrissage",
        "ESPACE VERT": "espace_vert",
        "UNI_BESS_Refroidissement": "citerne_refroidissement",
        "UNI_BESS_Zone_remise": "zone_remise",
    }
    for calque, categorie in attendus.items():
        assert categorie_proposee(calque) == categorie, calque
        assert categorie in CATEGORIES


def test_zone_d_implantation_potentielle_ecartee():
    """Seule la clôture délimite le projet au dossier."""
    from dp_socle.import_be import categorie_proposee, motif_ecart

    assert categorie_proposee("UNI_ZIP") is None
    assert "clôture" in motif_ecart("UNI_ZIP")


def test_remplissage_doublon_des_plateformes_ecarte():
    """« VAL-PDL » redouble « UNI_VRD_Plateforme ».

    Mesuré sur Sarnois : ses deux boucles ont le même centre et la même aire
    que celles des plateformes, à 1 m² près (81,8 contre 82,9 m² et 94,4 contre
    95,5 m²). Les importer dessinerait chaque plateforme deux fois.
    """
    from dp_socle.import_be import motif_ecart

    assert motif_ecart("VAL-PDL") == "doublon du remplissage des plateformes"


def test_calque_ecarte_ne_produit_pas_d_avertissement_de_hatch(tmp_path):
    """Un calque écarté l'est pour toutes ses formes, remplissages compris."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_700))
    remplissage = espace.add_hatch(dxfattribs={"layer": "VAL-PDL"})
    remplissage.paths.add_polyline_path(
        [(622_920, 6_750_700), (622_930, 6_750_700), (622_930, 6_750_710)],
        is_closed=True,
    )
    chemin = tmp_path / "ecarte_hatch.dxf"
    document.saveas(str(chemin))

    plan_ecarte = lire_plan_be(chemin)
    assert not any("VAL-PDL" in m and "HATCH" in m for m in plan_ecarte.avertissements)
    assert any("VAL-PDL" in m and "écarté" in m for m in plan_ecarte.avertissements)


def test_toutes_les_categories_ont_un_style_et_un_rang_de_dessin():
    """Le contrat et l'aperçu doivent rester alignés.

    Une catégorie sans style ne se verrait pas sur l'écran de validation, et
    c'est justement là qu'un appariement fautif doit sauter aux yeux.
    """
    from dp_socle.apercu_be import ORDRE_DESSIN, STYLES
    from dp_socle.import_be import CATEGORIES

    assert set(CATEGORIES) == set(STYLES)
    assert set(CATEGORIES) - {"ligne_coupe"} == set(ORDRE_DESSIN)


def test_calque_par_defaut_ecarte_mais_appariable(tmp_path):
    """Le calque 0 est le calque de travail du BE, censé être vide à l'export.

    Ce qui s'y trouve y est par oubli : sur l'indice A de Sarnois une zone de
    contention de 70 m², sur l'indice B des traits de construction. Il est donc
    écarté d'office — mais reste appariable à la main, ce qui permet de
    récupérer l'ouvrage égaré sans faire rééditer le fichier au BE.
    """
    from dp_socle.import_be import motif_ecart

    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_700))
    espace.add_lwpolyline(
        [(622_940, 6_750_700), (622_948, 6_750_700), (622_948, 6_750_710)],
        close=True,
        dxfattribs={"layer": "0"},
    )
    chemin = tmp_path / "calque_zero.dxf"
    document.saveas(str(chemin))

    assert "calque de travail" in motif_ecart("0")
    plan_ecarte = lire_plan_be(chemin)
    assert not plan_ecarte.par_categorie("zone_contention")
    assert any("0" in m and "écarté" in m for m in plan_ecarte.avertissements)

    # Rattrapage manuel, sans que le BE ait à rééditer son export.
    plan_rattrape = lire_plan_be(
        chemin,
        correspondance={CALQUE_TABLES: "tables_pv", "0": "zone_contention"},
    )
    recuperees = plan_rattrape.par_categorie("zone_contention")
    assert len(recuperees) == 1
    assert recuperees[0].geometrie.area == pytest.approx(40.0)


def test_les_trois_types_de_poste_sont_distincts():
    """PDL seul, PTR seul, PDL/PTR combiné : trois légendes au dossier.

    Le tableau bilan les compte séparément et leur donne des cotes distinctes —
    9 x 3 m pour le PDL, 10 x 3 m pour le PTR, 12 x 3 m pour le combiné. Les
    confondre dans une seule catégorie interdisait au lot 4 de les dessiner
    différemment.
    """
    from dp_socle.import_be import CATEGORIES, categorie_proposee

    assert {"pdl", "ptr", "pdl_ptr"} <= set(CATEGORIES)
    assert categorie_proposee("UNI_PDT") == "ptr"
    # Le poste de transformation de l'aire de charge BESS est un PTR.
    assert categorie_proposee("UNI_BESS_PTR") == "ptr"


def test_le_calque_pdl_porte_en_fait_un_poste_combine():
    """Le nom du calque dit « PDL », le tableau dit « PDL/PTR ».

    Sur Saint-Cyr comme sur Sarnois : 1 PDL/PTR, 0 PDL, 0 PTR. C'est le tableau
    qui tranche, pas le nom du calque — et un projet portant un PDL seul
    demandera de corriger l'appariement à l'écran.
    """
    from dp_socle.import_be import categorie_proposee

    assert categorie_proposee("UNI_PDL") == "pdl_ptr"
    tableau_scv = lire_tableau(TABLEAU, "IND06")
    assert tableau_scv.postes["nb_pdl_ptr"] == 1
    assert tableau_scv.postes["nb_pdl"] == 0
    assert tableau_scv.postes["nb_ptr"] == 0


def test_zone_de_contention_avec_les_elements_agrivoltaiques():
    from dp_socle.import_be import categorie_proposee

    assert categorie_proposee("UNI_Zone de contention") == "zone_contention"
    assert categorie_proposee("UNI_Limite paddock") == "limite_paddock"


# ---------------------------------------------------------------------------
# Exports triés du BE — livrés le 04/09/2026
# ---------------------------------------------------------------------------


def test_pistes_lourdes_et_legeres_distinguees():
    """Le BE sépare désormais ce que la légende du dossier sépare.

    Sans cette distinction, le lot 4 ne pouvait pas donner aux deux les deux
    gris de la légende de référence. Les intitulés retenus sont ceux du plan du
    BE : « voie lourde » et « piste légère ».
    """
    from dp_socle.import_be import CATEGORIES, categorie_proposee

    assert categorie_proposee("UNI_VRD_Pistes lourdes") == "piste_lourde"
    assert categorie_proposee("UNI_VRD_Pistes légères") == "piste_legere"
    assert categorie_proposee("UNI_VRD_Aire de grutage") == "aire_grutage"
    assert {"piste_lourde", "piste_legere", "aire_grutage"} <= set(CATEGORIES)


def test_arbres_existants_importes_et_non_ecartes():
    """« Arbres existant » figure à la légende du BE : c'est de la végétation
    en place, qui a sa place au dossier au même titre que les haies."""
    from dp_socle.import_be import categorie_proposee, motif_ecart

    assert categorie_proposee("PVcase Trees") == "arbre_existant"
    assert motif_ecart("PVcase Trees") is None


def test_maillage_3d_rendu_par_son_emprise_en_plan(tmp_path):
    """`ezdxf.path` refuse les maillages, et l'import échouait dessus.

    Les arbres du BE en sont. Ce qui se dessine d'un arbre sur un plan de masse
    est son houppier vu du dessus : l'enveloppe convexe des sommets le donne.
    Attention aux enregistrements de face, qui n'encodent que des indices —
    les compter comme des points donnait des houppiers de plusieurs millions
    de mètres carrés.
    """
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_700))
    maillage = espace.add_polyface(dxfattribs={"layer": "PVcase Trees"})
    maillage.append_face(
        [
            (622_950, 6_750_700, 0),
            (622_954, 6_750_700, 0),
            (622_954, 6_750_704, 3),
            (622_950, 6_750_704, 3),
        ]
    )
    chemin = tmp_path / "arbre.dxf"
    document.saveas(str(chemin))

    plan_arbre = lire_plan_be(chemin)
    arbres = plan_arbre.par_categorie("arbre_existant")
    assert len(arbres) == 1
    assert arbres[0].geometrie.geom_type == "Polygon"
    assert arbres[0].geometrie.area == pytest.approx(16.0, abs=0.1)


def _dxf_grutage(chemin: Path, diagonales: bool, largeur=12.0, longueur=15.0) -> Path:
    """DXF portant une aire de grutage, en croix de diagonales ou en polygone."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_800))
    x, y = 622_900.0, 6_750_700.0
    coins = [(x, y), (x + longueur, y), (x + longueur, y + largeur), (x, y + largeur)]
    if diagonales:
        espace.add_lwpolyline(
            [coins[0], coins[2]], dxfattribs={"layer": "UNI_VRD_Aire de grutage"}
        )
        espace.add_lwpolyline(
            [coins[1], coins[3]], dxfattribs={"layer": "UNI_VRD_Aire de grutage"}
        )
    else:
        espace.add_lwpolyline(
            coins, close=True, dxfattribs={"layer": "UNI_VRD_Aire de grutage"}
        )
    document.saveas(str(chemin))
    return chemin


def test_aire_de_grutage_reconstruite_depuis_ses_diagonales(tmp_path):
    """Le BE ne dessine parfois qu'une croix : les deux diagonales du rectangle.

    Un rectangle est le seul quadrilatère dont les diagonales se coupent en leur
    milieu et ont même longueur. Mesuré sur Sarnois : 12,0 x 19,0 m et
    12,0 x 15,0 m, la seconde étant l'aire « 12x15m » de la légende du plan, et
    la part débordant de la voie lourde donne 263 m² pour 255 déclarés.
    """
    plan_croix = lire_plan_be(_dxf_grutage(tmp_path / "croix.dxf", diagonales=True))
    aires = plan_croix.par_categorie("aire_grutage")
    assert len(aires) == 1
    assert aires[0].geometrie.geom_type == "Polygon"
    assert aires[0].geometrie.area == pytest.approx(180.0, abs=0.1)
    assert any("reconstruite" in m for m in plan_croix.avertissements)


def test_un_contour_dessine_fait_toujours_foi(tmp_path):
    """La reconstruction ne s'applique qu'à défaut de polygone."""
    plan_ferme = lire_plan_be(_dxf_grutage(tmp_path / "ferme.dxf", diagonales=False))
    aires = plan_ferme.par_categorie("aire_grutage")
    assert len(aires) == 1
    assert aires[0].geometrie.area == pytest.approx(180.0, abs=0.1)
    assert not any("reconstruite" in m for m in plan_ferme.avertissements)


def test_deux_traits_de_longueurs_differentes_ne_font_pas_un_rectangle(tmp_path):
    """Sans le contrôle des longueurs, deux traits quelconques se croisant en
    leur milieu passeraient pour une aire de grutage."""
    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_800))
    # Même milieu, longueurs très différentes : ce n'est pas un rectangle.
    espace.add_lwpolyline(
        [(622_900, 6_750_700), (622_920, 6_750_700)],
        dxfattribs={"layer": "UNI_VRD_Aire de grutage"},
    )
    espace.add_lwpolyline(
        [(622_910, 6_750_695), (622_910, 6_750_705)],
        dxfattribs={"layer": "UNI_VRD_Aire de grutage"},
    )
    chemin = tmp_path / "faux_rectangle.dxf"
    document.saveas(str(chemin))

    plan_faux = lire_plan_be(chemin)
    aires = plan_faux.par_categorie("aire_grutage")
    assert all(e.geometrie.geom_type == "LineString" for e in aires)
    assert not any("reconstruite" in m for m in plan_faux.avertissements)


def test_l_aire_de_grutage_se_dessine_en_voie_lourde(tmp_path):
    """Un élargissement de voie, que la légende du dossier ne distingue pas.

    Le tableau bilan la compte en « supplément piste lourde » et la planche de
    référence n'a pas cette entrée : deux catégories dessinées à l'identique
    n'apparaissent qu'une fois en légende, leurs objets comptés ensemble.
    """
    from dp_socle.apercu_be import STYLES, legende_presente

    assert STYLES["aire_grutage"].libelle == STYLES["piste_lourde"].libelle
    assert STYLES["aire_grutage"].remplissage == STYLES["piste_lourde"].remplissage

    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    _ajouter_table(espace, (622_900, 6_750_800))
    espace.add_lwpolyline(
        [(622_900, 6_750_700), (622_915, 6_750_700), (622_915, 6_750_712)],
        close=True,
        dxfattribs={"layer": "UNI_VRD_Pistes lourdes"},
    )
    espace.add_lwpolyline(
        [(622_940, 6_750_700), (622_955, 6_750_700), (622_955, 6_750_712)],
        close=True,
        dxfattribs={"layer": "UNI_VRD_Aire de grutage"},
    )
    chemin = tmp_path / "legende.dxf"
    document.saveas(str(chemin))

    entrees = legende_presente(lire_plan_be(chemin))
    voies = [(s.libelle, nb) for _, s, nb in entrees if s.libelle == "Voie lourde"]
    assert voies == [("Voie lourde", 2)]


# ---------------------------------------------------------------------------
# La voirie dont le type n'est pas encore tranché
# ---------------------------------------------------------------------------


def _plan_d_essai(categorie: str, calque: str, geometrie):
    """Un plan réduit à une seule entité, pour éprouver un contrôle isolé."""
    from dp_socle.import_be import EntiteBE, PlanBE

    return PlanBE(
        entites=[EntiteBE(categorie, calque, geometrie, False)],
        azimut_tables_deg=0.0,
        correspondance={calque: categorie},
        unite="m",
        facteur_unite=1.0,
        calques_ignores=[],
        calques_vides=[],
        source="essai.dxf",
    )


def test_une_voirie_non_tranchee_suspend_le_recoupement_des_surfaces():
    """Une alerte qui se déclenche sur le cas ordinaire n'alerte plus de rien.

    Mesuré le 17/09/2026 sur Sarnois : son calque de voirie ne dit pas si les
    voies sont lourdes ou légères, tout tombe donc dans la catégorie `voirie`
    que le chef de projet répartit plus tard (D5 du lot 4). Les deux catégories
    précises étant vides et le tableau déclarant ses surfaces, le recoupement
    annonçait un écart de 100 % — sur un plan parfaitement normal.
    """
    from shapely.geometry import Polygon

    from dp_socle.import_be import AVERTISSEMENT, _surfaces_de_piste

    plan = _plan_d_essai("voirie", "UNI_VRD_Voirie", Polygon([(0, 0), (40, 0), (40, 5), (0, 5)]))
    controles = _surfaces_de_piste(
        plan,
        {"surface_piste_lourde_m2": 200.0, "surface_piste_legere_interne_m2": 0.0},
    )
    voie_lourde = next(c for c in controles if c.libelle == "Surface de voie lourde")
    assert voie_lourde.statut == AVERTISSEMENT
    assert "n'ont pas encore de type" in voie_lourde.message
    # Le total dessiné (200 m²) couvre le total déclaré : rien d'autre à dire
    # que la suspension du détail.
    assert "Demandez-les" not in voie_lourde.message
    assert voie_lourde.valeur_dxf is None


def test_une_voirie_tranchee_se_recoupe_normalement():
    """Saint-Cyr nomme ses voies « à créer » : le contrôle conclut, et le doit."""
    from shapely.geometry import Polygon

    from dp_socle.import_be import OK, _surfaces_de_piste

    plan = _plan_d_essai(
        "piste_lourde_a_creer",
        "UNI_VRD_Piste_lourde_a_creer",
        Polygon([(0, 0), (100, 0), (100, 10), (0, 10)]),
    )
    controles = _surfaces_de_piste(plan, {"surface_piste_lourde_m2": 1000.0})
    voie_lourde = next(c for c in controles if c.libelle == "Surface de voie lourde")
    assert voie_lourde.statut == OK
    assert voie_lourde.valeur_dxf == pytest.approx(1000.0)


def test_une_voirie_non_tranchee_ne_masque_pas_un_plan_qui_ne_la_dessine_pas():
    """Suspendre le détail ne doit pas faire taire ce qui reste mesurable.

    Mesuré le 17/09/2026 sur Sarnois : ses trois objets de voirie sont deux
    linéaires sans surface et un polygone de 36 m², contre 3 870 m² déclarés au
    tableau. Le plan ne dessine quasiment pas ses voies — les planches DP 2 et
    DP 4 ne les montreront pas — et un contrôle simplement « suspendu » l'aurait
    passé sous silence.
    """
    from shapely.geometry import Polygon

    from dp_socle.import_be import AVERTISSEMENT, _surfaces_de_piste

    plan = _plan_d_essai(
        "voirie", "UNI_VRD_Voirie", Polygon([(0, 0), (12, 0), (12, 3), (0, 3)])
    )
    controles = _surfaces_de_piste(
        plan,
        {
            "surface_piste_lourde_m2": 1423.0,
            "surface_piste_legere_interne_m2": 2447.0,
        },
    )
    voie_lourde = next(c for c in controles if c.libelle == "Surface de voie lourde")
    assert voie_lourde.statut == AVERTISSEMENT
    assert "n'ont pas encore de type" in voie_lourde.message
    # La formule est celle que l'interface reconnaît pour rassembler les
    # demandes en un bloc recopiable : la changer ici sans la changer dans
    # `app.MARQUEUR_BUREAU_ETUDES` ferait disparaître la demande de cette liste.
    assert "à demander au bureau d'études" in voie_lourde.message
    # Le total est mesuré, lui : 36 m² dessinés contre 3 870 déclarés.
    assert voie_lourde.valeur_dxf == pytest.approx(36.0)
    assert voie_lourde.valeur_tableau == pytest.approx(3870.0)


def test_le_recoupement_se_rejoue_une_fois_le_type_tranche():
    """Suspendre un contrôle ne doit pas le suspendre pour toute la vie du dossier.

    Demandé le 17/09/2026 : sans cette reprise, l'écart que ce contrôle doit
    attraper — un plan mis à jour sans son tableau — passait inaperçu dès lors
    que le calque du bureau d'études ne nommait pas le type de ses voies.
    """
    from shapely.geometry import Polygon

    from dp_socle.contrat import VOIRIE_SANS_OBJET
    from dp_socle.import_be import OK, controler_surfaces_de_voirie

    plan = _plan_d_essai(
        "voirie", "UNI_VRD_Voirie", Polygon([(0, 0), (100, 0), (100, 10), (0, 10)])
    )
    pistes = {"surface_piste_lourde_m2": 1000.0}

    # Avant le choix : le contrôle ne conclut pas, et le dit.
    suspendu = next(
        c for c in controler_surfaces_de_voirie(plan, pistes)
        if c.libelle == "Surface de voie lourde"
    )
    assert "n'ont pas encore de type" in suspendu.message

    # Après : les 1 000 m² dessinés recoupent les 1 000 m² déclarés.
    rejoue = next(
        c for c in controler_surfaces_de_voirie(plan, pistes, ["piste_lourde"])
        if c.libelle == "Surface de voie lourde"
    )
    assert rejoue.statut == OK
    assert rejoue.valeur_dxf == pytest.approx(1000.0)

    # Et rangée du mauvais côté, elle fait ressortir l'écart des deux côtés.
    mal_rangee = {
        c.libelle: c
        for c in controler_surfaces_de_voirie(plan, pistes, ["piste_legere"])
    }
    assert mal_rangee["Surface de voie lourde"].statut != OK


def test_un_objet_sans_objet_ne_compte_dans_aucune_surface():
    """Un axe rangé « sans objet » n'entre ni dans l'une ni dans l'autre."""
    from shapely.geometry import LineString

    from dp_socle.contrat import VOIRIE_SANS_OBJET
    from dp_socle.import_be import controler_surfaces_de_voirie

    plan = _plan_d_essai("voirie", "UNI_VRD_Voirie", LineString([(0, 0), (50, 0)]))
    controles = controler_surfaces_de_voirie(
        plan, {"surface_piste_lourde_m2": 900.0}, [VOIRIE_SANS_OBJET]
    )
    voie_lourde = next(c for c in controles if c.libelle == "Surface de voie lourde")
    # Rien n'est dessiné : l'écart porte sur les 900 m² déclarés, et le message
    # ne parle plus d'un type en attente — il n'y en a plus.
    assert "n'ont pas encore de type" not in voie_lourde.message


# ---------------------------------------------------------------------------
# Fond cadastral du plan du BE — contrôle croisé du foncier
# ---------------------------------------------------------------------------
#
# Tous les plans n'en portent pas : Saint-Cyr n'a aucun calque cadastral, celui
# de Sarnois en a vingt-quatre, aux noms d'un export EDIGÉO de la DGFiP. Le DXF
# de Sarnois n'est pas versionné — 50 Mo — et ces DXF synthétiques rejouent ce
# qu'il a montré, y compris son enclave.

_ORIGINE = (622_900.0, 6_750_700.0)


def _rectangle(x, y, largeur, hauteur):
    return [(x, y), (x + largeur, y), (x + largeur, y + hauteur), (x, y + hauteur)]


def _dxf_cadastre(
    chemin: Path,
    section: str | None = "ZA",
    numeros=("58", "59"),
    enclave: bool = False,
    cloture=(20, 20, 100, 60),
) -> Path:
    """Deux parcelles mitoyennes de 200 x 100 m, clôture dans la première.

    `enclave` ajoute dans la première parcelle une parcelle de 10 x 10 m,
    dessinée deux fois comme le fait un export DGFiP : comme parcelle, et comme
    trou de celle qui l'entoure.
    """
    from dp_socle.import_be import (
        CALQUE_NUMEROS_PARCELLE,
        CALQUE_NUMEROS_SECTION,
        CALQUE_PARCELLES,
        CALQUE_SECTIONS,
        CALQUE_TROUS_PARCELLE,
    )

    document = ezdxf.new(setup=True)
    espace = document.modelspace()
    x, y = _ORIGINE

    dx, dy, largeur, hauteur = cloture
    espace.add_lwpolyline(
        _rectangle(x + dx, y + dy, largeur, hauteur),
        close=True,
        dxfattribs={"layer": "UNI_Clôture"},
    )
    _ajouter_table(espace, (x + 40, y + 40))

    for rang, numero in enumerate(numeros):
        espace.add_lwpolyline(
            _rectangle(x + rang * 200, y, 200, 100),
            close=True,
            dxfattribs={"layer": CALQUE_PARCELLES},
        )
        if numero:
            espace.add_text(
                numero, dxfattribs={"layer": CALQUE_NUMEROS_PARCELLE}
            ).set_placement((x + rang * 200 + 160, y + 90))

    if enclave:
        coin = _rectangle(x + 150, y + 10, 10, 10)
        espace.add_lwpolyline(coin, close=True, dxfattribs={"layer": CALQUE_PARCELLES})
        espace.add_lwpolyline(
            coin, close=True, dxfattribs={"layer": CALQUE_TROUS_PARCELLE}
        )
        espace.add_text(
            "99", dxfattribs={"layer": CALQUE_NUMEROS_PARCELLE}
        ).set_placement((x + 155, y + 15))

    if section is not None:
        espace.add_lwpolyline(
            _rectangle(x - 50, y - 50, 500, 300),
            close=True,
            dxfattribs={"layer": CALQUE_SECTIONS},
        )
        espace.add_text(
            section, dxfattribs={"layer": CALQUE_NUMEROS_SECTION}
        ).set_placement((x + 10, y + 10))

    document.saveas(str(chemin))
    return chemin


def _parcelle_du_projet():
    """Le foncier maîtrisé : la première parcelle, telle que le géomètre la rend."""
    from shapely.geometry import Polygon

    x, y = _ORIGINE
    return Polygon(_rectangle(x, y, 200, 100))


def test_le_fond_cadastral_du_be_est_lu_sans_etre_dessine(tmp_path):
    """Écarté du dessin comme les points de terrain, lu pour la mesure.

    Le cadastre de la planche DP 1-3 vient du WFS IGN, qui fait foi ; celui du
    DXF est une copie de travail. Il ne remplace donc aucune géométrie du plan.
    """
    from dp_socle.import_be import CALQUE_PARCELLES

    plan = lire_plan_be(_dxf_cadastre(tmp_path / "cadastre.dxf"))
    assert [p.reference for p in plan.parcelles] == ["ZA 58", "ZA 59"]
    assert not any(e.calque == CALQUE_PARCELLES for e in plan.entites)


def test_un_plan_sans_calque_cadastral_ne_porte_aucune_parcelle(tmp_path):
    """Le cas de Saint-Cyr : l'absence n'est pas une anomalie."""
    plan = lire_plan_be(_dxf_minimal(tmp_path / "sans_cadastre.dxf", "UNI_Clôture"))
    assert plan.parcelles == []


def test_la_parcelle_sous_la_cloture_est_nommee_quand_l_emprise_concorde(tmp_path):
    """Le shapefile dit ce qui est maîtrisé, le cadastre du BE dit lequel c'est."""
    from dp_socle.import_be import OK, _emprise

    plan = lire_plan_be(_dxf_cadastre(tmp_path / "cadastre.dxf"))
    controle = _emprise(plan, _parcelle_du_projet())
    assert controle.statut == OK
    assert "la parcelle ZA 58" in controle.message
    # La parcelle voisine n'est pas sous la clôture : elle n'est pas nommée.
    assert "ZA 59" not in controle.message


def test_un_debordement_est_situe_sur_la_parcelle_qu_il_mord(tmp_path):
    """« déborde de 1 200 m² » ne dit pas quoi négocier ; « sur ZA 59 », si."""
    from shapely.geometry import Polygon

    from dp_socle.import_be import AVERTISSEMENT, _emprise

    # La clôture est à cheval sur les deux parcelles ; le foncier maîtrisé
    # s'arrête à la limite de la première.
    plan = lire_plan_be(
        _dxf_cadastre(tmp_path / "cadastre.dxf", cloture=(150, 20, 100, 60))
    )
    x, y = _ORIGINE
    controle = _emprise(plan, Polygon(_rectangle(x, y, 200, 100)))
    assert controle.statut == AVERTISSEMENT
    assert "ZA 59" in controle.message
    assert "maîtrise foncière" in controle.message


def test_sans_shapefile_le_cadastre_du_be_nomme_sans_conclure(tmp_path):
    """Le cadastre ne connaît pas la maîtrise foncière : il ne la remplace pas.

    Le contrôle reste donc en avertissement — mais il donne la liste à
    confronter à l'acte, au lieu de dire seulement qu'il n'a pas eu lieu.
    """
    from dp_socle.import_be import AVERTISSEMENT, _emprise

    plan = lire_plan_be(_dxf_cadastre(tmp_path / "cadastre.dxf"))
    controle = _emprise(plan, None)
    assert controle.statut == AVERTISSEMENT
    assert "la parcelle ZA 58" in controle.message
    assert "shapefile" in controle.message


def test_une_enclave_ne_prend_pas_le_numero_de_la_parcelle_qui_l_entoure(tmp_path):
    """Mesuré sur Sarnois : la parcelle 46 héritait du numéro de son enclave.

    Un export DGFiP dessine une parcelle enclavée deux fois — comme parcelle, et
    comme trou de l'englobante. Sans ôter le trou, deux numéros tombent dans
    l'englobante et sa référence devient illisible.
    """
    plan = lire_plan_be(_dxf_cadastre(tmp_path / "enclave.dxf", enclave=True))
    references = sorted(p.reference or "?" for p in plan.parcelles)
    assert references == ["ZA 58", "ZA 59", "ZA 99"]


def test_une_parcelle_sans_numero_sort_sans_reference_plutot_qu_avec_une_fausse(
    tmp_path,
):
    """Une référence inventée irait telle quelle dans un message de contrôle."""
    plan = lire_plan_be(
        _dxf_cadastre(tmp_path / "anonyme.dxf", section=None, numeros=("58", ""))
    )
    assert [p.reference for p in plan.parcelles] == ["58", None]


def test_les_parcelles_sont_enumerees_dans_l_ordre_des_numeros():
    """Trié en texte, « ZA 10 » passerait avant « ZA 9 »."""
    from shapely.geometry import box

    from dp_socle.import_be import ParcelleBE, _enumerer_parcelles

    carre = box(0, 0, 1, 1)
    parcelles = [
        ParcelleBE(geometrie=carre, reference=f"ZA {n}") for n in (10, 9, 100, 2)
    ]
    assert _enumerer_parcelles(parcelles) == "les parcelles ZA 2, ZA 9, ZA 10 et ZA 100"


def test_une_longue_liste_de_parcelles_est_tronquee_et_comptee():
    """Une énumération de quarante références ne se lit pas."""
    from shapely.geometry import box

    from dp_socle.import_be import (
        PARCELLES_ENUMEREES_MAX,
        ParcelleBE,
        _enumerer_parcelles,
    )

    carre = box(0, 0, 1, 1)
    nombre = PARCELLES_ENUMEREES_MAX + 4
    parcelles = [
        ParcelleBE(geometrie=carre, reference=f"ZA {n}") for n in range(1, nombre + 1)
    ]
    phrase = _enumerer_parcelles(parcelles)
    assert phrase.startswith(f"{nombre} parcelles, dont ZA 1,")
    assert phrase.endswith("et 4 autres")


# ---------------------------------------------------------------------------
# Un contour replié sur lui-même (Saint-Aubin-sur-Loire, 01/10/2026)
# ---------------------------------------------------------------------------


def test_un_contour_sans_surface_ne_fait_pas_tomber_l_emprise_cloturee():
    """La clôture se mesure sur l'anneau qui entoure quelque chose, et le dit.

    Mesuré sur le plan de Saint-Aubin-sur-Loire : un tracé de clôture replié
    sur lui-même n'entoure aucune surface, et l'union des contours tombait
    dessus — « point array must contain 0 or >1 elements », sur le serveur
    seulement, la suite restant verte sous Windows. Le chef de projet n'avait
    ni emprise ni cause.
    """
    from shapely.geometry import LineString, MultiLineString

    from dp_socle.import_be import EntiteBE, PlanBE

    # Les deux anneaux arrivent dans la même entité, comme le plan les donne :
    # une clôture en plusieurs tracés.
    carre = LineString([(0, 0), (100, 0), (100, 100), (0, 100), (0, 0)])
    replie = LineString([(200, 200), (210, 200), (205, 200), (200, 200)])
    plan = PlanBE(
        entites=[
            EntiteBE("cloture", "CLOTURE", MultiLineString([carre, replie]), False),
        ],
        azimut_tables_deg=0.0,
        correspondance={"CLOTURE": "cloture"},
        unite="m",
        facteur_unite=1.0,
        calques_ignores=[],
        calques_vides=[],
        source="essai.dxf",
    )

    assert plan.surface_cloturee_m2 == pytest.approx(10_000.0)
    assert any("n'entourent aucune surface" in a for a in plan.avertissements)
    # Une seconde lecture ne redit pas la même chose.
    plan.polygone_cloture
    assert sum("n'entourent aucune surface" in a for a in plan.avertissements) == 1


def test_un_contour_qui_se_recoupe_est_mesure_et_signale():
    """Un tracé de clôture en nœud papillon se mesure, au lieu de tout arrêter.

    Mesuré sur le plan de Saint-Aubin-sur-Loire : l'union des contours tombait
    sur `TopologyException: side location conflict`, et le chef de projet
    n'avait ni dossier ni cause. GEOS découpe le nœud en anneaux : la surface
    est conservée, et le tracé reste à reprendre au plan.
    """
    from shapely.geometry import LineString

    from dp_socle.import_be import EntiteBE, PlanBE

    noeud = LineString([(0, 0), (100, 100), (100, 0), (0, 100), (0, 0)])
    plan = PlanBE(
        entites=[EntiteBE("cloture", "CLOTURE", noeud, False)],
        azimut_tables_deg=0.0,
        correspondance={"CLOTURE": "cloture"},
        unite="m",
        facteur_unite=1.0,
        calques_ignores=[],
        calques_vides=[],
        source="essai.dxf",
    )

    # Les deux lobes du nœud : 2 500 m² chacun.
    assert plan.surface_cloturee_m2 == pytest.approx(5_000.0)
    assert any("se recoupent" in a for a in plan.avertissements)


def test_le_controle_de_l_emprise_ne_parle_pas_du_dxf_a_qui_depose_un_pdf():
    """Le seul contrôle partagé par les deux lots nomme la bonne source.

    Le 2bis lit un DXF, le 2ter un plan PDF, et ce contrôle est le seul qu'ils
    partagent. Il annonçait « aucun contour de clôture dans le DXF » dans les
    deux cas — à qui n'avait déposé aucun DXF, et à qui voyait sa clôture sur
    son plan. Relevé le 01/10/2026 sur Saint-Aubin-sur-Loire, dont la clôture
    est bien là mais ne se referme pas.
    """
    from shapely.geometry import LineString, Polygon

    from dp_socle.import_be import AVERTISSEMENT, EntiteBE, _emprise

    def plan_avec(*entites):
        from dp_socle.import_be import PlanBE

        return PlanBE(
            entites=list(entites),
            azimut_tables_deg=0.0,
            correspondance={"CLOTURE": "cloture"},
            unite="m",
            facteur_unite=1.0,
            calques_ignores=[],
            calques_vides=[],
            source="Plan du projet.pdf sur HelioScope Export.zip",
        )

    # Rien du tout : le message ne nomme aucun format de fichier.
    sans_rien = _emprise(plan_avec(), None)
    assert sans_rien.statut == AVERTISSEMENT
    assert "DXF" not in sans_rien.message
    assert "Aucun contour de clôture dans le plan importé" in sans_rien.message

    # Une clôture qui ne se referme pas : dire « aucun contour » envoyait
    # chercher un calque manquant au lieu du tracé ouvert.
    ouverte = _emprise(
        plan_avec(
            EntiteBE(
                "cloture",
                "CLOTURE",
                LineString([(0, 0), (100, 0), (100, 100), (0, 100)]),
                False,
            )
        ),
        None,
    )
    assert ouverte.statut == AVERTISSEMENT
    assert "DXF" not in ouverte.message
    assert "ne se referme pas" in ouverte.message

    # Et quand la clôture se referme, le contrôle se fait comme avant.
    fermee = _emprise(
        plan_avec(
            EntiteBE(
                "cloture",
                "CLOTURE",
                Polygon([(0, 0), (100, 0), (100, 100), (0, 100)]),
                False,
            )
        ),
        None,
    )
    assert "ne se referme pas" not in fermee.message


# ---------------------------------------------------------------------------
# Le contour que ses deux bouts manquent de peu (Bédarieux, 02/10/2026)
# ---------------------------------------------------------------------------


def test_un_contour_ferme_a_quelques_millimetres_pres_est_ferme():
    """Neuf millimètres sur 748 m ne font pas une clôture ouverte.

    Mesuré sur le DXF de Bédarieux du 24/09/2026 : la clôture est une polyligne
    de 747,6 m dont les deux bouts se manquent de 9 mm. Fermée pour qui la
    regarde, ouverte pour `is_closed` — et l'exigence stricte posée la veille
    la refusait, ce qui privait le chef de projet de sa surface clôturée **et**
    faisait tomber la carte.

    Un vrai trou de clôture se compte en mètres : un portail en fait sept.
    """
    from shapely.geometry import LineString, MultiLineString

    from dp_socle.import_be import EntiteBE, PlanBE

    def plan_avec(trace):
        return PlanBE(
            entites=[EntiteBE("cloture", "CLOTURE", trace, False)],
            azimut_tables_deg=0.0,
            correspondance={"CLOTURE": "cloture"},
            unite="m",
            facteur_unite=1.0,
            calques_ignores=[],
            calques_vides=[],
            source="essai.dxf",
        )

    # 9 mm : le même point, à la précision du dessin près. Rien à dire.
    presque = plan_avec(
        LineString([(0, 0), (100, 0), (100, 100), (0, 100), (0.009, 0)])
    )
    assert presque.surface_cloturee_m2 == pytest.approx(10_000.0, rel=1e-3)
    assert not [a for a in presque.avertissements if "refermé" in a]

    # 5 cm : refermé aussi, mais dit — le tracé mérite d'être repris.
    large = plan_avec(LineString([(0, 0), (100, 0), (100, 100), (0, 100), (0.05, 0)]))
    assert large.surface_cloturee_m2 == pytest.approx(10_000.0, rel=1e-3)
    assert any("refermé" in a for a in large.avertissements)

    # 1 m : c'est un trou, et la surface ne se calcule pas.
    trouee = plan_avec(LineString([(0, 0), (100, 0), (100, 100), (0, 100), (1.0, 0)]))
    assert trouee.surface_cloturee_m2 is None
    assert any("ne se referment pas" in a for a in trouee.avertissements)


def test_la_carte_se_cadre_meme_sans_contour_de_cloture():
    """Un plan sans clôture fermée garde une carte, cadrée sur le reste.

    Relevé sur Bédarieux : `bornes_wgs84(None)` faisait tomber l'application
    entière, et le chef de projet perdait jusqu'au rapport de contrôle qui lui
    disait ce qui manquait — c'est précisément quand un plan est incomplet
    qu'on a besoin de l'outil.
    """
    from shapely.geometry import LineString, Polygon

    from dp_socle.apercu_be import cadre_de_la_carte
    from dp_socle.import_be import EntiteBE, PlanBE

    def plan_avec(*entites):
        return PlanBE(
            entites=list(entites),
            azimut_tables_deg=0.0,
            correspondance={},
            unite="m",
            facteur_unite=1.0,
            calques_ignores=[],
            calques_vides=[],
            source="essai.dxf",
        )

    cloture = Polygon([(0, 0), (100, 0), (100, 100), (0, 100)])
    tables = Polygon([(10, 10), (40, 10), (40, 40), (10, 40)])

    # Avec clôture : c'est elle qui cadre.
    avec = plan_avec(
        EntiteBE("cloture", "CLOTURE", cloture, False),
        EntiteBE("tables_pv", "TABLES", tables, False),
    )
    assert cadre_de_la_carte(avec).bounds == cloture.bounds

    # Sans : les tables et la piste suffisent à cadrer.
    piste = LineString([(200, 0), (260, 0)])
    sans = plan_avec(
        EntiteBE("tables_pv", "TABLES", tables, False),
        EntiteBE("piste_lourde_a_creer", "PISTE", piste, False),
    )
    assert cadre_de_la_carte(sans).bounds == (10.0, 0.0, 260.0, 40.0)

    # Rien du tout : il n'y a pas de carte à cadrer, et c'est à l'appelant de
    # le dire plutôt que de lever.
    assert cadre_de_la_carte(plan_avec()) is None


# ---------------------------------------------------------------------------
# Les remplissages dont le contour est fait d'arcs (Bédarieux, 02/10/2026)
# ---------------------------------------------------------------------------


def test_un_contour_de_remplissage_en_arcs_se_lit_au_lieu_d_etre_perdu():
    """PVcase dessine ses voiries avec des raccords courbes, et on les perdait.

    Mesuré sur le DXF de Bédarieux : le calque « PVcase Road » porte sept
    remplissages dont six ont des bords en arcs. Leur `vertices` est vide — un
    chemin fait d'arêtes ne porte pas de sommets —, ils étaient déclarés
    illisibles, et **3 595 m² de voirie sur 3 851** disparaissaient. Le chef de
    projet ne voyait pas ses pistes et croyait le calque non apparié.

    `ezdxf` sait aplatir ces bords à la flèche près, comme le module le fait
    déjà pour les arcs des portails et les bulges des plateformes.
    """
    import math

    from shapely.geometry import Polygon

    from dp_socle.import_be import _sommets_du_contour

    document = ezdxf.new()
    hachure = document.modelspace().add_hatch()
    chemin = hachure.paths.add_edge_path()
    # Un carré de 10 m dont le côté droit est un demi-disque de 5 m de rayon.
    chemin.add_line((0.0, 0.0), (10.0, 0.0))
    chemin.add_arc(center=(10.0, 5.0), radius=5.0, start_angle=-90.0, end_angle=90.0)
    chemin.add_line((10.0, 10.0), (0.0, 10.0))
    chemin.add_line((0.0, 10.0), (0.0, 0.0))

    # Le chemin ne porte aucun sommet : c'est ce qui le faisait déclarer perdu.
    assert not list(getattr(chemin, "vertices", []))

    sommets = _sommets_du_contour(chemin, 1.0)
    assert len(sommets) > 20, "l'arc doit être discrétisé, pas sauté"
    # 100 m² pour le carré, plus le demi-disque de 5 m : 139,3 m².
    assert Polygon(sommets).area == pytest.approx(100.0 + math.pi * 25.0 / 2.0, rel=1e-3)


def test_un_contour_que_ezdxf_ne_sait_pas_aplatir_reste_illisible_sans_lever():
    """Ce qu'on ne sait pas lire se compte, et se dit — il ne remonte pas.

    L'appelant tient le compte des remplissages illisibles et le porte au
    rapport. Une exception à cet endroit emporterait l'import entier pour un
    seul bord récalcitrant.
    """
    from dp_socle.import_be import _sommets_du_contour

    class _CheminCasse:
        vertices = []

    assert _sommets_du_contour(_CheminCasse(), 1.0) == []


# ---------------------------------------------------------------------------
# Les cotes normalisées des ouvrages dessinés
# ---------------------------------------------------------------------------


def test_le_tableau_de_reference_cote_tous_les_ouvrages_du_plan(plan, tableau):
    """Sur les fichiers de référence, rien ne manque — le contrôle se tait."""
    from dp_socle.import_be import _cotes_des_ouvrages

    controle = _cotes_des_ouvrages(plan, tableau)
    assert controle.statut == "ok", controle.message
    assert controle.valeur_dxf == controle.valeur_tableau


def test_un_ouvrage_dessine_sans_cote_part_en_demande_au_bureau_d_etudes(plan):
    """Le plan porte une citerne, le tableau ne la cote pas : il faut le dire.

    Mesuré le 06/10/2026 sur Auzainvilliers, dont le tableau bilan ne portait
    que 7 cotes — les quatre postes et les trois locaux — là où ceux de Sarnois
    et de Bray-Saint-Aignan en portent 17. Six catégories dessinées au DXF
    n'avaient aucune cote, dont la citerne incendie. Rien ne l'a dit à
    l'import : le chef de projet a produit le dossier entier avant que le
    rapport n'annonce les ouvrages non dessinés, et c'est la citerne absente du
    plan de coupe qui a mis sur la voie.

    Le message doit porter le marqueur de demande au bureau d'études : c'est lui
    qui le sort de la liste des remarques pour le poser dans le bandeau rouge
    que le chef de projet recopie dans son message.
    """
    from dp_socle.import_be import _cotes_des_ouvrages

    class TableauSansCiterne:
        """Le tableau de référence privé de ses lignes « Citerne incendie »."""

        cotes = [
            cote
            for cote in lire_tableau(TABLEAU, "IND06").cotes
            if not cote.ouvrage.startswith("Citerne incendie")
        ]

    controle = _cotes_des_ouvrages(plan, TableauSansCiterne())
    assert controle.statut == "avertissement"
    assert controle.valeur_tableau == controle.valeur_dxf - 1
    assert "bache_incendie" in controle.message
    assert "à demander au bureau d'études" in controle.message
    # Le constat énumère les quatre libellés cherchés — c'est ce qui permet de
    # retrouver la ligne dans le tableau. La **demande**, elle, nomme la
    # famille : réclamer « Citerne incendie — 30 » parce qu'elle vient en tête
    # désignerait la mauvaise sur un plan qui porte une 120.
    assert "Citerne incendie — 30" in controle.message
    _, _, demande = controle.message.partition("à demander au bureau d'études")
    assert "Citerne incendie (la variante du projet)" in demande
    assert "Citerne incendie — 30" not in demande


def test_la_demande_nomme_la_famille_et_non_sa_premiere_variante():
    """Une famille à variantes se demande par son nom commun, sans reliquat.

    « Local de stockage matériel — P » est le plus long préfixe commun aux trois
    tailles, le « P » de « P<=5MWc » compris : couper au tiret cadratin est ce
    qui rend une demande lisible.
    """
    from dp_socle.import_be import _famille_de_cote

    assert _famille_de_cote(("Aire d'aspiration",)) == "Aire d'aspiration"
    assert _famille_de_cote(
        ("Local de stockage matériel — P<=5MWc", "Local de stockage matériel — P>5MWc")
    ) == "Local de stockage matériel (la variante du projet)"


def test_le_controle_des_cotes_est_rendu_a_chaque_import(plan, tableau):
    """Il est dans la liste que l'interface affiche, pas seulement appelable."""
    from dp_socle.import_be import controler

    libelles = [c.libelle for c in controler(plan, tableau)]
    assert "Cotes normalisées des ouvrages" in libelles


# ---------------------------------------------------------------------------
# Ce que le plan pose de travers : ouvrages à cheval, plateformes détachées
# ---------------------------------------------------------------------------


def _plan_synthetique(objets):
    """Un `PlanBE` minimal porté par une liste de (catégorie, calque, géométrie)."""
    from dp_socle.import_be import EntiteBE, PlanBE

    return PlanBE(
        entites=[
            EntiteBE(categorie=c, calque=calque, geometrie=g, z_reel=False)
            for c, calque, g in objets
        ],
        azimut_tables_deg=0.0,
        correspondance={calque: c for c, calque, _g in objets},
        unite="m",
        facteur_unite=1.0,
        calques_ignores=[],
        calques_vides=[],
        source="synthétique",
    )


def _enceinte(x0, y0, x1, y1):
    """La clôture en polyligne fermée, comme le DXF la livre.

    Une `LinearRing` ne conviendrait pas : `_polygoniser` ne traite que les
    polylignes et les polygones, et l'anneau rendait une enceinte vide — le
    contrôle annonçait alors qu'il n'avait rien pu vérifier, ce qui est le bon
    comportement pour une mauvaise raison.
    """
    from shapely.geometry import LineString, box

    contour = LineString(list(box(x0, y0, x1, y1).exterior.coords))
    return ("cloture", "UNI_Clôture", contour)


def test_le_plan_de_reference_ne_declenche_aucun_des_deux_constats(plan):
    """Saint-Cyr est le mètre étalon : un faux positif s'y verrait tout de suite.

    Son aire d'aspiration est **entièrement hors clôture**, et c'est sa place —
    c'est le point d'eau des pompiers, posé là où leur engin se gare. La
    signaler aurait condamné le contrôle dès le premier dossier.
    """
    from dp_socle.import_be import _ouvrages_dans_l_enceinte, _plateformes_raccordees

    assert _ouvrages_dans_l_enceinte(plan).statut == "ok"
    assert _plateformes_raccordees(plan).statut == "ok"
    aspiration = plan.geometries("aire_aspiration")
    assert aspiration, "le plan de référence porte bien une aire d'aspiration"
    assert aspiration[0].difference(plan.polygone_cloture).area > 1.0


def test_un_ouvrage_a_cheval_sur_la_cloture_est_dit_et_groupe_par_calque():
    """Mesuré le 06/10/2026 sur Auzainvilliers, dont le PDL déborde de 38 %.

    Le calque du poste y porte six objets — le poste, sa plateforme et les
    bandes de terre autour —, tous à cheval. Six lignes pour un seul poste mal
    posé se lisent mal, et le bureau d'études corrige un calque, pas un objet :
    le constat groupe.
    """
    from shapely.geometry import box
    from dp_socle.import_be import _ouvrages_dans_l_enceinte

    plan = _plan_synthetique(
        [
            _enceinte(0, 0, 100, 100),
            # Le poste à cheval : 12 x 3 m, dont la moitié dehors.
            ("pdl_ptr", "UNI_PDL-PTR", box(94.0, 40.0, 106.0, 43.0)),
            # Sa plateforme, sur le même calque et à cheval elle aussi.
            ("pdl_ptr", "UNI_PDL-PTR", box(94.0, 44.0, 106.0, 48.0)),
            # Un poste bien posé, dedans : il ne doit pas être compté.
            ("ptr", "UNI_PDT", box(10.0, 10.0, 20.0, 13.0)),
        ]
    )
    controle = _ouvrages_dans_l_enceinte(plan)
    assert controle.statut == "avertissement"
    assert controle.valeur_dxf == 2, "les deux objets du calque sont comptés"
    assert controle.message.count("UNI_PDL-PTR") == 2, "une seule ligne de constat"
    assert "UNI_PDT" not in controle.message
    assert "à demander au bureau d'études" in controle.message


def test_un_ouvrage_qui_affleure_la_cloture_ne_declenche_rien():
    """Un pouillème de débord est de l'imprécision de tracé entre deux calques.

    Le seuil est celui de la clôture dans l'emprise cadastrale, pour la même
    raison : les cas réels se comptent en dizaines de pour cent.
    """
    from shapely.geometry import box
    from dp_socle.import_be import DEBORD_OUVRAGE_M2, _ouvrages_dans_l_enceinte

    plan = _plan_synthetique(
        [
            _enceinte(0, 0, 100, 100),
            # 3 m de long, débordant de 10 cm : 0,3 m², sous le seuil de 1 m².
            ("ptr", "UNI_PDT", box(99.9, 40.0, 110.0, 43.0).intersection(
                box(0, 0, 100.1, 100)
            )),
        ]
    )
    assert DEBORD_OUVRAGE_M2 == 1.0
    assert _ouvrages_dans_l_enceinte(plan).statut == "ok"


def test_une_plateforme_detachee_de_la_voirie_est_dite_avec_son_ecart():
    """Mesuré le 06/10/2026 : 1,05 m entre la plateforme du PTR et la piste.

    L'écart va au message : un mètre est un raccord manqué, trente mètres un
    accès absent, et c'est au chef de projet d'en juger.
    """
    from shapely.geometry import box
    from dp_socle.import_be import _plateformes_raccordees

    plan = _plan_synthetique(
        [
            ("piste_lourde_a_creer", "UNI_VRD_Piste", box(0.0, 0.0, 100.0, 4.0)),
            # Raccordée : elle touche la piste.
            ("plateforme", "UNI_VRD_Plateforme", box(10.0, 4.0, 20.0, 14.0)),
            # Détachée de 1,05 m, comme à Auzainvilliers.
            ("plateforme", "UNI_VRD_Plateforme", box(40.0, 5.05, 50.0, 15.0)),
        ]
    )
    controle = _plateformes_raccordees(plan)
    assert controle.statut == "avertissement"
    assert controle.valeur_dxf == 1
    assert "1.05 m" in controle.message
    assert "UNI_VRD_Plateforme" in controle.message
    assert "à demander au bureau d'études" in controle.message


def test_sans_voirie_dessinee_le_raccord_est_dit_non_controle():
    """Pas de repli muet : ce qui n'a pas pu être vérifié se dit."""
    from shapely.geometry import box
    from dp_socle.import_be import _plateformes_raccordees

    plan = _plan_synthetique(
        [("plateforme", "UNI_VRD_Plateforme", box(10.0, 4.0, 20.0, 14.0))]
    )
    controle = _plateformes_raccordees(plan)
    assert controle.statut == "avertissement"
    assert "n'a pas pu être contrôlé" in controle.message


def test_sans_cloture_fermee_le_debord_est_dit_non_controle():
    """Même règle : une clôture ouverte ne vaut pas un plan sans débord."""
    from shapely.geometry import box
    from dp_socle.import_be import _ouvrages_dans_l_enceinte

    plan = _plan_synthetique([("ptr", "UNI_PDT", box(10.0, 10.0, 20.0, 13.0))])
    controle = _ouvrages_dans_l_enceinte(plan)
    assert controle.statut == "avertissement"
    assert "n'aurait pas été vu" in controle.message


def test_les_deux_constats_sont_rendus_a_chaque_import(plan, tableau):
    """Ils sont dans la liste que l'interface affiche, pas seulement appelables."""
    from dp_socle.import_be import controler

    libelles = [c.libelle for c in controler(plan, tableau)]
    assert "Ouvrages dans l'enceinte" in libelles
    assert "Raccordement des plateformes" in libelles


# ---------------------------------------------------------------------------
# Les cotes complétées depuis le classeur des gabarits UNITe
# ---------------------------------------------------------------------------


def test_un_tableau_complet_n_emprunte_rien_au_classeur(tableau):
    """Saint-Cyr porte ses 17 cotes : rien à compléter, rien à annoncer."""
    assert tableau.cotes_completees == ()
    assert not [m for m in tableau.avertissements if "gabarits UNITe" in m]


def test_un_onglet_tronque_se_complete_et_le_dit(tmp_path):
    """Un tableau bâti sur un modèle ancien ne porte pas tout le catalogue.

    Mesuré le 06/10/2026 sur Auzainvilliers, dont l'onglet des dimensions
    s'arrête après « Local de stockage matériel » : il manquait les sections
    « Aire de charge (BESS) » et « Citerne incendie » et la ligne de l'aire
    d'aspiration, soit six ouvrages dessinés au plan et absents du dossier —
    dont la citerne incendie, que le chef de projet cherchait sur la coupe.

    L'onglet n'est pas une donnée de projet : c'est le catalogue des gabarits
    UNITe, le même dans tous les tableaux, et l'outil en livre une copie. Les
    sections absentes s'y prennent donc, et le complément s'annonce.
    """
    from dp_socle.tableau_bilan import lire_tableau

    chemin = _tableau_fabrique(
        tmp_path / "tronque.xlsx", decoupees=True, onglet_cotes="Dimensions postes"
    )
    tableau = lire_tableau(chemin, "IND03")

    ouvrages = {c.ouvrage for c in tableau.cotes}
    assert "Citerne incendie — 120" in ouvrages
    assert "Aire de charge (BESS) — Citerne de refroidissement" in ouvrages
    assert "Citerne incendie — 120" in tableau.cotes_completees
    annonces = [m for m in tableau.avertissements if "gabarits UNITe" in m]
    assert len(annonces) == 1, annonces
    assert "Citerne incendie — 120" in annonces[0]


def test_le_tableau_garde_la_main_sur_ce_qu_il_porte(tmp_path):
    """Seuls les libellés absents sont ajoutés : le projet prime sur le standard.

    Un tableau qui déclare un poste hors cotes de catalogue garde le sien — le
    complément n'est pas un alignement sur le standard, c'est un bouche-trou.
    """
    from dp_socle.tableau_bilan import lire_tableau

    tableau = lire_tableau(_tableau_fabrique(
            tmp_path / "propre.xlsx", decoupees=True,
            onglet_cotes="Dimensions postes",
        ), "IND03")

    ptr = [c for c in tableau.cotes if c.ouvrage == "Poste de transformation - PTR"]
    assert len(ptr) == 1, "pas de doublon entre le tableau et le classeur"
    assert ptr[0].dimensions == "10 x 3 x 3m"
    assert "Poste de transformation - PTR" not in tableau.cotes_completees


def test_la_variante_de_citerne_reste_choisie_sur_l_emprise_mesuree(tmp_path):
    """Rien à saisir : c'est le plan qui dit laquelle des quatre.

    La citerne d'Auzainvilliers mesure 103,9 m² au DXF, et tombe sur la 120
    (11,7 x 9,3 m, soit 108,8 m², à 4,5 % près). Demander le volume au chef de
    projet lui ferait ressaisir ce que le plan donne déjà au centimètre.
    """
    from dp_socle.contrat import Contrat
    from dp_socle.tableau_bilan import lire_tableau

    tableau = lire_tableau(_tableau_fabrique(
            tmp_path / "variante.xlsx", decoupees=True,
            onglet_cotes="Dimensions postes",
        ), "IND03")
    contrat = Contrat(
        dossier=tmp_path,
        donnees={
            "cotes_normalisees": [
                {
                    "ouvrage": c.ouvrage,
                    "dimensions": c.dimensions,
                    "ordre_cotes": c.ordre_cotes,
                    "surface_m2": c.surface_m2,
                    "surface_plateforme_m2": c.surface_plateforme_m2,
                }
                for c in tableau.cotes
            ]
        },
        couches={},
    )
    assert contrat.cote_de_categorie("bache_incendie", surface_m2=103.9).ouvrage == (
        "Citerne incendie — 120"
    )
    assert contrat.cote_de_categorie("bache_incendie", surface_m2=35.0).ouvrage == (
        "Citerne incendie — 30"
    )


def test_la_provenance_des_cotes_redescend_dans_le_contrat(plan, tmp_path):
    """Sur un dossier déjà déposé, on doit pouvoir dire d'où vient une hauteur."""
    from dp_socle.import_be import parametres_json
    from dp_socle.tableau_bilan import lire_tableau

    tableau = lire_tableau(_tableau_fabrique(
            tmp_path / "trace.xlsx", decoupees=True,
            onglet_cotes="Dimensions postes",
        ), "IND03")
    donnees = parametres_json(plan, tableau, [])
    assert "Citerne incendie — 120" in donnees["cotes_completees"]
    assert len(donnees["cotes_normalisees"]) == len(tableau.cotes)


def test_un_seul_chemin_pour_le_classeur_des_gabarits():
    """Deux définitions du même chemin finiraient par désigner deux fichiers."""
    from dp_socle.plan_pdf import CHEMIN_GABARITS as cote_plan_pdf
    from dp_socle.tableau_bilan import CHEMIN_GABARITS as cote_tableau

    assert cote_plan_pdf is cote_tableau
    assert cote_tableau.exists()
