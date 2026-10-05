"""La sortie PowerPoint du lot 8, mesurée dans le fichier produit.

« Tester le résultat, pas l'exécution » : le critère n'est pas que le `.pptx`
s'écrive, mais qu'il porte les bonnes géométries. Le fichier est donc enregistré,
relu avec `python-pptx`, et mesuré en EMU convertis en millimètres — la liste que
le brief du lot 8 fixe :

- la taille de diapo, 420 x 297 mm ;
- le nombre de diapos et leur ordre ;
- la position et le rapport de chaque cadre photo, contre `disposition` ;
- un groupe repère + cône par vue attendue, à 2,4 et 8 mm ;
- le dénominateur du cartouche de chaque diapo de repérage, contre le cadrage ;
- l'égalité des numéros de page des diapos alternatives, et le sommaire qui n'en
  compte qu'une.

Rien de tout cela ne prouve que PowerPoint l'ouvre correctement : c'est l'objet de
la liste de contrôles à l'écran que le brief tient, et qui se refait à chaque
changement du montage.

AUCUN SERVICE EN LIGNE
-----------------------
Les trois points d'entrée de la Géoplateforme sont remplacés par des doublures :
ce qui est mesuré ici est le montage, pas le service — `test_ign.py` s'occupe du
service. C'est ce qui permet de produire un dossier entier en quelques secondes,
là où la vraie génération demande deux à trois minutes de téléchargements.
"""

from __future__ import annotations

import zipfile

import pytest
from PIL import Image
from shapely.geometry import box

from dp_socle import sortie_pptx
from dp_socle.ign import FondRaster, Parcelle
from dp_socle.planche import HAUTEUR_MM, LARGEUR_MM, Planche
from dp_socle.planches.photographies import disposition
from dp_socle.planches.reperage_vues import RAYON_CONE_MM, RAYON_REPERE_MM
from dp_socle.projet import Projet

from . import contrat_synthetique as synthese

pytest.importorskip(
    "cairosvg", reason="CairoSVG et libcairo sont nécessaires pour rendre les planches"
)

EMU_PAR_MM = 36000.0


def _mm(emu) -> float:
    return round(emu / EMU_PAR_MM, 2)


@pytest.fixture(scope="module")
def sans_geoplateforme():
    """Remplace les trois appels IGN par des doublures, sans rien masquer d'autre.

    La doublure rend une image à la bbox **exacte** demandée : c'est la garantie
    que `Planche.ajouter_fond_raster` exige, et une doublure qui la violerait
    fausserait les mesures d'échelle au lieu de les éprouver.

    Portée module, comme le dossier produit : la substitution doit tenir pendant
    la génération, qui n'a lieu qu'une fois.

    Les planches écrivent `from ..ign import telecharger_fond` : le nom est lié
    à l'import, et remplacer le seul attribut du module `ign` ne les atteint
    pas. Mesuré le 29/09/2026 — la doublure n'était jamais appelée, ces tests
    interrogeaient la Géoplateforme à chaque passage sans porter la marque
    `reseau`, et la CI est tombée deux fois de suite le 28/09 pendant que le
    poste restait vert. La substitution suit donc chaque module qui tient une
    référence, et le compteur rendu ici sert à vérifier qu'elle a servi.
    """
    import sys
    from collections import Counter

    from dp_socle import ign

    appels = Counter()

    def fond(couche, bbox, largeur_mm, hauteur_mm, dpi=200, format_image="image/jpeg",
             timeout=120):
        appels["telecharger_fond"] += 1
        return FondRaster(
            image=Image.new("RGB", (240, 170), (232, 230, 226)),
            bbox=tuple(bbox), couche=couche, dpi=dpi, format=format_image,
        )

    def parcelles(bbox, *args, **kwargs):
        """Un damier de parcelles couvrant la fenêtre demandée.

        Pas une liste vide : le parcellaire est le fond des plans de repérage
        depuis le 26/09/2026, et c'est lui qui décide du poids d'une planche
        photographique, donc de la voie de rendu retenue. Un fond vide aurait
        fait passer ces planches en vectoriel sans rien prouver.
        """
        appels["telecharger_parcelles"] += 1
        minx, miny, maxx, maxy = bbox
        cote = 60.0
        return [
            Parcelle(
                geometrie=box(
                    minx + rang_x * cote, miny + rang_y * cote,
                    minx + (rang_x + 1) * cote, miny + (rang_y + 1) * cote,
                ),
                idu=f"{rang_x:03d}{rang_y:03d}", section="AB",
                numero=f"{rang_x}-{rang_y}", contenance_m2=cote * cote,
                commune="00000",
            )
            for rang_x in range(int((maxx - minx) // cote) + 1)
            for rang_y in range(int((maxy - miny) // cote) + 1)
        ]

    def batiments(*args, **kwargs):
        appels["telecharger_batiments"] += 1
        return []

    doublures = {
        "telecharger_fond": fond,
        "telecharger_parcelles": parcelles,
        "telecharger_batiments": batiments,
    }
    originaux = {nom: getattr(ign, nom) for nom in doublures}

    with pytest.MonkeyPatch.context() as substitution:
        for module in list(sys.modules.values()):
            if not getattr(module, "__name__", "").startswith("dp_socle"):
                continue
            for nom, doublure in doublures.items():
                # `is` et non `hasattr` : on ne remplace que la vraie fonction,
                # jamais une doublure déjà posée ni un homonyme.
                if getattr(module, nom, None) is originaux[nom]:
                    substitution.setattr(module, nom, doublure)
        yield appels


@pytest.fixture(scope="module")
def dossier_projet(tmp_path_factory):
    """Un projet complet : emprise cadastrale, contrat du bureau d'études, notice.

    L'emprise cadastrale est élargie de 60 m autour du site, comme dans les tests
    du lot 4 : c'est elle que cadre DP 2, et la clôture du contrat est ce que
    cadrent les plans de repérage.
    """
    import geopandas as gpd

    tmp_path = tmp_path_factory.mktemp("lot8")
    dossier = tmp_path / "sortie" / "PV-TÉMOIN"
    synthese.ecrire(dossier)

    x0, y0 = synthese.ORIGINE_L93
    emprise = tmp_path / "emprise.geojson"
    gpd.GeoDataFrame(
        geometry=[
            box(
                x0 - 60, y0 - 60,
                x0 + synthese.LARGEUR_SITE_M + 60,
                y0 + synthese.HAUTEUR_SITE_M + 60,
            )
        ],
        crs="EPSG:2154",
    ).to_file(emprise, driver="GeoJSON")

    projet = Projet(
        nom="PV-TÉMOIN",
        commune="Témoin",
        code_postal="45000",
        date="2026-09-26",
        emprise=str(emprise),
        notice=str(_notice(tmp_path / "notice.pdf")),
    )
    return projet, tmp_path / "sortie"


def _notice(chemin):
    """Une notice de deux pages, pour mesurer la pagination d'une pièce épaisse."""
    from pypdf import PdfWriter

    ecrivain = PdfWriter()
    for _ in range(2):
        ecrivain.add_blank_page(width=595, height=842)
    with open(chemin, "wb") as fichier:
        ecrivain.write(fichier)
    return chemin


@pytest.fixture(scope="module")
def rapport(dossier_projet, sans_geoplateforme):
    """Le dossier produit **une fois** pour tout le fichier.

    Une génération complète compose une quinzaine de planches et les rend deux
    fois chacune, pour comparer les deux voies : la refaire à chaque test coûtait
    plus de dix minutes. Les tests ne font que mesurer le fichier, ils ne le
    modifient pas — le partager est sans risque.
    """
    projet, sortie = dossier_projet
    return sortie_pptx.generer_pptx(projet, dossier_sortie=sortie)


@pytest.fixture(scope="module")
def presentation(rapport):
    from pptx import Presentation

    return Presentation(str(rapport.fichier))


# ---------------------------------------------------------------------------
# Le fichier, et ce qu'il porte
# ---------------------------------------------------------------------------


def test_les_doublures_ont_bien_remplace_la_geoplateforme(rapport, sans_geoplateforme):
    """Ce fichier mesure le montage, pas le service : il ne demande rien dehors.

    Sans ce contrôle, une doublure qui n'atteint plus les planches les laisse
    repartir vers la Géoplateforme, et tout reste vert — jusqu'au jour où le
    service tousse. C'est ce qui est arrivé le 28/09/2026 : deux CI rouges
    d'affilée, vertes le lendemain, sans qu'une ligne du dépôt ait changé.
    """
    assert sans_geoplateforme["telecharger_fond"] >= 2, sans_geoplateforme
    assert sans_geoplateforme["telecharger_parcelles"] >= 1, sans_geoplateforme


def test_le_fichier_se_nomme_pour_ce_qu_il_est(rapport):
    """Un dossier à finir, pas un dossier déposable (décision D9)."""
    assert rapport.fichier.name == "PV-TÉMOIN_DP_a_finaliser.pptx"
    assert rapport.fichier.exists()


def test_la_generation_ne_laisse_que_le_pptx_dans_le_dossier_du_projet(rapport):
    """Les PDF intermédiaires ne doivent pas remplacer ceux de la voie PDF.

    Les modules de pièce écrivent toujours un PDF, et ceux d'ici portent la
    pagination du PPTX — DP 11 y tombe plus loin, puisque les pièces
    photographiques y figurent. Les laisser dans le dossier du projet mélangerait
    deux dossiers vrais en un dossier faux.
    """
    ecrits = {chemin.name for chemin in rapport.dossier.iterdir()}

    assert rapport.fichier.name in ecrits
    assert not [nom for nom in ecrits if nom.lower().endswith(".pdf")]


def test_la_diapo_mesure_une_a3_paysage(presentation):
    """420 x 297 mm : c'est ce qui rend l'échelle du cartouche vraie."""
    assert (_mm(presentation.slide_width), _mm(presentation.slide_height)) == (
        LARGEUR_MM, HAUTEUR_MM,
    )


def test_les_pieces_se_suivent_dans_l_ordre_du_dossier(rapport):
    """L'ordre du sommaire, celui que les services instructeurs lisent."""
    codes = [diapo.code for diapo in rapport.diapos]

    attendus = [
        "", "DP 1-1", "DP 1-2", "DP 1-3", "DP 2", "DP 3",
        "DP 4-1", "DP 4-2",
        "DP 6", "DP 6", "DP 7", "DP 8", "DP 8", "DP 8",
        "DP 11", "DP 11",
    ]
    assert codes == attendus


def test_chaque_diapo_a_sa_mise_en_page_et_son_fond_hors_d_atteinte(presentation):
    """Le fond est sur la mise en page : le chef de projet ne peut pas le bouger.

    C'est ce qui protège l'échelle. Une planche redimensionnée rendrait faux le
    dénominateur écrit à son cartouche, et rien ne le signalerait.
    """
    noms = set()
    for diapositive in presentation.slides:
        mise_en_page = diapositive.slide_layout
        # Une mise en page par diapo, et aucune partagée : un fond partagé
        # afficherait la mauvaise planche.
        assert mise_en_page.name not in noms
        noms.add(mise_en_page.name)
        fonds = [
            forme for forme in mise_en_page.shapes
            if forme.element.tag.endswith("}pic")
        ]
        assert len(fonds) == 1
        assert (_mm(fonds[0].width), _mm(fonds[0].height)) == (LARGEUR_MM, HAUTEUR_MM)
        assert not any(
            forme.element.tag.endswith("}pic") for forme in diapositive.shapes
        )


# ---------------------------------------------------------------------------
# Les cadres photo
# ---------------------------------------------------------------------------


def _diapos_de(presentation, prefixe: str):
    return [
        d for d in presentation.slides if d.slide_layout.name.startswith(prefixe)
    ]


def _planche_temoin() -> Planche:
    return Planche(titre="T", numero="1", projet="PV-TÉMOIN", date="26/09/2026")


@pytest.mark.parametrize(
    "prefixe, nombre_cadres",
    [("DP 6 — 2 cadres", 2), ("DP 6 — 3 cadres", 3), ("DP 7", 2), ("DP 8", 2)],
)
def test_les_cadres_photo_tombent_ou_la_planche_les_pose(
    presentation, prefixe, nombre_cadres
):
    """Les réservations d'image reprennent la géométrie de `disposition`.

    Même arithmétique que la planche du dossier PDF : c'est ce qui garantit qu'un
    dossier fini dans PowerPoint ressemble à un dossier fini par l'outil.
    """
    attendus = [
        pose.image for pose in disposition(_planche_temoin(), nombre_cadres).emplacements
    ]

    for diapositive in _diapos_de(presentation, prefixe):
        cadres = sorted(diapositive.placeholders, key=lambda f: f.top)
        assert len(cadres) == nombre_cadres
        for cadre, attendu in zip(cadres, attendus):
            mesure = (
                _mm(cadre.left), _mm(cadre.top), _mm(cadre.width), _mm(cadre.height),
            )
            assert mesure == pytest.approx(attendu, abs=0.05)


def test_la_page_de_garde_reserve_la_perspective(presentation):
    """Le cadre du photomontage est une réservation, comme les cadres photo.

    Le chef de projet insère sa perspective dans le fichier plutôt que de la
    déposer avant de générer. La planche garde dessous son cadre tireté et sa
    mention « Photomontage à insérer », que l'image posée recouvre.
    """
    from dp_socle.planches.page_garde import zone_image

    garde = presentation.slides[0]

    cadres = list(garde.placeholders)
    assert len(cadres) == 1
    mesure = (
        _mm(cadres[0].left), _mm(cadres[0].top),
        _mm(cadres[0].width), _mm(cadres[0].height),
    )
    assert mesure == pytest.approx(zone_image(), abs=0.05)


def test_les_cadres_sont_verrouilles_en_position(rapport):
    """Un cadre déplacé fausse la planche : le verrou le dit dans le fichier."""
    with zipfile.ZipFile(rapport.fichier) as paquet:
        diapos = [
            paquet.read(nom).decode("utf-8")
            for nom in paquet.namelist()
            if nom.startswith("ppt/slides/slide") and nom.endswith(".xml")
        ]

    avec_cadres = [xml for xml in diapos if "<p:ph " in xml]
    assert avec_cadres
    for xml in avec_cadres:
        assert 'noMove="1"' in xml and 'noResize="1"' in xml


def test_chaque_cadre_porte_son_intitule_editable(presentation):
    """Seul reste vivant de la décision D1 : les intitulés sont des zones de texte.

    Ils portent le repère de la vue, qui est tout ce qui force encore
    l'appariement entre une photographie et son cône.
    """
    diapositive = _diapos_de(presentation, "DP 7")[0]

    textes = [
        forme.text_frame.text
        for forme in diapositive.shapes
        if forme.has_text_frame and forme.text_frame.text
    ]
    assert "PC7-1 : Photographie de l'environnement proche" in textes
    assert "PC7-2 : Photographie de l'environnement proche" in textes


# ---------------------------------------------------------------------------
# Les repères de vue
# ---------------------------------------------------------------------------


def _groupes(diapositive):
    return [f for f in diapositive.shapes if f.element.tag.endswith("}grpSp")]


def _xml_d_une_diapo_a_reperes(fichier) -> str:
    """L'XML d'une diapo qui porte des groupes de repérage.

    C'est `<p:grpSp>` qu'il faut chercher, et non « grpSp » : tout `p:spTree`
    porte un `p:grpSpPr`, et une recherche laxiste trouvait la page de garde.
    """
    with zipfile.ZipFile(fichier) as paquet:
        for nom in sorted(paquet.namelist()):
            if not nom.startswith("ppt/slides/slide"):
                continue
            xml = paquet.read(nom).decode("utf-8")
            if "<p:grpSp>" in xml:
                return xml
    raise AssertionError("aucune diapo ne porte de groupe de repérage")


@pytest.mark.parametrize(
    "prefixe, reperes",
    [
        ("DP 6", ["Vue A"]),
        ("DP 7", ["PC7-1", "PC7-2"]),
        ("DP 8", ["PC8-1", "PC8-2"]),
    ],
)
def test_une_vue_attendue_egale_un_groupe_repere_et_cone(
    presentation, prefixe, reperes
):
    for diapositive in _diapos_de(presentation, prefixe):
        groupes = _groupes(diapositive)
        assert [g.name for g in groupes] == [f"Repère {r}" for r in reperes]


def test_le_groupe_tourne_autour_du_point_de_vue(presentation):
    """Le cadre du groupe fait 16 x 16 mm, centré sur le repère.

    Le pivot d'un groupe est le centre de son cadre : décalé, viser ferait
    glisser le repère hors du point de vue. C'est le `prstGeom` « pie » du cône
    qui tient ce cadre — il occupe la boîte du cercle entier, soit 2 x
    `RAYON_CONE_MM` de côté — et le disque comme l'étiquette tiennent dedans.
    """
    diapositive = _diapos_de(presentation, "DP 7")[0]

    for groupe in _groupes(diapositive):
        assert (_mm(groupe.width), _mm(groupe.height)) == (
            2 * RAYON_CONE_MM, 2 * RAYON_CONE_MM,
        )
        centre = (
            _mm(groupe.left) + RAYON_CONE_MM, _mm(groupe.top) + RAYON_CONE_MM,
        )
        enfants = list(groupe.shapes)
        disque = next(f for f in enfants if _mm(f.width) == 2 * RAYON_REPERE_MM)
        assert _mm(disque.left) + RAYON_REPERE_MM == pytest.approx(centre[0], abs=0.05)
        assert _mm(disque.top) + RAYON_REPERE_MM == pytest.approx(centre[1], abs=0.05)


def test_le_cone_est_un_secteur_de_50_degres_translucide(rapport):
    """50° et 55 % d'opacité, comme le marqueur de `photos-geoloc` et le PDF.

    Les poignées d'angle s'écrivent en 60 000e de degré : c'est la mesure qui l'a
    dit, un secteur réglé sans ce facteur sortait à 0,12°.
    """
    xml = _xml_d_une_diapo_a_reperes(rapport.fichier)

    assert 'prst="pie"' in xml
    assert '<a:gd name="adj1" fmla="val 14700000"/>' in xml
    assert '<a:gd name="adj2" fmla="val 17700000"/>' in xml
    assert (17700000 - 14700000) / 60000 == 50.0
    assert '<a:alpha val="55000"/>' in xml


def test_l_etiquette_du_repere_reste_horizontale(rapport):
    xml = _xml_d_une_diapo_a_reperes(rapport.fichier)

    assert 'upright="1"' in xml
    assert 'wrap="none"' in xml


# ---------------------------------------------------------------------------
# Les alternatives, et la pagination
# ---------------------------------------------------------------------------


def test_les_alternatives_partagent_le_numero_de_leur_piece(rapport):
    """Supprimer les surnuméraires laisse la pagination juste (décision D3)."""
    for code, attendu in (("DP 6", 2), ("DP 8", 3)):
        numeros = {d.numero for d in rapport.diapos if d.code == code}
        rangs = [d.alternative for d in rapport.diapos if d.code == code]
        assert len(numeros) == 1
        assert rangs == [(rang, attendu) for rang in range(1, attendu + 1)]


def test_le_sommaire_ne_compte_qu_une_diapo_par_piece(rapport):
    lignes = [(l["numero"], l["page"]) for l in rapport.sommaire]

    assert lignes == [
        ("—", 1), ("DP 1-1", 2), ("DP 1-2", 3), ("DP 1-3", 4), ("DP 2", 5),
        ("DP 3", 6), ("DP 4-1", 7), ("DP 4-2", 8), ("DP 6", 9), ("DP 7", 10),
        ("DP 8", 11), ("DP 11", 12),
    ]


def test_la_notice_occupe_autant_de_diapos_que_de_pages(rapport):
    """Seule pièce qui s'étend, et seule dont les numéros se suivent."""
    notice = [d for d in rapport.diapos if d.code == "DP 11"]

    assert len(notice) == 2
    assert [d.numero for d in notice] == [12, 13]


def test_seules_les_alternatives_portent_un_bandeau(presentation):
    """Posé sur la diapo, donc supprimable sans toucher au fond."""
    for diapositive in presentation.slides:
        bandeaux = [f for f in diapositive.shapes if f.name == "Bandeau à supprimer"]
        alternative = diapositive.slide_layout.name.startswith(("DP 6 —", "DP 8 au"))
        assert bool(bandeaux) is alternative
        if bandeaux:
            assert "supprimez les autres" in bandeaux[0].text_frame.text


def test_les_trois_cadrages_de_dp8_sont_trois_echelles_croissantes(rapport):
    """Le chef de projet garde celle où son point de vue tombe (décision D2)."""
    echelles = [d.echelle for d in rapport.diapos if d.code == "DP 8"]

    assert len(echelles) == 3
    assert echelles == sorted(echelles)
    assert len(set(echelles)) == 3


def test_le_cartouche_annonce_l_echelle_du_cadrage(dossier_projet, rapport):
    """Le dénominateur écrit sur la planche est celui que le cadrage a retenu.

    Mesuré dans le PDF de la planche, pas dans la valeur intermédiaire : c'est le
    texte du cartouche qui trompe l'instruction, pas l'attribut Python.
    """
    from pypdf import PdfReader

    from dp_socle.echelle import formater_echelle

    projet, sortie = dossier_projet
    planches, _ = sortie_pptx._planches_photo(
        projet, "DP 8", 11, (2,), 2, 3,
        _emprise_du_contrat(sortie / projet.nom, projet),
        _contrat(sortie / projet.nom, projet), fond_ign=False,
    )
    assert len(planches) == 3
    for planches_alternative in planches:
        chemin = planches_alternative.planche.rendre_pdf(
            sortie / f"temoin_{planches_alternative.echelle}.pdf"
        )
        texte = PdfReader(str(chemin)).pages[0].extract_text()
        attendu = formater_echelle(planches_alternative.echelle)
        assert _espaces_normalisees(attendu) in _espaces_normalisees(texte)


def _espaces_normalisees(texte: str) -> str:
    """Toutes les espaces ramenées à l'espace ordinaire.

    `formater_echelle` sépare les milliers par une insécable, et l'extraction du
    PDF ne rend pas forcément la même : comparer les codes de caractères ferait
    échouer un cartouche juste.
    """
    for insecable in (" ", " ", " "):
        texte = texte.replace(insecable, " ")
    return texte


def _contrat(dossier, projet):
    from dp_socle.contrat import charger_contrat

    return charger_contrat(dossier, voirie=projet.voirie)


def _emprise_du_contrat(dossier, projet):
    from dp_socle.planches.primitives import union_valide

    return union_valide(_contrat(dossier, projet).geometries("cloture"))


# ---------------------------------------------------------------------------
# Ce que le rapport doit dire
# ---------------------------------------------------------------------------


def test_le_rapport_dit_les_garanties_que_cette_sortie_ne_donne_pas(rapport):
    """Taire ce qui n'est plus vérifié serait le repli silencieux le plus coûteux."""
    for phrase in sortie_pptx.AVERTISSEMENTS_DE_PRINCIPE:
        assert phrase in rapport.avertissements


def test_les_planches_de_reperage_partent_en_vectoriel(rapport):
    """Elles partaient en image, et se voyaient pixellisées à l'écran.

    C'est la conséquence du fond : tant que le plan de repérage portait une image
    du Plan IGN v2, la planche entière était dominée par du raster et la mesure
    de `rendu_pptx` retenait la voie matricielle — 200 dpi, le plafond de
    l'export PowerPoint. Le parcellaire vectoriel l'a fait basculer
    (retour d'usage du 26/09/2026).
    """
    photographiques = [
        diapo for diapo in rapport.diapos if diapo.code in ("DP 6", "DP 7", "DP 8")
    ]

    assert photographiques
    assert all(d.voie == "vectorielle" for d in photographiques), [
        (d.libelle, d.voie) for d in photographiques
    ]


def test_le_rapport_dit_la_voie_et_le_poids_de_chaque_diapo(rapport):
    """Une voie choisie sans qu'on sache pourquoi est une décision illisible."""
    for diapo in rapport.diapos:
        assert diapo.voie in ("vectorielle", "matricielle")
        assert diapo.octets > 0
    assert rapport.taille_mo > 0


# ---------------------------------------------------------------------------
# Une pièce absente se dit à part des remarques (05/10/2026)
# ---------------------------------------------------------------------------


def test_une_piece_absente_se_reconnait_a_sa_formule():
    """Saint-Aubin-sur-Loire : le `.pptx` est parti sans sa coupe.

    Son contrat n'avait pas de profil de terrain — la coupe A-A' n'avait pas été
    relevée —, DP 3 n'a donc pas été produite, et le message qui le disait
    arrivait au milieu des « points à savoir », parmi les avertissements de
    principe que tout dossier porte. Une pièce qui manque n'est pas une remarque
    sur le dossier : c'est un trou dedans, et l'écran la montre à part.

    Ce test tient la formule, qui est ce qui permet de les reconnaître sans
    énumérer les pièces une à une.
    """
    from dp_socle.sortie_pptx import annonce_une_piece_absente

    assert annonce_une_piece_absente(
        "Aucun profil de terrain au contrat : DP 3 n'est pas produite. "
        "Tracez la ligne de coupe A-A' et relevez le profil à l'import."
    )
    assert annonce_une_piece_absente(
        "Aucune prise de vue enregistrée : DP 6, DP 7 et DP 8 ne sont pas produites."
    )
    assert annonce_une_piece_absente("DP 4-3 n'est pas produite : aucun ouvrage.")
    # Et ce qui n'est pas une pièce manquante ne doit pas s'y glisser : les
    # avertissements de principe accompagnent **tous** les dossiers, et les
    # faire remonter en bloquants rendrait le signal muet.
    from dp_socle.sortie_pptx import AVERTISSEMENTS_DE_PRINCIPE

    for message in AVERTISSEMENTS_DE_PRINCIPE:
        assert not annonce_une_piece_absente(message), message


def test_les_deux_producteurs_emploient_la_meme_formule():
    """`sortie_pptx` et `assemblage` annoncent une pièce absente des mêmes mots.

    Sans quoi la sortie PowerPoint signalerait le trou et le PDF non, ou
    l'inverse — et le défaut ne se verrait que sur un dossier réel.
    """
    from pathlib import Path

    from dp_socle.sortie_pptx import MARQUEURS_PIECE_ABSENTE

    racine = Path(__file__).resolve().parent.parent
    for module in ("sortie_pptx.py", "assemblage.py"):
        texte = (racine / "dp_socle" / module).read_text(encoding="utf-8")
        assert any(m in texte for m in MARQUEURS_PIECE_ABSENTE), module
