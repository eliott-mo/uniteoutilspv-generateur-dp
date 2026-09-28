"""Verser la notice dans un dossier déjà fini (lot 7).

Ce qui se mesure ici : qu'un `.pptx` produit par l'outil **dise ce qu'il est**
— son projet, sa pagination — et qu'on puisse y ajouter la notice sans toucher
au travail que le chef de projet y a mis.

Les dossiers d'essai sont fabriqués par les mêmes recettes OOXML que la
génération (`ooxml.nouvelle_mise_en_page`, `ooxml.poser_fond`) : un fixture qui
monterait ses diapos autrement cesserait de mesurer ce que l'outil produit.
"""

from __future__ import annotations

import io

import pytest

pytest.importorskip("cairosvg", reason="CairoSVG est nécessaire pour la page de garde")

from dp_socle.environnement import preparer_cairo  # noqa: E402

preparer_cairo()

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: E402
from pptx.util import Mm  # noqa: E402

from dp_socle import ajout_notice, ooxml  # noqa: E402
from dp_socle.erreurs import ErreurReprise  # noqa: E402
from dp_socle.planche import HAUTEUR_MM, LARGEUR_MM  # noqa: E402
from dp_socle.projet import Projet  # noqa: E402

#: Ce qu'un dossier fini porte, alternatives déjà tranchées par le chef de
#: projet — c'est l'état dans lequel il revient.
LIBELLES = (
    "Page de garde",
    "DP 1-1",
    "DP 1-2",
    "DP 2",
    "DP 6 — 2 cadres",
    "DP 8 au 1/5000",
)


def _projet() -> Projet:
    return Projet(
        nom="PV-ESSAI-IND01",
        commune="ESSAI",
        code_postal="01000",
        date="2026-09-17",
        emprise=None,
        libelle="PV ESSAI IND01",
    )


def _image(couleur=(230, 230, 230)) -> bytes:
    from PIL import Image

    tampon = io.BytesIO()
    Image.new("RGB", (400, 283), couleur).save(tampon, format="PNG")
    return tampon.getvalue()


def _notice_pdf(pages: int = 3) -> bytes:
    """Une notice de synthèse, en autant de pages qu'on veut."""
    import cairosvg

    feuilles = "".join(
        f'<svg width="210mm" height="297mm" viewBox="0 0 210 297" '
        f'xmlns="http://www.w3.org/2000/svg">'
        f'<text x="20" y="40" font-size="8">Notice, page {rang}</text></svg>'
        for rang in range(1, pages + 1)
    )
    # CairoSVG ne fusionne pas : une page par appel, réunies par pypdf.
    from pypdf import PdfReader, PdfWriter

    ecrivain = PdfWriter()
    for feuille in feuilles.split("</svg>")[:-1]:
        tampon = io.BytesIO()
        cairosvg.svg2pdf(bytestring=(feuille + "</svg>").encode(), write_to=tampon)
        ecrivain.append(PdfReader(io.BytesIO(tampon.getvalue())))
    sortie = io.BytesIO()
    ecrivain.write(sortie)
    return sortie.getvalue()


def _dossier_fini(avec_identite: bool = True, libelles=LIBELLES) -> bytes:
    """Un `.pptx` comme la génération en produit, monté par les mêmes recettes."""
    presentation = Presentation()
    presentation.slide_width = Mm(LARGEUR_MM)
    presentation.slide_height = Mm(HAUTEUR_MM)
    for rang, libelle in enumerate(libelles):
        mise_en_page = ooxml.nouvelle_mise_en_page(presentation, libelle)
        ooxml.poser_fond(mise_en_page, _image())
        if rang == 0:
            # La réservation de la perspective, que le chef de projet remplit
            # dans PowerPoint : c'est elle qui doit survivre au redessin.
            ooxml.ajouter_reservation_image(
                mise_en_page, 10, "Cadre photo 1", 20.0, 60.0, 180.0, 120.0
            )
        presentation.slides.add_slide(mise_en_page)
    if avec_identite:
        ajout_notice.ecrire_identite(presentation, _projet())
    tampon = io.BytesIO()
    presentation.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------------------
# Ce que le fichier dit de lui-même
# ---------------------------------------------------------------------------


def test_l_identite_du_projet_survit_a_l_enregistrement():
    """Cent caractères gravés à la génération, relus à la reprise.

    `core_properties` plafonne à 255 caractères par champ : le `projet.json`
    entier, 2 000 octets, ne passe pas. Ces quatre champs-là suffisent à
    redessiner la page de garde, et c'est tout ce que la reprise demande.
    """
    presentation = Presentation(io.BytesIO(_dossier_fini()))
    identite = ajout_notice.lire_identite(presentation)

    assert identite == {
        "libelle": "PV ESSAI IND01",
        "commune": "ESSAI",
        "code_postal": "01000",
        "date": "2026-09-17",
    }
    charge = presentation.core_properties.comments
    assert len(charge) <= 255, f"{len(charge)} caractères, le champ en tient 255"


def test_un_fichier_sans_identite_est_refuse_en_le_disant():
    """Un dossier produit avant ce lot ne peut pas être complété.

    Le refus renvoie au parcours normal plutôt que de demander au chef de
    projet de ressaisir commune et date : une date retapée de travers donnerait
    une couverture qui ment sur ses propres planches.
    """
    with pytest.raises(ErreurReprise, match="identité de son projet"):
        ajout_notice.lire_le_dossier(_dossier_fini(avec_identite=False))


@pytest.mark.parametrize(
    "nom, code",
    [
        ("Page de garde", ""),
        ("DP 1-1", "DP 1-1"),
        ("DP 6 — 2 cadres", "DP 6"),
        ("DP 8 au 1/5000", "DP 8"),
        ("DP 11, page 3", "DP 11"),
    ],
)
def test_le_code_de_la_piece_se_lit_dans_le_nom_de_la_mise_en_page(nom, code):
    """Les alternatives et les pages d'une même pièce rendent le code nu.

    C'est ce nom-là qui porte la pagination : le serveur ne garde rien, et le
    fichier déposé est la seule source. Relevé le 28/09/2026 sur un dossier
    revenu de PowerPoint — les noms survivent à l'enregistrement.
    """
    assert ajout_notice.code_de_la_mise_en_page(nom) == code


def test_la_pagination_lue_est_celle_du_fichier_depose():
    """Les diapos supprimées par le chef de projet ont disparu de la pagination.

    C'est tout l'intérêt de lire le fichier plutôt qu'une mémoire : il porte
    l'état réel du dossier, alternatives déjà tranchées.
    """
    dossier = ajout_notice.lire_le_dossier(_dossier_fini())

    assert dossier.nb_diapos == 6
    assert dossier.pagination() == [
        ("", "Page de garde", 1),
        ("DP 1-1", "Plan de situation", 2),
        ("DP 1-2", "Photo aérienne", 3),
        ("DP 2", "Plan de masse", 4),
        ("DP 6", "Insertions paysagères", 5),
        ("DP 8", "Photographie paysage lointain", 6),
    ]
    assert not dossier.porte_la_notice
    assert dossier.premiere_page_de_la_notice == 7


# ---------------------------------------------------------------------------
# L'ajout lui-même
# ---------------------------------------------------------------------------


def test_la_notice_s_ajoute_sans_toucher_au_travail_du_chef_de_projet(tmp_path):
    """Le geste entier : trois pages ajoutées, le reste intact.

    Ce qui est vérifié en propre, parce que c'est ce qui serait perdu : la
    réservation d'image de la page de garde — où le chef de projet dépose sa
    perspective — survit au redessin du sommaire. Le fond est sur la mise en
    page, la réservation sur la diapo : échanger l'un ne touche pas l'autre.
    """
    fini = _dossier_fini()
    notice = tmp_path / "notice.pdf"
    notice.write_bytes(_notice_pdf(pages=3))

    complet = ajout_notice.ajouter_la_notice(fini, notice, tmp_path)

    relu = Presentation(io.BytesIO(complet))
    noms = [diapo.slide_layout.name for diapo in relu.slides]
    assert noms == list(LIBELLES) + [
        "DP 11, page 1", "DP 11, page 2", "DP 11, page 3",
    ]
    # La perspective du chef de projet est toujours là. La forme prend le nom
    # que PowerPoint donne à une réservation héritée — « Picture Placeholder 1 »
    # et non celui de la mise en page ; c'est son type qui l'identifie, et c'est
    # bien ce qu'on retrouve dans les fichiers revenus des chefs de projet.
    formes = list(relu.slides[0].shapes)
    assert len(formes) == 1
    assert formes[0].is_placeholder
    assert formes[0].placeholder_format.type == PP_PLACEHOLDER.PICTURE
    # Et l'identité voyage encore : un second ajout serait possible.
    assert ajout_notice.lire_identite(relu) is not None


def test_le_sommaire_de_la_couverture_porte_la_notice(tmp_path):
    """La page de garde est redessinée, et c'est la seule.

    Son sommaire est dans l'image : sans redessin, la ligne DP 11 resterait à
    « — » alors que la notice est là, et le dossier se contredirait.
    """
    fini = _dossier_fini()
    notice = tmp_path / "notice.pdf"
    notice.write_bytes(_notice_pdf(pages=2))

    avant = Presentation(io.BytesIO(fini)).slides[0].slide_layout
    fond_avant = next(
        forme.image.blob for forme in avant.shapes if forme.shape_type == 13
    )

    complet = ajout_notice.ajouter_la_notice(fini, notice, tmp_path)
    apres = Presentation(io.BytesIO(complet)).slides[0].slide_layout
    fonds = [forme for forme in apres.shapes if forme.shape_type == 13]

    assert len(fonds) == 1, "le fond précédent n'a pas été retiré"
    assert fonds[0].image.blob != fond_avant, "la couverture n'a pas été redessinée"


def test_un_dossier_qui_porte_deja_sa_notice_est_refuse(tmp_path):
    """Pas de seconde notice ajoutée en silence.

    Remplacer une notice est un autre geste que l'ajouter, et le refus le dit
    plutôt que d'empiler les pages.
    """
    fini = _dossier_fini(libelles=LIBELLES + ("DP 11, page 1", "DP 11, page 2"))
    notice = tmp_path / "notice.pdf"
    notice.write_bytes(_notice_pdf(pages=1))

    with pytest.raises(ErreurReprise, match="déjà une notice de 2 page"):
        ajout_notice.ajouter_la_notice(fini, notice, tmp_path)


def test_une_notice_vide_ne_produit_rien(tmp_path):
    """Un PDF illisible ne doit pas rendre un dossier inchangé sans le dire."""
    notice = tmp_path / "notice.pdf"
    notice.write_bytes(b"%PDF-1.4 ce n'est pas un PDF")

    with pytest.raises(Exception):
        ajout_notice.ajouter_la_notice(_dossier_fini(), notice, tmp_path)
