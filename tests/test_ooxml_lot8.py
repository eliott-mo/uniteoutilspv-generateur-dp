"""Les six recettes OOXML du lot 8, mesurées dans le fichier écrit.

Le critère n'est pas « le `.pptx` s'écrit » mais « il porte ce que le montage
voulait ». Chaque mesure se fait donc sur un fichier **enregistré puis relu**, et
non sur l'objet en mémoire : c'est l'enregistrement qui sérialise l'XML, et c'est
le fichier relu qui dit ce que PowerPoint lira.

Ce que ces mesures ne prouvent pas : que PowerPoint l'ouvre correctement. Cela
s'est vérifié à l'écran les 25 et 26/09/2026, et se revérifie à chaque
changement du montage — le brief du lot en tient la liste.

LE PIÈGE QUE CE FICHIER SURVEILLE
----------------------------------
`python-pptx` accepte sans lever qu'on lui règle une propriété qui n'existe pas :
`fill.transparency = 0.45` ne faisait rien, et le cône sortait opaque. Les tests
d'alpha et de verrou ne vérifient donc pas un appel, mais la présence de
l'attribut **dans l'XML du fichier**.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Mm

from dp_socle import ooxml
from dp_socle.erreurs import ErreurMontagePPTX
from dp_socle.planche import HAUTEUR_MM, LARGEUR_MM

#: Un EMU vaut 1/36 000 de millimètre. La conversion inverse de `pptx.util.Mm`,
#: écrite ici pour que les mesures se lisent en millimètres.
EMU_PAR_MM = 36000.0


def _mm(emu) -> float:
    return emu / EMU_PAR_MM


@pytest.fixture
def fond(tmp_path) -> Path:
    """Une image de fond quelconque : sa taille n'entre dans aucune mesure.

    Le fond est posé pleine page, aux dimensions de la diapo et non aux siennes.
    """
    chemin = tmp_path / "fond.png"
    Image.new("RGB", (60, 42), (240, 240, 240)).save(chemin)
    return chemin


@pytest.fixture
def presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(LARGEUR_MM), Mm(HAUTEUR_MM)
    return prs


def _enregistrer_et_relire(prs, tmp_path, nom="sonde.pptx"):
    chemin = tmp_path / nom
    prs.save(str(chemin))
    return chemin, Presentation(str(chemin))


# ---------------------------------------------------------------------------
# Les mises en page, une par diapo
# ---------------------------------------------------------------------------


def test_une_mise_en_page_de_plus_survit_a_l_enregistrement(presentation, tmp_path):
    """Le masque doit lister la mise en page, sinon PowerPoint répare le fichier."""
    depart = len(presentation.slide_layouts)
    for nom in ("DP 2", "DP 7", "DP 8 au 1/5 000"):
        ooxml.nouvelle_mise_en_page(presentation, nom)

    _, relu = _enregistrer_et_relire(presentation, tmp_path)

    assert len(relu.slide_layouts) == depart + 3
    assert relu.slide_layouts.get_by_name("DP 8 au 1/5 000") is not None
    # Le masque les liste, et chaque entrée porte un identifiant distinct : deux
    # `p:sldLayoutId` de même `id` font un fichier que PowerPoint répare.
    entrees = relu.slide_master.element.get_or_add_sldLayoutIdLst().sldLayoutId_lst
    assert len(entrees) == depart + 3
    identifiants = [e.get("id") for e in entrees]
    assert len(set(identifiants)) == len(identifiants)


def test_la_mise_en_page_neuve_ne_porte_aucune_reservation_heritee(presentation):
    """Les `idx` 10 à 12 du modèle « Blank » doivent être libérés.

    C'est la mesure du 25/09/2026 : la case de date occupe l'`idx` 10, et un
    cadre photo posé à cet `idx` héritait de ses cotes — 59 x 10 mm au lieu de
    222 x 115 mm. Une géométrie fausse qui s'affiche correctement.
    """
    assert len(presentation.slide_layouts[6].placeholders) == 3

    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")

    assert len(mise_en_page.placeholders) == 0
    assert len(mise_en_page.shapes) == 0


# ---------------------------------------------------------------------------
# Le fond de planche
# ---------------------------------------------------------------------------


def test_le_fond_couvre_la_diapo_entiere_et_reste_sur_la_mise_en_page(
    presentation, fond, tmp_path
):
    """Le fond est posé pleine page sur la mise en page, pas sur la diapo.

    C'est ce qui le met hors d'atteinte du chef de projet, et donc ce qui protège
    l'échelle : une planche redimensionnée rend faux le dénominateur de son
    cartouche.
    """
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 2")
    ooxml.poser_fond(mise_en_page, fond)
    presentation.slides.add_slide(mise_en_page)

    _, relu = _enregistrer_et_relire(presentation, tmp_path)
    diapo = relu.slides[0]

    assert len(diapo.shapes) == 0, "rien ne doit être posé sur la diapo elle-même"
    images = [f for f in diapo.slide_layout.shapes if f.shape_type is not None]
    assert len(images) == 1
    image = images[0]
    assert (_mm(image.left), _mm(image.top)) == (0.0, 0.0)
    assert (_mm(image.width), _mm(image.height)) == (LARGEUR_MM, HAUTEUR_MM)


def test_le_svg_voyage_a_cote_du_matriciel(presentation, fond, tmp_path):
    """La voie vectorielle : un `asvg:svgBlip` dans l'extension du `a:blip`.

    Mesuré dans le paquet et non par l'API : `python-pptx` ne connaît pas ce type
    de contenu, et c'est précisément pour cela que l'extension s'écrit à la main.
    """
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 2")
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="420" height="297"/>'
    ooxml.poser_fond(mise_en_page, fond, svg=svg)

    chemin, _ = _enregistrer_et_relire(presentation, tmp_path)

    with zipfile.ZipFile(chemin) as paquet:
        svgs = [n for n in paquet.namelist() if n.endswith(".svg")]
        assert len(svgs) == 1
        assert paquet.read(svgs[0]) == svg
        # Le type de contenu doit être déclaré, sinon PowerPoint refuse le paquet.
        types = paquet.read("[Content_Types].xml").decode("utf-8")
        assert "image/svg+xml" in types
        mises_en_page = [
            n for n in paquet.namelist() if n.startswith("ppt/slideLayouts/slideLayout")
        ]
        porteuses = [
            n for n in mises_en_page
            if ooxml.URI_EXTENSION_SVG in paquet.read(n).decode("utf-8")
        ]
        assert len(porteuses) == 1
        xml = paquet.read(porteuses[0]).decode("utf-8")
        assert "asvg:svgBlip" in xml
        # Le matriciel de repli reste dans le `a:blip` : c'est ce que voit une
        # version de PowerPoint antérieure à 2016.
        assert "<a:blip" in xml and "r:embed" in xml


def test_le_fond_ne_se_selectionne_pas(presentation, fond, tmp_path):
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 2")
    ooxml.poser_fond(mise_en_page, fond)

    chemin, _ = _enregistrer_et_relire(presentation, tmp_path)

    xml = _xml_de_la_mise_en_page_du_fond(chemin)
    assert 'noSelect="1"' in xml
    assert 'noMove="1"' in xml and 'noResize="1"' in xml


def _xml_de_la_mise_en_page_du_fond(chemin: Path) -> str:
    """L'XML de la seule mise en page qui porte une image.

    Le rang de la partie n'est pas nommé en dur : il dépend du nombre de mises en
    page du modèle, qui appartient à `python-pptx` et non à nous.
    """
    with zipfile.ZipFile(chemin) as paquet:
        porteuses = [
            paquet.read(nom).decode("utf-8")
            for nom in paquet.namelist()
            if nom.startswith("ppt/slideLayouts/slideLayout")
            and nom.endswith(".xml")
            and "<p:pic>" in paquet.read(nom).decode("utf-8")
        ]
    assert len(porteuses) == 1
    return porteuses[0]


# ---------------------------------------------------------------------------
# Les réservations d'image
# ---------------------------------------------------------------------------


#: Cotes des deux cadres d'une planche photographique à deux emplacements,
#: relevées par la sonde du 25/09/2026 et retrouvées par `photographies`.
CADRES_DEUX_EMPLACEMENTS = (
    (185.5, 18.1, 222.0, 115.4),
    (185.5, 150.1, 222.0, 115.4),
)


def test_la_diapo_herite_des_cotes_de_ses_cadres(presentation, tmp_path):
    """La diapo porte les réservations de la mise en page, aux mêmes cotes.

    Le clonage ne reprend que l'identifiant et le type : les cotes sont
    **héritées**, et c'est ce qui garantit qu'un cadre ne peut pas dériver d'une
    diapo à l'autre.
    """
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    for rang, cotes in enumerate(CADRES_DEUX_EMPLACEMENTS):
        ooxml.ajouter_reservation_image(mise_en_page, 10 + rang, f"Cadre {rang}", *cotes)
    presentation.slides.add_slide(mise_en_page)

    _, relu = _enregistrer_et_relire(presentation, tmp_path)
    cadres = list(relu.slides[0].placeholders)

    assert len(cadres) == len(CADRES_DEUX_EMPLACEMENTS)
    for cadre, attendues in zip(cadres, CADRES_DEUX_EMPLACEMENTS):
        mesurees = (
            _mm(cadre.left), _mm(cadre.top), _mm(cadre.width), _mm(cadre.height),
        )
        assert mesurees == pytest.approx(attendues, abs=0.05)
        # Le rapport du cadre est ce qui impose le rognage de la photographie :
        # 1,92:1 à deux emplacements.
        assert cadre.width / cadre.height == pytest.approx(1.924, abs=0.005)


def test_un_idx_deja_pris_est_refuse(presentation):
    """Deux réservations du même `idx` donneraient à l'une les cotes de l'autre."""
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    ooxml.ajouter_reservation_image(mise_en_page, 10, "Cadre 1", 185.5, 18.1, 222.0, 115.4)

    with pytest.raises(ErreurMontagePPTX, match="idx 10"):
        ooxml.ajouter_reservation_image(
            mise_en_page, 10, "Cadre 2", 185.5, 150.1, 222.0, 115.4
        )


def test_la_photographie_deposee_entre_dans_le_cadre_sans_le_deformer(
    presentation, tmp_path
):
    """Une réservation d'image rogne la photographie au format du cadre.

    C'est le défaut que le montage corrige : une image simplement posée, quand on
    la remplace, étire son cadre au format de l'image — une photographie en
    portrait cassait la planche (mesuré le 25/09/2026). Ici c'est PowerPoint qui
    rogne, et le test le mesure par `insert_picture`, qui applique la même règle.
    """
    portrait = tmp_path / "portrait.jpg"
    Image.new("RGB", (1200, 1600), (120, 140, 160)).save(portrait)
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    ooxml.ajouter_reservation_image(
        mise_en_page, 10, "Cadre PC7-1", *CADRES_DEUX_EMPLACEMENTS[0]
    )
    diapo = presentation.slides.add_slide(mise_en_page)

    diapo.placeholders[10].insert_picture(str(portrait))

    _, relu = _enregistrer_et_relire(presentation, tmp_path)
    image = relu.slides[0].shapes[0]
    attendues = CADRES_DEUX_EMPLACEMENTS[0]
    assert _mm(image.width) == pytest.approx(attendues[2], abs=0.05)
    assert _mm(image.height) == pytest.approx(attendues[3], abs=0.05)
    # Rognée sur un seul axe : une image en portrait perd de la hauteur, jamais
    # de la largeur.
    assert image.crop_left == pytest.approx(0.0, abs=1e-6)
    assert image.crop_top > 0.0


# ---------------------------------------------------------------------------
# Le verrou, l'alpha, l'étiquette droite
# ---------------------------------------------------------------------------


def test_le_verrou_se_relit_dans_la_diapo(presentation, tmp_path):
    """`a:spLocks` se pose sur la forme de la diapo, et s'y retrouve."""
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    diapo = presentation.slides.add_slide(mise_en_page)
    titre = diapo.shapes.add_textbox(Mm(10), Mm(10), Mm(60), Mm(8))
    ooxml.verrouiller(titre)

    chemin, _ = _enregistrer_et_relire(presentation, tmp_path)

    with zipfile.ZipFile(chemin) as paquet:
        xml = paquet.read("ppt/slides/slide1.xml").decode("utf-8")
    assert 'noMove="1"' in xml and 'noResize="1"' in xml
    # Le titre de cadre reste sélectionnable : le chef de projet doit pouvoir le
    # récrire, c'est ce que D1 lui promet.
    assert "noSelect" not in xml


def test_l_alpha_s_ecrit_dans_le_remplissage(presentation, tmp_path):
    """Le cône doit laisser voir le plan sous lui.

    `fill.transparency` n'existe pas : le réglage se mesure dans l'XML, seul
    endroit où son absence se verrait.
    """
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    diapo = presentation.slides.add_slide(mise_en_page)
    cone = diapo.shapes.add_shape(MSO_SHAPE.PIE, Mm(50), Mm(50), Mm(16), Mm(16))
    cone.fill.solid()
    cone.fill.fore_color.rgb = RGBColor.from_string("D81B8C")
    ooxml.regler_alpha(cone, 0.55)

    chemin, _ = _enregistrer_et_relire(presentation, tmp_path)

    with zipfile.ZipFile(chemin) as paquet:
        xml = paquet.read("ppt/slides/slide1.xml").decode("utf-8")
    assert '<a:alpha val="55000"/>' in xml


def test_un_remplissage_absent_est_refuse_plutot_que_laisse_opaque(
    presentation,
):
    """Sans `a:solidFill`, l'alpha n'a nulle part à s'écrire — et on le dit."""
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    diapo = presentation.slides.add_slide(mise_en_page)
    forme = diapo.shapes.add_shape(MSO_SHAPE.PIE, Mm(50), Mm(50), Mm(16), Mm(16))
    forme.fill.background()

    with pytest.raises(ErreurMontagePPTX, match="remplissage uni"):
        ooxml.regler_alpha(forme, 0.55)


@pytest.mark.parametrize("opacite", [0.0, -0.2, 1.5])
def test_une_opacite_hors_bornes_est_refusee(presentation, opacite):
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    diapo = presentation.slides.add_slide(mise_en_page)
    forme = diapo.shapes.add_shape(MSO_SHAPE.PIE, Mm(50), Mm(50), Mm(16), Mm(16))
    forme.fill.solid()

    with pytest.raises(ErreurMontagePPTX, match="Opacité"):
        ooxml.regler_alpha(forme, opacite)


def test_l_etiquette_reste_horizontale_dans_un_groupe_qui_tourne(
    presentation, tmp_path
):
    """`upright` **et** `wrap="none"` : le premier seul replie le texte.

    Mesuré le 25/09/2026 : avec `upright` seul, l'étiquette d'un cône tourné de
    côté se replie à une lettre par ligne — la boîte tourne, et le texte
    horizontal n'y tient plus.
    """
    mise_en_page = ooxml.nouvelle_mise_en_page(presentation, "DP 7")
    diapo = presentation.slides.add_slide(mise_en_page)
    etiquette = diapo.shapes.add_textbox(Mm(50), Mm(50), Mm(20), Mm(5))
    etiquette.text_frame.text = "PC7-1"
    ooxml.garder_horizontal(etiquette)

    chemin, _ = _enregistrer_et_relire(presentation, tmp_path)

    with zipfile.ZipFile(chemin) as paquet:
        xml = paquet.read("ppt/slides/slide1.xml").decode("utf-8")
    assert 'upright="1"' in xml
    assert 'wrap="none"' in xml
